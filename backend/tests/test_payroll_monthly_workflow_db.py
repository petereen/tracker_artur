"""End-to-end monthly payroll workflow against a disposable PostgreSQL database.

Runs only when PAYROLL_TEST_DATABASE_URL points at a throwaway database
(for example ``postgresql+asyncpg://payroll@127.0.0.1:5439/payroll_test``).
It creates just the payroll tables and their foreign-key dependencies, then
drives the §18 acceptance flow through the HTTP API: advance → final →
sequencing → edits → import → approval → export → close → archive → unlock.
"""

from __future__ import annotations

import asyncio
import os
from contextlib import asynccontextmanager
from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest

DATABASE_URL = os.environ.get("PAYROLL_TEST_DATABASE_URL")
pytestmark = pytest.mark.skipif(not DATABASE_URL, reason="PAYROLL_TEST_DATABASE_URL is not set")

if DATABASE_URL:
    os.environ.setdefault("DATABASE_URL", DATABASE_URL)
    os.environ.setdefault("SYNC_DATABASE_URL", DATABASE_URL.replace("+asyncpg", "+psycopg2"))
    os.environ.setdefault("SECRET_KEY", "payroll-workflow-test-secret-key-0123456789")


D = Decimal


def _payroll_tables():
    from app.models import models as m

    seeds = [
        m.AttendanceLog, m.Department, m.Employee, m.EmployeeBankAccount, m.EmployeeDetails, m.HolidayRecord, m.MonthlyPayrollArchive,
        m.MonthlyPayrollCalendarDay, m.MonthlyPayrollCompanySettings, m.MonthlyPayrollMonth, m.MonthlyPayrollProfile, m.MonthlyPayrollRowAudit,
        m.MonthlyPayrollRuleSet, m.MonthlyPayrollRun, m.MonthlyPayrollRunRow, m.ERPAccount, m.MonthlyPayrollSalaryHistory, m.Organization,
        m.Schedule, m.TimeOff, m.WorkTimeEntry,
    ]
    tables = {model.__table__ for model in seeds}
    grown = True
    while grown:
        grown = False
        for table in list(tables):
            for foreign_key in table.foreign_keys:
                if foreign_key.column.table not in tables:
                    tables.add(foreign_key.column.table)
                    grown = True
    return list(tables)


@asynccontextmanager
async def _payroll_api(monkeypatch):
    from fastapi import FastAPI
    from httpx import ASGITransport, AsyncClient
    from sqlalchemy import text
    from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

    from app.core.database import get_db
    from app.core.enterprise_deps import ActorContext, get_actor
    from app.models import models as m
    from app.payroll import monthly_workflow
    from app.services.secret_box import encrypt_secret

    engine = create_async_engine(DATABASE_URL)
    tables = _payroll_tables()
    async with engine.begin() as connection:
        # Only this test's tables are reset; CASCADE covers the employees ↔ user_accounts FK cycle.
        await connection.execute(text("DROP TABLE IF EXISTS " + ", ".join(f'"{table.name}"' for table in tables) + " CASCADE"))
        await connection.run_sync(lambda sync: m.Base.metadata.create_all(sync, tables=tables))
    sessions = async_sessionmaker(engine, expire_on_commit=False)

    async with sessions() as db:
        organization = m.Organization(name="Оюунс")
        db.add(organization)
        await db.flush()
        account = m.UserAccount(organization_id=organization.id, email="payroll@test.mn", password_hash="x")
        finance = m.Department(organization_id=organization.id, code="FIN", name="Санхүү")
        sales = m.Department(organization_id=organization.id, code="SAL", name="Борлуулалт")
        db.add_all([account, finance, sales])
        await db.flush()

        async def worker(name, department, *, start=date(2025, 1, 1), profile=None, salary=None, salary_from=date(2025, 1, 1), bank=False):
            employee = m.Employee(organization_id=organization.id, name=name)
            db.add(employee)
            await db.flush()
            db.add(m.EmployeeDetails(organization_id=organization.id, employee_id=employee.id, department_id=department.id, job_title="Мэргэжилтэн", start_date=start))
            if profile is not None:
                row = m.MonthlyPayrollProfile(organization_id=organization.id, employee_id=employee.id, **profile)
                db.add(row)
                await db.flush()
                db.add(m.MonthlyPayrollSalaryHistory(profile_id=row.id, monthly_salary=salary, valid_from=salary_from))
            if bank:
                db.add(m.EmployeeBankAccount(employee_id=employee.id, bank_code="KHAN", account_number_ciphertext=encrypt_secret("5000123456"), account_fingerprint="fp-" + name, account_last4="3456", valid_from=date(2025, 1, 1), is_primary=True))
            return employee.id

        ids = {
            # §18 row: BIWEEKLY 10/25, 40% advance, meal + commute 120,000.
            "oyun": await worker("Бат Оюун-Эрдэнэ", finance, bank=True, salary=D("1500000"), profile={"salary_type": "PRORATION", "meal_allowance": D("70000"), "commute_allowance": D("50000"), "allowance_basis": "MONTHLY", "payment_frequency": "BIWEEKLY", "pay_days": [10, 25], "advance_basis": "PERCENT", "advance_values": ["40"], "tax_relief_eligible": True}),
            # WEEKLY FIXED instalments must not net each other off.
            "weekly": await worker("Дорж Сараа", sales, salary=D("1500000"), profile={"salary_type": "FIXED", "payment_frequency": "WEEKLY", "pay_days": [7, 14, 21, 28], "advance_basis": "FIXED", "advance_values": ["200000", "200000", "200000"], "tax_relief_eligible": True}),
            # Hired mid-month: salary history starts on the hire date.
            "hire": await worker("Ганаа Тэмүүлэн", sales, start=date(2026, 8, 17), salary=D("1500000"), salary_from=date(2026, 8, 17), profile={"salary_type": "FIXED", "meal_allowance": D("10000"), "commute_allowance": D("5000"), "allowance_basis": "FIXED", "payment_frequency": "MONTHLY", "pay_days": [25], "tax_relief_eligible": True}),
            # No payroll profile yet: must surface as a blocking row, not vanish.
            "missing": await worker("Профайлгүй Ажилтан", finance),
        }
        await db.commit()
        actor = ActorContext(account_id=account.id, organization_id=organization.id, employee_id=None, email="payroll@test.mn", locale="mn", roles=frozenset({"admin"}))

    async def no_capability_check(*args, **kwargs):
        return None

    async def no_event(*args, **kwargs):
        return None

    monkeypatch.setattr(monthly_workflow, "require_capability", no_capability_check)
    monkeypatch.setattr(monthly_workflow, "record_change", no_event)

    async def test_db():
        async with sessions() as session:
            yield session

    app = FastAPI()
    app.include_router(monthly_workflow.router, prefix="/m")
    app.dependency_overrides[get_db] = test_db
    app.dependency_overrides[get_actor] = lambda: actor
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        yield client, ids, sessions, organization.id
    await engine.dispose()


def _row(run: dict, employee_id: int) -> dict:
    return next(row for row in run["rows"] if row["employee_id"] == employee_id)


async def _ok(response, status: int = 200):
    assert response.status_code == status, response.text
    return response.json()


def test_monthly_payroll_end_to_end(monkeypatch):
    asyncio.run(_end_to_end(monkeypatch))


async def _end_to_end(monkeypatch):
    async with _payroll_api(monkeypatch) as (client, ids, sessions, organization_id):
        await _scenario(client, ids, sessions, organization_id)


async def _scenario(client, ids, sessions, organization_id):

    # Opening a month seeds the default 2026 rule set for a new organization.
    month = await _ok(await client.post("/m/months", json={"year": 2026, "month": 8}), 201)
    assert month["rule_snapshot"]["minimum_wage"] in ("792000", "792000.0000")
    dates = await _ok(await client.get(f"/m/months/{month['id']}/advance-dates"))
    assert {item["day"]: item["workers"] for item in dates} == {7: 1, 10: 1, 14: 1, 21: 1}

    # Sequencing: a final run before any approved advance shows «Урьдчилгаа бодоогүй» (a warning, not a block).
    final = await _ok(await client.post(f"/m/months/{month['id']}/runs", json={"run_type": "final", "pay_date": "2026-08-31"}), 201)
    final_run = await _ok(await client.get(f"/m/runs/{final['id']}"))
    assert len(final_run["rows"]) == 4
    assert "advance_not_calculated" in _row(final_run, ids["oyun"])["warnings"]
    assert _row(final_run, ids["missing"])["warnings"] == ["profile_missing"]
    hire = _row(final_run, ids["hire"])
    assert hire["profile"]["complete"] is True and hire["result"]["base_pay"] == "785714"
    # FIXED daily allowance: 15,000 × 11 planned workdays from the 17th.
    assert (hire["result"]["allowance_days"], hire["result"]["meal_commute"]) == ("11", "165000")

    # Advance run on the 10th: 40% of 1,500,000. Approve one row, then «Бүгдийг батлах».
    advance10 = await _ok(await client.post(f"/m/months/{month['id']}/runs", json={"run_type": "advance", "pay_date": "2026-08-10"}), 201)
    run10 = await _ok(await client.get(f"/m/runs/{advance10['id']}"))
    assert [row["employee_id"] for row in run10["rows"]] == [ids["oyun"]]
    assert _row(run10, ids["oyun"])["result"]["advance"] == "600000"
    # Reference column: a full month on the profile, not hours worked so far:
    # gross 1,620,000 − НДШ 186,300 − ХХОАТ (143,370 − 16,000 relief) = 1,306,330.
    assert _row(run10, ids["oyun"])["result"]["estimated_net"] == "1306330"
    # No time data yet → the projection assumes the planned month (168 h):
    # Суутгалын дүн = 600,000 урьдчилгаа + 186,300 НДШ + 127,370 ХХОАТ.
    projection = _row(run10, ids["oyun"])["result"]["projection"]
    assert _row(run10, ids["oyun"])["inputs"]["worked_normal_hours"] == "168"
    assert (projection["base_pay"], projection["meal_commute"], projection["gross"]) == ("1500000", "120000", "1620000")
    assert (projection["employee_shi"], projection["pit"], projection["advance"]) == ("186300", "127370", "600000")
    assert (projection["total_deductions"], projection["net_pay"]) == ("913670", "706330")
    await _ok(await client.post(f"/m/runs/{advance10['id']}/rows/{ids['oyun']}/approve"))
    approved = await _ok(await client.post(f"/m/runs/{advance10['id']}/approve"))
    assert approved["status"] == "approved" and approved["approval_summary"]["skipped"] == []

    # WEEKLY FIXED instalments: the second one is not reduced by the first.
    for pay_day in ("2026-08-07", "2026-08-14"):
        advance = await _ok(await client.post(f"/m/months/{month['id']}/runs", json={"run_type": "advance", "pay_date": pay_day}), 201)
        run = await _ok(await client.get(f"/m/runs/{advance['id']}"))
        assert _row(run, ids["weekly"])["result"]["advance"] == "200000"
        assert (await _ok(await client.post(f"/m/runs/{advance['id']}/approve")))["status"] == "approved"

    # One-off advance for a MONTHLY worker via «Ажилтан нэмэх» no longer crashes.
    advance21 = await _ok(await client.post(f"/m/months/{month['id']}/runs", json={"run_type": "advance", "pay_date": "2026-08-21"}), 201)
    added = await _ok(await client.post(f"/m/runs/{advance21['id']}/add-workers", json={"employee_ids": [ids["hire"]]}))
    one_off = _row(added["run"], ids["hire"])
    assert one_off["inputs"]["one_off_advance"] is True and "advance_not_due" not in one_off["warnings"]

    # The approved advances changed after the final run was created → «Урьдчилгаа өөрчлөгдсөн» until re-pulled.
    final_run = await _ok(await client.get(f"/m/runs/{final['id']}"))
    assert "advance_changed" in _row(final_run, ids["oyun"])["warnings"]
    final_run = await _ok(await client.post(f"/m/runs/{final['id']}/refresh-advances"))
    oyun = _row(final_run, ids["oyun"])
    assert "advance_changed" not in oyun["warnings"] and oyun["result"]["advance"] == "600000"
    assert [line["amount"] for line in oyun["result"]["advance_lines"]] == ["600000"]
    assert _row(final_run, ids["weekly"])["result"]["advance"] == "400000"

    # §18 Excel row: 168/168 hours + 168,000 leave pay → Q 1,788,000, then a 50,000 penalty after tax.
    row = await _ok(await client.put(f"/m/runs/{final['id']}/rows/{ids['oyun']}", json={"worked_normal_hours": "168", "leave_pay": "168000", "reason": "Excel мөр"}))
    result = row["result"]
    assert (result["gross"], result["employee_shi"], result["taxable_income"], result["pit_before_relief"], result["relief"], result["pit"]) == ("1788000", "205620", "1582380", "158238", "14000", "144238")
    assert (result["total_deductions"], result["net_pay"], result["employer_shi"]) == ("949858", "838142", "223500")
    row = await _ok(await client.put(f"/m/runs/{final['id']}/rows/{ids['oyun']}", json={"other_deductions": [{"type": "Торгууль", "amount": "50000", "note": "хоцролт"}], "reason": "Торгууль"}))
    assert (row["result"]["employee_shi"], row["result"]["pit"], row["result"]["total_deductions"], row["result"]["net_pay"]) == ("205620", "144238", "999858", "788142")

    # Excel import: only filled cells change; blank cells keep the register's values.
    from openpyxl import load_workbook
    template = await client.get(f"/m/runs/{final['id']}/import-template")
    assert template.status_code == 200
    workbook = load_workbook(BytesIO(template.content))
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    target = next(row for row in sheet.iter_rows(min_row=3) if row[0].value == ids["weekly"])
    target[headers.index("bonus")].value = 100000
    target[headers.index("reason")].value = "Урамшуулал"
    buffer = BytesIO()
    workbook.save(buffer)
    imported = await _ok(await client.post(f"/m/runs/{final['id']}/import-xlsx", files={"file": ("inputs.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}))
    assert imported["updated_rows"] == 1
    final_run = await _ok(await client.get(f"/m/runs/{final['id']}"))
    assert _row(final_run, ids["weekly"])["inputs"]["bonus"] == "100000"
    assert _row(final_run, ids["oyun"])["inputs"]["worked_normal_hours"] == "168"

    # Partial «Бүгдийг батлах»: the profile-less worker stays in review; the run stays draft.
    outcome = await _ok(await client.post(f"/m/runs/{final['id']}/approve"))
    assert outcome["status"] == "draft"
    assert [item["employee_id"] for item in outcome["approval_summary"]["skipped"]] == [ids["missing"]]

    # HR fills the profile; sync parks it as «HR changed» until the accountant accepts.
    from app.models import models as m
    async with sessions() as db:
        profile = m.MonthlyPayrollProfile(organization_id=organization_id, employee_id=ids["missing"], salary_type="FIXED", payment_frequency="MONTHLY", pay_days=[25], tax_relief_eligible=True)
        db.add(profile)
        await db.flush()
        db.add(m.MonthlyPayrollSalaryHistory(profile_id=profile.id, monthly_salary=D("900000"), valid_from=date(2025, 1, 1)))
        await db.commit()
    synced = await _ok(await client.post(f"/m/runs/{final['id']}/sync-workers"))
    assert synced["hr_changed_rows"] == 1 and "hr_changed" in _row(synced["run"], ids["missing"])["warnings"]
    accepted = await _ok(await client.post(f"/m/runs/{final['id']}/rows/{ids['missing']}/accept-hr"))
    assert accepted["profile"]["complete"] is True and accepted["result"]["gross"] == "900000"
    # Editing an approved row returns it to draft; «Батлалт цуцлах» works per row.
    approved_row = await _ok(await client.post(f"/m/runs/{final['id']}/rows/{ids['hire']}/unapprove"))
    assert approved_row["status"] == "draft"
    final_state = await _ok(await client.post(f"/m/runs/{final['id']}/approve"))
    assert final_state["status"] == "approved" and final_state["approval_summary"]["skipped"] == []

    # Export: Excel layout sheets with the approved numbers.
    export = await client.get(f"/m/runs/{final['id']}/export")
    assert export.status_code == 200
    exported = load_workbook(BytesIO(export.content))
    assert exported.sheetnames == ["Цалингийн хүснэгт", "Дүн", "Илүү цаг", "НДШ задаргаа", "Төлбөрийн жагсаалт", "Бусад суутгал"]
    register_values = [cell.value for row_cells in exported["Цалингийн хүснэгт"].iter_rows() for cell in row_cells]
    assert 788142 in register_values and 1788000 in register_values

    # Paid toggles both ways before closing.
    assert (await _ok(await client.post(f"/m/runs/{final['id']}/paid")))["status"] == "paid"
    assert (await _ok(await client.post(f"/m/runs/{final['id']}/unpaid")))["status"] == "approved"
    await _ok(await client.post(f"/m/runs/{final['id']}/paid"))

    # Close: the unapproved one-off advance run blocks until waived with a reason.
    preview = await _ok(await client.get(f"/m/months/{month['id']}/closing-stats"))
    assert "all_runs_must_be_approved" in preview["close_issues"] and preview["close_issues_with_waivers"] == []
    assert preview["accounting"]["account_a"]["equals_gross"] is True and preview["accounting"]["advance_reconciliation"]["matches"] is True
    blocked = await client.post(f"/m/months/{month['id']}/close", json={"waivers": {}})
    assert blocked.status_code == 409
    closed = await _ok(await client.post(f"/m/months/{month['id']}/close", json={"waivers": {str(advance21["id"]): "Төлөхгүй болсон"}}))
    assert closed["archive_version"] == 1
    assert closed["closing_stats"]["totals"]["advance_total"] == "1000000"
    assert (await client.post(f"/m/months/{month['id']}/close", json={})).status_code == 409
    # A closed month is read-only, including the waived draft advance run.
    locked = await client.put(f"/m/runs/{advance21['id']}/rows/{ids['hire']}", json={"fixed_advance": "1000", "reason": "x"})
    assert locked.status_code == 409

    archives = await _ok(await client.get(f"/m/months/{month['id']}/archives"))
    stored = await client.get(f"/m/archives/{archives[0]['id']}/runs/{final['id']}/export")
    assert stored.status_code == 200 and stored.content[:2] == b"PK"
    index = await _ok(await client.get("/m/archives"))
    assert index[0]["month"] == "2026-08" and index[0]["runs_paid"] == 1

    # Unlock keeps Approved/Paid (the waived run returns to draft); an admin reopens one run with a reason.
    await _ok(await client.post(f"/m/months/{month['id']}/unlock", json={"reason": "Засвар"}))
    month_state = await _ok(await client.get(f"/m/months/{month['id']}"))
    statuses = {run["id"]: run["status"] for run in month_state["runs"]}
    assert statuses[final["id"]] == "paid" and statuses[advance10["id"]] == "approved" and statuses[advance21["id"]] == "draft"
    reopened = await _ok(await client.post(f"/m/runs/{final['id']}/reopen", json={"reason": "Урамшуулал засна"}))
    assert reopened["status"] == "draft"
    final_run = await _ok(await client.get(f"/m/runs/{final['id']}"))
    assert all(row["status"] == "draft" for row in final_run["rows"])

    dashboard = await _ok(await client.get("/m/dashboard", params={"month": "2026-08"}))
    assert dashboard["has_final"] is True and dashboard["trend"][-1]["month"] == "2026-08"


def test_advance_projection_carries_into_final(monkeypatch):
    asyncio.run(_advance_projection_scenario(monkeypatch))


async def _advance_projection_scenario(monkeypatch):
    async with _payroll_api(monkeypatch) as (client, ids, _sessions, _organization_id):
        month = await _ok(await client.post("/m/months", json={"year": 2026, "month": 8}), 201)
        advance = await _ok(await client.post(f"/m/months/{month['id']}/runs", json={"run_type": "advance", "pay_date": "2026-08-10"}), 201)
        # Extra pay typed in the advance table: 4 h weekday overtime, leave pay, bonus.
        row = await _ok(await client.put(f"/m/runs/{advance['id']}/rows/{ids['oyun']}", json={"overtime_hours": {"weekday": "4"}, "leave_pay": "100000", "bonus": "50000", "reason": "Сарын төлөв"}))
        projection = row["result"]["projection"]
        assert row["result"]["advance"] == "600000"
        # 1,500,000 + 53,571 илүү цаг + 100,000 + 120,000 хоол унаа + 50,000 = 1,823,571.
        assert (projection["overtime_pay"], projection["gross"]) == ("53571", "1823571")
        assert (projection["employee_shi"], projection["relief"], projection["pit"]) == ("209711", "14000", "147386")
        assert (projection["total_deductions"], projection["net_pay"]) == ("957097", "866474")
        await _ok(await client.post(f"/m/runs/{advance['id']}/approve"))

        # The final run starts from what the accountant already typed.
        final = await _ok(await client.post(f"/m/months/{month['id']}/runs", json={"run_type": "final", "pay_date": "2026-08-31"}), 201)
        oyun = _row(await _ok(await client.get(f"/m/runs/{final['id']}")), ids["oyun"])
        assert (oyun["inputs"]["leave_pay"], oyun["inputs"]["bonus"]) == ("100000", "50000")
        assert oyun["inputs"]["overtime_hours"]["weekday"] == "4"
        assert oyun["inputs"]["carried_from_advance_run"] == advance["id"]
        assert oyun["result"]["overtime_pay"] == "53571" and oyun["result"]["advance"] == "600000"
