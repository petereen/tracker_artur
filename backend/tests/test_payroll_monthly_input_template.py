"""Monthly payroll Excel input template: register-shaped, pre-filled, change-only import."""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook, load_workbook

from app.payroll.monthly_input_template import KEY_ROW, TemplateError, build_input_template, parse_input_template


def _row(employee_id: int, name: str, department: str, **inputs):
    return {
        "employee_id": employee_id,
        "identity": {"name": name, "last_name": name.split()[0], "first_name": name.split()[-1], "rd": "УБ90010101", "job_title": "Нягтлан", "department": department},
        "profile": {"base_salary": "1500000", "salary_type": "PRORATION", "advance_basis": "PERCENT"},
        "inputs": {"worked_normal_hours": "168", "leave_pay": "0", "bonus": "0", "overtime_hours": {"weekday": "2"}, **inputs},
        "result": {"planned_days": 21, "planned_hours": "168", "advance_basis": "PERCENT", "advance_value": "40"},
    }


ROWS = [_row(1, "Бат Болд", "Санхүү"), _row(2, "Сараа Дорж", "Борлуулалт", other_deductions=[{"type": "Торгууль", "amount": "50000", "note": "хоцролт"}])]


def _build(run_type: str, rows=ROWS) -> bytes:
    return build_input_template(run_type=run_type, rows=rows, company="Оюунс", year=2026, month=9, pay_date_label="2026-09-25")


def _sheet(content: bytes):
    workbook = load_workbook(io.BytesIO(content))
    sheet = workbook.worksheets[0]
    keys = [cell.value for cell in sheet[KEY_ROW]]
    return workbook, sheet, keys


def _data_row(sheet, keys, employee_id):
    return next(row for row in sheet.iter_rows(min_row=KEY_ROW + 1) if row[keys.index("employee_id")].value == employee_id)


def _save(workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_final_template_follows_register_order_and_is_prefilled():
    workbook, sheet, keys = _sheet(_build("final"))
    assert keys[:7] == ["index", "employee_id", "last_name", "first_name", "rd", "job_title", "base_salary"]
    assert keys.index("planned_days") < keys.index("worked_normal_hours") < keys.index("overtime_weekday") < keys.index("leave_pay") < keys.index("bonus") < keys.index("deduction_amount") < keys.index("reason")
    assert "advance_basis" not in keys and sheet.row_dimensions[KEY_ROW].hidden
    assert sheet.cell(KEY_ROW + 1, keys.index("worked_normal_hours") + 1).value == "Ажилласан цаг"
    assert sheet.cell(KEY_ROW + 1, keys.index("deduction_type") + 1).value == "Бусад суутгал"
    # Department separators precede workers, sorted like the register.
    department_rows = [row[0].value for row in sheet.iter_rows(min_row=KEY_ROW + 3) if row[1].value is None and row[0].value]
    assert department_rows == ["Борлуулалт", "Санхүү"]
    sara = _data_row(sheet, keys, 2)
    assert sara[keys.index("worked_normal_hours")].value == 168
    assert sara[keys.index("overtime_weekday")].value == 2
    assert (sara[keys.index("deduction_type")].value, sara[keys.index("deduction_amount")].value) == ("Торгууль", 50000)
    assert workbook.sheetnames == ["Оролт", "Заавар"]


def test_untouched_template_imports_nothing():
    for run_type in ("final", "advance"):
        assert parse_input_template(_build(run_type), {row["employee_id"]: row for row in ROWS}, run_type) == []


def test_only_changed_cells_become_inputs_with_default_reason():
    workbook, sheet, keys = _sheet(_build("final"))
    bold = _data_row(sheet, keys, 1)
    bold[keys.index("bonus")].value = 100000
    bold[keys.index("overtime_rest_day")].value = 4
    sara = _data_row(sheet, keys, 2)
    sara[keys.index("deduction_amount")].value = 0
    sara[keys.index("reason")].value = "Торгууль цуцалсан"
    changes = dict(parse_input_template(_save(workbook), {row["employee_id"]: row for row in ROWS}, "final"))
    assert changes[1] == {"bonus": "100000", "overtime_hours": {"weekday": "2", "rest_day": "4", "public_holiday": "0"}, "reason": "Excel оролт"}
    assert changes[2] == {"other_deductions": [], "reason": "Торгууль цуцалсан"}


def test_advance_basis_by_label_and_value():
    workbook, sheet, keys = _sheet(_build("advance"))
    assert keys.index("advance_basis") < keys.index("advance_value") < keys.index("planned_days")
    assert "deduction_amount" not in keys
    bold = _data_row(sheet, keys, 1)
    assert (bold[keys.index("advance_basis")].value, bold[keys.index("advance_value")].value) == ("Үндсэн цалингийн хувь", 40)
    bold[keys.index("advance_basis")].value = "Тогтмол дүн"
    bold[keys.index("advance_value")].value = 500000
    _data_row(sheet, keys, 2)[keys.index("advance_value")].value = 50
    changes = dict(parse_input_template(_save(workbook), {row["employee_id"]: row for row in ROWS}, "advance"))
    assert changes[1] == {"advance_basis": "FIXED", "fixed_advance": "500000", "reason": "Excel оролт"}
    assert changes[2] == {"advance_basis": "PERCENT", "advance_percent": "50", "reason": "Excel оролт"}


def test_legacy_flat_template_still_imports():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["employee_id", "employee_name", "worked_normal_hours", "bonus", "reason"])
    sheet.append(["Ажилтны ID", "Ажилтан", "Ердийн цаг", "Урамшуулал", "Шалтгаан"])
    sheet.append([1, "Бат Болд", None, 20000, "Урамшуулал"])
    assert parse_input_template(_save(workbook), {row["employee_id"]: row for row in ROWS}, "final") == [(1, {"bonus": "20000", "reason": "Урамшуулал"})]


def test_rejects_unknown_worker_and_bad_numbers():
    workbook, sheet, keys = _sheet(_build("final"))
    _data_row(sheet, keys, 1)[keys.index("bonus")].value = "abc"
    with pytest.raises(TemplateError, match="тоо биш"):
        parse_input_template(_save(workbook), {row["employee_id"]: row for row in ROWS}, "final")
    with pytest.raises(TemplateError, match="байхгүй"):
        parse_input_template(_build("final"), {1: ROWS[0]}, "final")
    with pytest.raises(TemplateError):
        parse_input_template(b"not an xlsx", {}, "final")
