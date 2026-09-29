"""Excel layout for budget grids and analysis (d161: “Excel-ээс импортлох”).

Row 1 is a title, row 2 a hidden row of machine keys (so users may rename the
visible headers), row 3 the visible headers, data from row 4. Budget columns
are keyed by the ISO start date of each planning period.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FIXED_KEYS = ("account_code", "account_name", "project_code", "party_group_code", "note")
FIXED_HEADERS = ("Төсөвт дансны код", "Төсөвт данс", "Төслийн код", "Харилцагчийн бүлгийн код", "Тайлбар")
KEY_ROW, HEADER_ROW, FIRST_DATA_ROW = 2, 3, 4
MAX_IMPORT_ROWS = 2000
_HEADER_FILL = PatternFill("solid", fgColor="E8EEFD")
_MONEY_FORMAT = "#,##0.00;[Red]-#,##0.00"


@dataclass
class ParsedRow:
    row_number: int
    account_code: str
    project_code: str | None
    party_group_code: str | None
    note: str | None
    amounts: dict[date, Decimal] = field(default_factory=dict)


@dataclass
class ParseResult:
    rows: list[ParsedRow]
    errors: list[dict[str, Any]]


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    return text or None


def _decimal(value: Any) -> Decimal | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    try:
        return Decimal(str(value).replace(",", "").replace(" ", "").replace("₮", ""))
    except InvalidOperation:
        raise ValueError(f"“{value}” тоо биш байна") from None


def _style_header(sheet: Any, width: int) -> None:
    for column in range(1, width + 1):
        cell = sheet.cell(row=HEADER_ROW, column=column)
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[KEY_ROW].hidden = True
    sheet.freeze_panes = sheet.cell(row=FIRST_DATA_ROW, column=len(FIXED_KEYS) + 1)


def build_budget_workbook(*, title: str, columns: list[dict[str, Any]], rows: list[dict[str, Any]]) -> bytes:
    """``columns``: [{start: date, label: str}], ``rows``: fixed keys + ``amounts`` {date: Decimal}."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Төсөв"
    sheet.cell(row=1, column=1, value=title).font = Font(bold=True, size=13)
    keys = [*FIXED_KEYS, *[column["start"].isoformat() for column in columns], "total"]
    headers = [*FIXED_HEADERS, *[column["label"] for column in columns], "Нийт"]
    for index, (key, header) in enumerate(zip(keys, headers), start=1):
        sheet.cell(row=KEY_ROW, column=index, value=key)
        sheet.cell(row=HEADER_ROW, column=index, value=header)
    first_amount = len(FIXED_KEYS) + 1
    last_amount = first_amount + len(columns) - 1
    for offset, row in enumerate(rows):
        excel_row = FIRST_DATA_ROW + offset
        for index, key in enumerate(FIXED_KEYS, start=1):
            sheet.cell(row=excel_row, column=index, value=row.get(key))
        for index, column in enumerate(columns, start=first_amount):
            amount = row.get("amounts", {}).get(column["start"])
            cell = sheet.cell(row=excel_row, column=index, value=float(amount) if amount else None)
            cell.number_format = _MONEY_FORMAT
        if columns:
            total = sheet.cell(row=excel_row, column=last_amount + 1,
                               value=f"=SUM({get_column_letter(first_amount)}{excel_row}:{get_column_letter(last_amount)}{excel_row})")
            total.number_format = _MONEY_FORMAT
    _style_header(sheet, len(keys))
    for index, width in enumerate((14, 34, 14, 18, 24), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for index in range(first_amount, last_amount + 2):
        sheet.column_dimensions[get_column_letter(index)].width = 15
    notes = workbook.create_sheet("Заавар")
    for line_no, line in enumerate((
        "Орлогыг эерэг (+), ББӨ болон зардлыг сөрөг (−) утгаар оруулна. Ашиг = Орлого + Зардал.",
        "Төсөвт дансны кодыг “Төсөвт данс” тохиргооноос авна. Төсөл болон харилцагчийн бүлэг сонголттой.",
        "Нуусан 2-р мөрний түлхүүрийг бүү өөрчил — баганын гарчгийг чөлөөтэй засаж болно.",
        "Импорт хийхэд тухайн төсвийн бүх мөр файлын агуулгаар солигдоно.",
    ), start=1):
        notes.cell(row=line_no, column=1, value=line)
    notes.column_dimensions["A"].width = 110
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def parse_budget_workbook(content: bytes, column_starts: Iterable[date]) -> ParseResult:
    allowed = set(column_starts)
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # corrupt / not an xlsx
        raise ValueError("Excel файлыг уншиж чадсангүй. .xlsx форматтай эсэхийг шалгана уу.") from exc
    sheet = workbook.worksheets[0]
    grid = list(sheet.iter_rows(values_only=True))
    if len(grid) < KEY_ROW:
        raise ValueError("Файл хоосон байна")
    keys = [_text(value) for value in grid[KEY_ROW - 1]]
    if [key for key in keys[: len(FIXED_KEYS)]] != list(FIXED_KEYS):
        raise ValueError("Загвар таарахгүй байна. Төсвийн “Excel татах” загварыг ашиглана уу.")
    period_columns: dict[int, date] = {}
    errors: list[dict[str, Any]] = []
    for index, key in enumerate(keys):
        if index < len(FIXED_KEYS) or not key or key == "total":
            continue
        try:
            start = date.fromisoformat(key)
        except ValueError:
            continue
        if start not in allowed:
            errors.append({"row": KEY_ROW, "message": f"{key} хугацаа энэ төсвийн хугацаанд хамаарахгүй"})
            continue
        period_columns[index] = start
    rows: list[ParsedRow] = []
    data = grid[FIRST_DATA_ROW - 1:]
    if len(data) > MAX_IMPORT_ROWS:
        raise ValueError(f"Нэг удаад {MAX_IMPORT_ROWS}-аас ихгүй мөр импортлоно")
    for offset, values in enumerate(data):
        row_number = FIRST_DATA_ROW + offset
        values = list(values) + [None] * max(0, len(keys) - len(values))
        code = _text(values[0])
        if not code and all(value in (None, "") for value in values):
            continue
        if not code:
            errors.append({"row": row_number, "message": "Төсөвт дансны код хоосон байна"})
            continue
        parsed = ParsedRow(row_number=row_number, account_code=code, project_code=_text(values[2]), party_group_code=_text(values[3]), note=_text(values[4]))
        for index, start in period_columns.items():
            try:
                amount = _decimal(values[index])
            except ValueError as exc:
                errors.append({"row": row_number, "message": str(exc)})
                continue
            if amount:
                parsed.amounts[start] = amount
        rows.append(parsed)
    workbook.close()
    return ParseResult(rows=rows, errors=errors)


def build_analysis_workbook(*, title: str, subtitle: str, headers: list[str], rows: list[list[Any]], money_columns: set[int]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Төсөв анализ"
    sheet.cell(row=1, column=1, value=title).font = Font(bold=True, size=13)
    sheet.cell(row=2, column=1, value=subtitle)
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=HEADER_ROW, column=index, value=header)
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        sheet.column_dimensions[get_column_letter(index)].width = 30 if index == 1 else 16
    for offset, values in enumerate(rows):
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row=FIRST_DATA_ROW + offset, column=index, value=value)
            if index in money_columns:
                cell.number_format = _MONEY_FORMAT
    sheet.freeze_panes = sheet.cell(row=FIRST_DATA_ROW, column=2)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def stamp() -> str:
    return datetime.now().strftime("%Y%m%d-%H%M")
