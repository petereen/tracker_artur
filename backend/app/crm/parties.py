"""Customer master endpoints (Dayansoft d026 “Харилцагч”)."""

from __future__ import annotations

import csv
import io
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Literal

from fastapi import APIRouter, Depends, File, Query, UploadFile, status
from fastapi.responses import Response
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.core.enterprise_deps import ActorContext, get_actor
from app.crm.common import (
    ATTACHMENT_OBJECT_TYPES,
    assert_version,
    bank_account_out,
    contact_out,
    crm_error,
    get_party,
    iso,
    list_object_attachments,
    money,
    organization_timezone,
    party_context,
    party_out,
    store_object_attachment,
    validate_references,
)
from app.crm.schemas import BankAccountInput, BankAccountPatch, ContactInput, ContactPatch, PartyCreate, PartyPatch
from app.crm.service import (
    ACTIVITY_REGISTER,
    PARTY_RESOURCE,
    PARTY_TRACKED_FIELDS,
    derive_party_type,
    descendant_party_ids,
    assert_parent_is_not_descendant,
    ensure_crm_defaults,
    field_changes,
    find_duplicates,
    local_today,
    next_party_code,
    normalize_tags,
    party_references,
    utcnow,
)
from app.erp.service import require_capability, validate_custom_fields
from app.models.crm import FINISHED_STATUS_CATEGORIES, CRMActivity, ERPPartyBankAccount, ERPPartyContact, ERPPartyGroup, ERPStatus
from app.models.models import AuditLog, ERPDocument, ERPParty, UserAccount, Employee
from app.services.ebarimt_lookup import TaxpayerLookupError, lookup_taxpayer
from app.services.enterprise_events import record_change

router = APIRouter()

PARTY_REFERENCE_FIELDS = ("group_id", "responsible_employee_id", "parent_party_id", "payment_term_id", "price_list_id", "settlement_account_id")
IMPORT_COLUMNS = (
    "code", "name", "name_en", "business_name", "registry_no", "tax_id", "group", "is_customer", "is_supplier", "is_individual",
    "is_foreign", "vat_payer", "city_tax_payer", "phone", "email", "website", "legal_address", "location", "informal_address",
    "tags", "customer_since", "credit_limit", "currency",
)
EXPORT_COLUMNS = IMPORT_COLUMNS + ("responsible", "head_customer", "payment_term", "is_active")
_TRUE = {"1", "true", "yes", "y", "тийм", "т", "x", "✓"}
MAX_IMPORT_ROWS = 5000


def _snapshot(party: ERPParty) -> dict[str, Any]:
    return {field: getattr(party, field) for field in PARTY_TRACKED_FIELDS}


def _duplicate_payload(matches: list[Any]) -> list[dict[str, Any]]:
    return [{"party_id": match.party_id, "code": match.code, "name": match.name, "reason": match.reason} for match in matches]


async def _duplicate_sets(db: AsyncSession, organization_id: int) -> tuple[set[str], set[str]]:
    tins = set((await db.execute(select(ERPParty.tax_id).where(ERPParty.organization_id == organization_id, ERPParty.tax_id.is_not(None), ERPParty.tax_id != "")
                                 .group_by(ERPParty.tax_id).having(func.count(ERPParty.id) > 1))).scalars().all())
    names = set((await db.execute(select(func.lower(ERPParty.name)).where(ERPParty.organization_id == organization_id)
                                  .group_by(func.lower(ERPParty.name)).having(func.count(ERPParty.id) > 1))).scalars().all())
    return tins, names


def _with_duplicate_flags(item: dict[str, Any], tins: set[str], names: set[str]) -> dict[str, Any]:
    item["duplicate_tax_id"] = bool(item.get("tax_id") and item["tax_id"] in tins)
    item["duplicate_name"] = item["name"].lower() in names
    return item


def _list_statement(actor: ActorContext, *, search: str | None, kind: str, group_id: int | None, parent_party_id: int | None,
                    responsible_employee_id: int | None, is_active: bool | None, since_from: date | None, since_to: date | None):
    statement = select(ERPParty).where(ERPParty.organization_id == actor.organization_id)
    if search:
        pattern = f"%{search.strip()}%"
        statement = statement.where(or_(
            ERPParty.name.ilike(pattern), ERPParty.code.ilike(pattern), ERPParty.tax_id.ilike(pattern), ERPParty.registry_no.ilike(pattern),
            ERPParty.business_name.ilike(pattern), ERPParty.name_en.ilike(pattern), ERPParty.phone.ilike(pattern), ERPParty.email.ilike(pattern),
            func.array_to_string(ERPParty.tags, " ").ilike(pattern),
        ))
    if kind == "customer":
        statement = statement.where(ERPParty.is_customer.is_(True))
    elif kind == "supplier":
        statement = statement.where(ERPParty.is_supplier.is_(True))
    elif kind == "prospect":
        statement = statement.where(ERPParty.party_type == "prospect")
    if group_id is not None:
        statement = statement.where(ERPParty.group_id == group_id)
    if parent_party_id is not None:
        statement = statement.where(ERPParty.parent_party_id == parent_party_id)
    if responsible_employee_id is not None:
        statement = statement.where(ERPParty.responsible_employee_id == responsible_employee_id)
    if is_active is not None:
        statement = statement.where(ERPParty.status == "active") if is_active else statement.where(ERPParty.status != "active")
    if since_from is not None:
        statement = statement.where(ERPParty.customer_since >= since_from)
    if since_to is not None:
        statement = statement.where(ERPParty.customer_since <= since_to)
    return statement


@router.get("/parties")
async def list_parties(
    search: str | None = Query(default=None, max_length=160),
    kind: Literal["all", "customer", "supplier", "prospect"] = "all",
    group_id: int | None = None,
    parent_party_id: int | None = None,
    responsible_employee_id: int | None = None,
    is_active: bool | None = None,
    since_from: date | None = None,
    since_to: date | None = None,
    duplicates_only: bool = False,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=50, ge=1, le=500),
    db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor),
):
    await require_capability(db, actor, PARTY_RESOURCE, "view")
    statement = _list_statement(actor, search=search, kind=kind, group_id=group_id, parent_party_id=parent_party_id,
                                responsible_employee_id=responsible_employee_id, is_active=is_active, since_from=since_from, since_to=since_to)
    tins, names = await _duplicate_sets(db, actor.organization_id)
    if duplicates_only:
        statement = statement.where(or_(ERPParty.tax_id.in_(tins or {""}), func.lower(ERPParty.name).in_(names or {""})))
    total = int(await db.scalar(select(func.count()).select_from(statement.subquery())) or 0)
    rows = (await db.execute(statement.order_by(ERPParty.name, ERPParty.id).offset((page - 1) * page_size).limit(page_size))).scalars().all()
    ctx = await party_context(db, rows)
    return {"items": [_with_duplicate_flags(party_out(row, ctx), tins, names) for row in rows], "total": total, "page": page, "page_size": page_size}


@router.get("/parties/lookup-taxpayer")
async def lookup_party_taxpayer(registry_no: str | None = Query(default=None, max_length=20), tin: str | None = Query(default=None, max_length=20),
                                db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    """РД / ТТД-аар татварын системээс нэр, НӨАТ, НХАТ татах."""
    await require_capability(db, actor, PARTY_RESOURCE, "view")
    try:
        info = await lookup_taxpayer(registry_no=registry_no, tin=tin)
    except TaxpayerLookupError as exc:
        raise crm_error(404 if exc.code in {"ebarimt_not_found", "ebarimt_number_required"} else 502, exc.code, exc.message) from exc
    duplicates = await find_duplicates(db, actor.organization_id, tax_id=info.tin, registry_no=info.registry_no, name=None)
    return {**info.as_dict(), "duplicates": _duplicate_payload(duplicates)}


@router.get("/parties/export.csv")
async def export_parties(
    search: str | None = Query(default=None, max_length=160), kind: Literal["all", "customer", "supplier", "prospect"] = "all",
    group_id: int | None = None, responsible_employee_id: int | None = None, is_active: bool | None = None,
    db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor),
):
    await require_capability(db, actor, PARTY_RESOURCE, "view")
    statement = _list_statement(actor, search=search, kind=kind, group_id=group_id, parent_party_id=None,
                                responsible_employee_id=responsible_employee_id, is_active=is_active, since_from=None, since_to=None)
    rows = (await db.execute(statement.order_by(ERPParty.code))).scalars().all()
    ctx = await party_context(db, rows)
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(EXPORT_COLUMNS)
    for row in rows:
        item = party_out(row, ctx)
        writer.writerow([
            item["code"], item["name"], item["name_en"] or "", item["business_name"] or "", item["registry_no"] or "", item["tax_id"] or "",
            item["group_name"] or "", int(item["is_customer"]), int(item["is_supplier"]), int(item["is_individual"]), int(item["is_foreign"]),
            int(item["vat_payer"]), int(item["city_tax_payer"]), item["phone"] or "", item["email"] or "", item["website"] or "",
            item["legal_address"] or "", item["location"] or "", item["informal_address"] or "", ";".join(item["tags"]),
            item["customer_since"] or "", item["credit_limit"] or "", item["currency"], item["responsible_name"] or "",
            item["parent_name"] or "", item["payment_term_name"] or "", int(item["is_active"]),
        ])
    await record_change(db, actor=actor, topic="crm", aggregate_type="erp_party_export", aggregate_id=0, operation="exported", after={"rows": len(rows)})
    await db.commit()
    return Response(("﻿" + buffer.getvalue()).encode("utf-8"), media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": 'attachment; filename="customers.csv"'})


@router.get("/parties/import-template.csv")
async def party_import_template(db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, PARTY_RESOURCE, "create")
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(IMPORT_COLUMNS)
    writer.writerow(["", "Жишээ ХХК", "Example LLC", "Жишээ дэлгүүр", "1234567", "", "Харилцагчид", "1", "0", "0", "0", "1", "0",
                     "99112233", "info@example.mn", "https://example.mn", "Улаанбаатар, СБД", "Улаанбаатар", "Их дэлгүүрийн хойно",
                     "VIP;Бөөний", date.today().isoformat(), "2000000", "MNT"])
    return Response(("﻿" + buffer.getvalue()).encode("utf-8"), media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": 'attachment; filename="customer-import-template.csv"'})


def _read_import_rows(file_name: str, content: bytes) -> list[dict[str, Any]]:
    if file_name.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        workbook.close()
        if not values:
            return []
        header = [str(cell or "").strip().lower() for cell in values[0]]
        return [{header[index]: cell for index, cell in enumerate(row) if index < len(header) and header[index]} for row in values[1:]
                if any(cell not in (None, "") for cell in row)]
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    return [{(key or "").strip().lower(): value for key, value in row.items()} for row in reader if any((value or "").strip() for value in row.values() if isinstance(value, str))]


def _cell(row: dict[str, Any], key: str) -> str | None:
    value = row.get(key)
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _flag(row: dict[str, Any], key: str, default: bool = False) -> bool:
    value = _cell(row, key)
    return default if value is None else value.casefold() in _TRUE


@router.post("/parties/import")
async def import_parties(dry_run: bool = True, file: UploadFile = File(...), db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    """Excel/CSV импорт. ``dry_run`` validates and flags duplicates without saving."""
    await require_capability(db, actor, PARTY_RESOURCE, "create")
    content = await file.read(10 * 1024 * 1024 + 1)
    if len(content) > 10 * 1024 * 1024:
        raise crm_error(413, "crm_import_too_large", "Файл 10MB-аас их байна")
    try:
        raw_rows = _read_import_rows(file.filename or "import.csv", content)
    except Exception as exc:
        raise crm_error(422, "crm_import_unreadable", "Файлыг уншиж чадсангүй. CSV эсвэл XLSX загвар ашиглана уу.") from exc
    if len(raw_rows) > MAX_IMPORT_ROWS:
        raise crm_error(422, "crm_import_too_many_rows", f"Нэг удаад {MAX_IMPORT_ROWS}-аас ихгүй мөр импортлоно")
    await ensure_crm_defaults(db, actor.organization_id)
    groups = (await db.execute(select(ERPPartyGroup).where(ERPPartyGroup.organization_id == actor.organization_id))).scalars().all()
    group_lookup = {group.name.casefold(): group.id for group in groups} | {group.code.casefold(): group.id for group in groups}
    default_group = next((group.id for group in groups if group.is_default), None)
    existing_codes = set((await db.execute(select(ERPParty.code).where(ERPParty.organization_id == actor.organization_id))).scalars().all())
    existing_tins = set((await db.execute(select(ERPParty.tax_id).where(ERPParty.organization_id == actor.organization_id, ERPParty.tax_id.is_not(None)))).scalars().all())
    seen_codes: set[str] = set()
    valid: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    for index, row in enumerate(raw_rows, start=2):
        name = _cell(row, "name")
        if not name:
            errors.append({"row": index, "code": "missing_name", "message": "Нэр хоосон байна"}); continue
        code = _cell(row, "code")
        if code and (code in existing_codes or code in seen_codes):
            errors.append({"row": index, "code": "duplicate_code", "message": f"“{code}” код давхардсан байна"}); continue
        group_name = _cell(row, "group")
        group_id = group_lookup.get(group_name.casefold()) if group_name else default_group
        if group_name and group_id is None:
            errors.append({"row": index, "code": "unknown_group", "message": f"“{group_name}” бүлэг бүртгэгдээгүй байна"}); continue
        try:
            credit_limit = Decimal(_cell(row, "credit_limit")) if _cell(row, "credit_limit") else None
            customer_since = date.fromisoformat(_cell(row, "customer_since")) if _cell(row, "customer_since") else None
        except (InvalidOperation, ValueError):
            errors.append({"row": index, "code": "invalid_value", "message": "Тоо эсвэл огнооны формат буруу"}); continue
        tax_id = _cell(row, "tax_id")
        if tax_id and tax_id in existing_tins:
            warnings.append({"row": index, "code": "duplicate_tax_id", "message": f"ТТД {tax_id} өмнө бүртгэгдсэн байна"})
        if code:
            seen_codes.add(code)
        is_supplier = _flag(row, "is_supplier")
        is_customer = _flag(row, "is_customer", default=not is_supplier)
        valid.append({
            "code": code, "name": name, "name_en": _cell(row, "name_en"), "business_name": _cell(row, "business_name"),
            "registry_no": (_cell(row, "registry_no") or "").upper() or None, "tax_id": tax_id, "group_id": group_id,
            "is_customer": is_customer, "is_supplier": is_supplier, "is_individual": _flag(row, "is_individual"), "is_foreign": _flag(row, "is_foreign"),
            "vat_payer": _flag(row, "vat_payer"), "city_tax_payer": _flag(row, "city_tax_payer"), "phone": _cell(row, "phone"), "email": _cell(row, "email"),
            "website": _cell(row, "website"), "legal_address": _cell(row, "legal_address"), "location": _cell(row, "location"),
            "informal_address": _cell(row, "informal_address"), "tags": normalize_tags((_cell(row, "tags") or "").split(";")),
            "customer_since": customer_since, "credit_limit": credit_limit, "currency": (_cell(row, "currency") or "MNT").upper()[:3],
        })
    result = {"dry_run": dry_run, "total_rows": len(raw_rows), "valid_rows": len(valid), "errors": errors, "warnings": warnings, "created": 0}
    if dry_run or errors:
        return result
    for values in valid:
        values["code"] = values["code"] or await next_party_code(db, actor.organization_id)
        values["party_type"] = derive_party_type(is_customer=values["is_customer"], is_supplier=values["is_supplier"])
        db.add(ERPParty(organization_id=actor.organization_id, status="active", **values))
    await db.flush()
    await record_change(db, actor=actor, topic="crm", aggregate_type="erp_party_import", aggregate_id=0, operation="committed", after={"created": len(valid)})
    await db.commit()
    return {**result, "created": len(valid)}


@router.post("/parties", status_code=status.HTTP_201_CREATED)
async def create_party(data: PartyCreate, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, PARTY_RESOURCE, "create")
    await ensure_crm_defaults(db, actor.organization_id)
    values = data.model_dump(exclude={"confirm_duplicate_tin", "party_type", "code", "name", "is_active", "tags", "links", "custom"}, exclude_none=True)
    await validate_references(db, actor.organization_id, {field: values.get(field) for field in PARTY_REFERENCE_FIELDS})
    code = data.code or await next_party_code(db, actor.organization_id)
    if await db.scalar(select(ERPParty.id).where(ERPParty.organization_id == actor.organization_id, ERPParty.code == code)):
        raise crm_error(409, "crm_party_code_taken", f"“{code}” код өөр харилцагчид олгогдсон байна")
    duplicates = await find_duplicates(db, actor.organization_id, tax_id=data.tax_id, registry_no=data.registry_no, name=data.name)
    tin_matches = [match for match in duplicates if match.reason in {"tax_id", "registry_no"}]
    if tin_matches and not data.confirm_duplicate_tin:
        raise crm_error(409, "crm_party_duplicate_tin", "Энэ ТТД/РД-тай харилцагч бүртгэгдсэн байна. Салбар бол баталгаажуулж үргэлжлүүлнэ үү.",
                        matches=_duplicate_payload(tin_matches))
    if values.get("group_id") is None:
        values["group_id"] = await db.scalar(select(ERPPartyGroup.id).where(ERPPartyGroup.organization_id == actor.organization_id, ERPPartyGroup.is_default.is_(True), ERPPartyGroup.is_active.is_(True)).limit(1))
    is_supplier = data.is_supplier if data.is_supplier is not None else data.party_type == "supplier"
    is_customer = data.is_customer if data.is_customer is not None else (data.party_type != "supplier")
    values.update(is_customer=is_customer, is_supplier=is_supplier)
    active = data.is_active is not False
    if not active and values.get("inactive_since") is None:
        values["inactive_since"] = local_today(await organization_timezone(db, actor.organization_id))
    party = ERPParty(
        organization_id=actor.organization_id, code=code, name=data.name,
        party_type=derive_party_type(is_customer=is_customer, is_supplier=is_supplier, requested=data.party_type),
        status="active" if active else "inactive", tags=normalize_tags(data.tags),
        links=[link.model_dump() for link in data.links or []],
        custom=await validate_custom_fields(db, actor.organization_id, "party", data.custom or {}),
        **values,
    )
    db.add(party)
    await db.flush()
    await record_change(db, actor=actor, topic="crm", aggregate_type="erp_party", aggregate_id=party.id, operation="created", version=party.version,
                        after={key: value for key, value in _snapshot(party).items() if value is not None})
    await db.commit()
    await db.refresh(party)
    return {**party_out(party, await party_context(db, [party])), "duplicates": _duplicate_payload([match for match in duplicates if match.reason == "name"])}


async def _party_stats(db: AsyncSession, organization_id: int, party_ids: set[int]) -> dict[str, Any]:
    finished = select(ERPStatus.id).where(ERPStatus.organization_id == organization_id, ERPStatus.register == ACTIVITY_REGISTER, ERPStatus.category.in_(FINISHED_STATUS_CATEGORIES))
    open_clause = [CRMActivity.organization_id == organization_id, CRMActivity.party_id.in_(party_ids), CRMActivity.is_active.is_(True),
                   CRMActivity.is_closed.is_(False), CRMActivity.completed_at.is_(None),
                   or_(CRMActivity.status_id.is_(None), CRMActivity.status_id.not_in(finished))]
    now = utcnow()
    return {
        "activities_total": int(await db.scalar(select(func.count(CRMActivity.id)).where(CRMActivity.organization_id == organization_id, CRMActivity.party_id.in_(party_ids))) or 0),
        "activities_open": int(await db.scalar(select(func.count(CRMActivity.id)).where(*open_clause)) or 0),
        "activities_overdue": int(await db.scalar(select(func.count(CRMActivity.id)).where(*open_clause, CRMActivity.due_at < now)) or 0),
        "expected_revenue_open": money(await db.scalar(select(func.sum(CRMActivity.expected_revenue)).where(*open_clause))),
        "last_activity_at": iso(await db.scalar(select(func.max(CRMActivity.activity_at)).where(CRMActivity.organization_id == organization_id, CRMActivity.party_id.in_(party_ids)))),
        "documents_total": int(await db.scalar(select(func.count(ERPDocument.id)).where(ERPDocument.organization_id == organization_id, ERPDocument.party_id.in_(party_ids))) or 0),
    }


@router.get("/parties/{party_id}")
async def get_party_detail(party_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, PARTY_RESOURCE, "view")
    party = await get_party(db, actor, party_id)
    contacts = (await db.execute(select(ERPPartyContact).where(ERPPartyContact.party_id == party.id).order_by(ERPPartyContact.is_default.desc(), ERPPartyContact.name))).scalars().all()
    banks = (await db.execute(select(ERPPartyBankAccount).where(ERPPartyBankAccount.party_id == party.id).order_by(ERPPartyBankAccount.is_default.desc(), ERPPartyBankAccount.id))).scalars().all()
    children = (await db.execute(select(ERPParty).where(ERPParty.organization_id == actor.organization_id, ERPParty.parent_party_id == party.id).order_by(ERPParty.name))).scalars().all()
    tree = await descendant_party_ids(db, actor.organization_id, party.id)
    tins, names = await _duplicate_sets(db, actor.organization_id)
    return {
        **_with_duplicate_flags(party_out(party, await party_context(db, [party])), tins, names),
        "contacts": [contact_out(row) for row in contacts],
        "bank_accounts": [bank_account_out(row) for row in banks],
        "children": [{"id": row.id, "code": row.code, "name": row.name, "is_active": row.status == "active"} for row in children],
        "stats": await _party_stats(db, actor.organization_id, {party.id}),
        "group_stats": await _party_stats(db, actor.organization_id, tree) if len(tree) > 1 else None,
    }


@router.patch("/parties/{party_id}")
async def update_party(party_id: int, data: PartyPatch, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, PARTY_RESOURCE, "edit")
    party = await get_party(db, actor, party_id)
    assert_version(party, data.version)
    before = _snapshot(party)
    changes = data.model_dump(exclude_unset=True, exclude={"version", "confirm_duplicate_tin"})
    if "name" in changes and not changes["name"]:
        raise crm_error(422, "crm_party_name_required", "Нэр хоосон байж болохгүй")
    await validate_references(db, actor.organization_id, {field: changes.get(field) for field in PARTY_REFERENCE_FIELDS if field in changes})
    if "parent_party_id" in changes:
        await assert_parent_is_not_descendant(db, actor.organization_id, party.id, changes["parent_party_id"])
    if changes.get("code") and changes["code"] != party.code:
        if await db.scalar(select(ERPParty.id).where(ERPParty.organization_id == actor.organization_id, ERPParty.code == changes["code"], ERPParty.id != party.id)):
            raise crm_error(409, "crm_party_code_taken", f"“{changes['code']}” код өөр харилцагчид олгогдсон байна")
    elif "code" in changes:
        changes.pop("code")
    tin_changed = ("tax_id" in changes and changes["tax_id"] != party.tax_id) or ("registry_no" in changes and changes["registry_no"] != party.registry_no)
    if tin_changed:
        matches = await find_duplicates(db, actor.organization_id, tax_id=changes.get("tax_id", party.tax_id), registry_no=changes.get("registry_no", party.registry_no), name=None, exclude_id=party.id)
        if matches and not data.confirm_duplicate_tin:
            raise crm_error(409, "crm_party_duplicate_tin", "Энэ ТТД/РД-тай харилцагч бүртгэгдсэн байна. Салбар бол баталгаажуулж үргэлжлүүлнэ үү.", matches=_duplicate_payload(matches))
    if "tags" in changes:
        changes["tags"] = normalize_tags(changes["tags"])
    if "links" in changes:
        changes["links"] = [link for link in changes["links"] or []]
    if "custom" in changes:
        changes["custom"] = await validate_custom_fields(db, actor.organization_id, "party", changes["custom"] or {})
    if "is_active" in changes:
        active = changes.pop("is_active")
        if active is not None and active != (party.status == "active"):
            party.status = "active" if active else "inactive"
            if not active and "inactive_since" not in changes:
                changes["inactive_since"] = local_today(await organization_timezone(db, actor.organization_id))
            if active:
                changes.setdefault("inactive_since", None)
    requested_type = changes.pop("party_type", None)
    for boolean in ("is_customer", "is_supplier", "is_individual", "is_foreign", "vat_payer", "city_tax_payer", "settle_via_parent"):
        if boolean in changes and changes[boolean] is None:
            changes.pop(boolean)
    if "currency" in changes and not changes["currency"]:
        changes.pop("currency")
    for field, value in changes.items():
        setattr(party, field, value)
    if requested_type or {"is_customer", "is_supplier"} & changes.keys():
        # A prospect stays a prospect until the user changes its role flags.
        party.party_type = derive_party_type(is_customer=party.is_customer, is_supplier=party.is_supplier, requested=requested_type)
    party.version += 1
    diff = field_changes(before, _snapshot(party), PARTY_TRACKED_FIELDS)
    await record_change(db, actor=actor, topic="crm", aggregate_type="erp_party", aggregate_id=party.id, operation="updated", version=party.version,
                        before={key: value["from"] for key, value in diff.items()}, after={key: value["to"] for key, value in diff.items()})
    await db.commit()
    await db.refresh(party)
    return party_out(party, await party_context(db, [party]))


@router.post("/parties/{party_id}/refresh-tax-status")
async def refresh_party_tax_status(party_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    """Татварын мэдээллийг дахин шинэчилж шалгах (НӨАТ/НХАТ төлөв өөрчлөгдсөн эсэх)."""
    await require_capability(db, actor, PARTY_RESOURCE, "edit")
    party = await get_party(db, actor, party_id)
    try:
        info = await lookup_taxpayer(registry_no=party.registry_no, tin=party.tax_id)
    except TaxpayerLookupError as exc:
        raise crm_error(404 if exc.code in {"ebarimt_not_found", "ebarimt_number_required"} else 502, exc.code, exc.message) from exc
    before = _snapshot(party)
    party.tax_id = party.tax_id or info.tin
    party.vat_payer, party.city_tax_payer, party.tax_status_checked_at = info.vat_payer, info.city_tax_payer, utcnow()
    party.version += 1
    diff = field_changes(before, _snapshot(party), PARTY_TRACKED_FIELDS)
    await record_change(db, actor=actor, topic="crm", aggregate_type="erp_party", aggregate_id=party.id, operation="tax_status_refreshed", version=party.version,
                        before={key: value["from"] for key, value in diff.items()}, after={key: value["to"] for key, value in diff.items()})
    await db.commit()
    await db.refresh(party)
    return {**party_out(party, await party_context(db, [party])), "official_name": info.name, "changed": sorted(diff)}


@router.delete("/parties/{party_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_party(party_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, PARTY_RESOURCE, "archive")
    party = await get_party(db, actor, party_id)
    references = await party_references(db, party)
    if references:
        raise crm_error(409, "crm_party_in_use", "Харилцагч дээр гүйлгээ, харилцаа эсвэл тохиргоо үүссэн тул устгах боломжгүй. Идэвхгүй болгоно уу.", references=references)
    snapshot = {"code": party.code, "name": party.name, "tax_id": party.tax_id}
    await db.delete(party)
    await record_change(db, actor=actor, topic="crm", aggregate_type="erp_party", aggregate_id=party_id, operation="deleted", before=snapshot)
    await db.commit()


@router.get("/parties/{party_id}/history")
async def party_history(party_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    """Харилцагчийн өөрчлөлтийн түүх (Лог)."""
    await require_capability(db, actor, PARTY_RESOURCE, "view")
    await get_party(db, actor, party_id)
    rows = (await db.execute(
        select(AuditLog, Employee.name, UserAccount.email)
        .outerjoin(Employee, Employee.id == AuditLog.actor_employee_id)
        .outerjoin(UserAccount, UserAccount.id == AuditLog.actor_account_id)
        .where(AuditLog.organization_id == actor.organization_id, AuditLog.entity_type == "erp_party", AuditLog.entity_id == party_id)
        .order_by(AuditLog.created_at.desc(), AuditLog.id.desc()).limit(200)
    )).all()
    return [{"id": log.id, "action": log.action, "actor_name": name or email, "before": log.before_data or {}, "after": log.after_data or {},
             "created_at": iso(log.created_at)} for log, name, email in rows]


@router.get("/parties/{party_id}/documents")
async def party_documents(party_id: int, include_children: bool = False, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    """Харилцагчтай холбоотой нэхэмжлэх, үнийн санал, борлуулалт, худалдан авалт."""
    await require_capability(db, actor, PARTY_RESOURCE, "view")
    await get_party(db, actor, party_id)
    ids = await descendant_party_ids(db, actor.organization_id, party_id) if include_children else {party_id}
    rows = (await db.execute(select(ERPDocument).where(ERPDocument.organization_id == actor.organization_id, ERPDocument.party_id.in_(ids))
                             .order_by(ERPDocument.posting_date.desc(), ERPDocument.id.desc()).limit(300))).scalars().all()
    return [{"id": row.id, "number": row.number, "document_type": row.document_type, "status": row.status, "posting_date": iso(row.posting_date),
             "due_date": iso(row.due_date), "grand_total": money(row.grand_total), "outstanding_amount": money(row.outstanding_amount),
             "currency": row.currency, "party_id": row.party_id} for row in rows]


# ─── Contacts ─────────────────────────────────────────────────────────────────

async def _clear_default(db: AsyncSession, model: Any, party_id: int, keep_id: int | None) -> None:
    rows = (await db.execute(select(model).where(model.party_id == party_id, model.is_default.is_(True)))).scalars().all()
    for row in rows:
        if row.id != keep_id:
            row.is_default = False


@router.get("/parties/{party_id}/contacts")
async def list_contacts(party_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, PARTY_RESOURCE, "view")
    await get_party(db, actor, party_id)
    rows = (await db.execute(select(ERPPartyContact).where(ERPPartyContact.party_id == party_id).order_by(ERPPartyContact.is_default.desc(), ERPPartyContact.name))).scalars().all()
    return [contact_out(row) for row in rows]


@router.post("/parties/{party_id}/contacts", status_code=status.HTTP_201_CREATED)
async def create_contact(party_id: int, data: ContactInput, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, PARTY_RESOURCE, "edit")
    party = await get_party(db, actor, party_id)
    is_first = not await db.scalar(select(ERPPartyContact.id).where(ERPPartyContact.party_id == party.id).limit(1))
    contact = ERPPartyContact(organization_id=actor.organization_id, party_id=party.id, **{**data.model_dump(), "is_default": data.is_default or is_first})
    db.add(contact)
    await db.flush()
    if contact.is_default:
        await _clear_default(db, ERPPartyContact, party.id, contact.id)
    await record_change(db, actor=actor, topic="crm", aggregate_type="erp_party", aggregate_id=party.id, operation="contact_added", after={"contact": data.name})
    await db.commit()
    return contact_out(contact)


async def _contact(db: AsyncSession, actor: ActorContext, party_id: int, contact_id: int) -> ERPPartyContact:
    contact = await db.scalar(select(ERPPartyContact).where(ERPPartyContact.id == contact_id, ERPPartyContact.party_id == party_id, ERPPartyContact.organization_id == actor.organization_id))
    if not contact:
        raise crm_error(404, "crm_contact_not_found", "Холбоо барих хүн олдсонгүй")
    return contact


@router.patch("/parties/{party_id}/contacts/{contact_id}")
async def update_contact(party_id: int, contact_id: int, data: ContactPatch, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, PARTY_RESOURCE, "edit")
    contact = await _contact(db, actor, party_id, contact_id)
    for field, value in data.model_dump(exclude_unset=True).items():
        if field in {"name", "is_default", "is_active"} and value is None:
            continue
        setattr(contact, field, value)
    if contact.is_default:
        await _clear_default(db, ERPPartyContact, party_id, contact.id)
    await db.commit()
    return contact_out(contact)


@router.delete("/parties/{party_id}/contacts/{contact_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_contact(party_id: int, contact_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, PARTY_RESOURCE, "edit")
    contact = await _contact(db, actor, party_id, contact_id)
    if await db.scalar(select(CRMActivity.id).where(CRMActivity.contact_id == contact.id).limit(1)):
        raise crm_error(409, "crm_contact_in_use", "Харилцаа холбооны бүртгэлд ашиглагдсан тул идэвхгүй болгоно уу")
    await db.delete(contact)
    await record_change(db, actor=actor, topic="crm", aggregate_type="erp_party", aggregate_id=party_id, operation="contact_removed", after={"contact": contact.name})
    await db.commit()


# ─── Bank accounts ────────────────────────────────────────────────────────────

@router.post("/parties/{party_id}/bank-accounts", status_code=status.HTTP_201_CREATED)
async def create_bank_account(party_id: int, data: BankAccountInput, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, PARTY_RESOURCE, "edit")
    party = await get_party(db, actor, party_id)
    is_first = not await db.scalar(select(ERPPartyBankAccount.id).where(ERPPartyBankAccount.party_id == party.id).limit(1))
    account = ERPPartyBankAccount(organization_id=actor.organization_id, party_id=party.id,
                                  **{**data.model_dump(), "currency": data.currency.upper(), "is_default": data.is_default or is_first})
    db.add(account)
    await db.flush()
    if account.is_default:
        await _clear_default(db, ERPPartyBankAccount, party.id, account.id)
    await record_change(db, actor=actor, topic="crm", aggregate_type="erp_party", aggregate_id=party.id, operation="bank_account_added", after={"bank_name": data.bank_name, "account_no": data.account_no})
    await db.commit()
    return bank_account_out(account)


async def _bank_account(db: AsyncSession, actor: ActorContext, party_id: int, account_id: int) -> ERPPartyBankAccount:
    account = await db.scalar(select(ERPPartyBankAccount).where(ERPPartyBankAccount.id == account_id, ERPPartyBankAccount.party_id == party_id, ERPPartyBankAccount.organization_id == actor.organization_id))
    if not account:
        raise crm_error(404, "crm_bank_account_not_found", "Банкны данс олдсонгүй")
    return account


@router.patch("/parties/{party_id}/bank-accounts/{account_id}")
async def update_bank_account(party_id: int, account_id: int, data: BankAccountPatch, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, PARTY_RESOURCE, "edit")
    account = await _bank_account(db, actor, party_id, account_id)
    for field, value in data.model_dump(exclude_unset=True).items():
        if field in {"bank_name", "account_no", "currency", "is_default", "is_active"} and value is None:
            continue
        setattr(account, field, value.upper() if field == "currency" else value)
    if account.is_default:
        await _clear_default(db, ERPPartyBankAccount, party_id, account.id)
    before_after = {"bank_name": account.bank_name, "account_no": account.account_no}
    await record_change(db, actor=actor, topic="crm", aggregate_type="erp_party", aggregate_id=party_id, operation="bank_account_updated", after=before_after)
    await db.commit()
    return bank_account_out(account)


@router.delete("/parties/{party_id}/bank-accounts/{account_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_bank_account(party_id: int, account_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, PARTY_RESOURCE, "edit")
    account = await _bank_account(db, actor, party_id, account_id)
    await db.delete(account)
    await record_change(db, actor=actor, topic="crm", aggregate_type="erp_party", aggregate_id=party_id, operation="bank_account_removed", after={"account_no": account.account_no})
    await db.commit()


# ─── Files ────────────────────────────────────────────────────────────────────

@router.get("/parties/{party_id}/files")
async def list_party_files(party_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, PARTY_RESOURCE, "view")
    await get_party(db, actor, party_id)
    return await list_object_attachments(db, actor, ATTACHMENT_OBJECT_TYPES["party"], party_id)


@router.post("/parties/{party_id}/files", status_code=status.HTTP_201_CREATED)
async def upload_party_file(party_id: int, file: UploadFile = File(...), db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, PARTY_RESOURCE, "edit")
    await get_party(db, actor, party_id)
    return await store_object_attachment(db, actor, ATTACHMENT_OBJECT_TYPES["party"], party_id, file)
