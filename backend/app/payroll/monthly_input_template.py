"""Excel input template for a monthly payroll run.

The sheet mirrors the run register (§7.1 / §7.2 column order, department
groups, two-row grouped header) and is pre-filled with the row's current
values. Editable cells are shaded; everything else is reference only. A hidden
key row carries the machine column names, so parsing never depends on the
visible labels or their position.

Import rule: a cell that is blank or still equal to the register's value keeps
the register as is; only changed cells become row inputs.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from typing import Any, Callable, Mapping, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .monthly_exports import (
    BASIS_LABELS, BORDER, GROUP_FILL, HOURS, MONEY, SALARY_TYPE_LABELS,
    _by_department, _d, _grouped_header, _num, company_title, split_name,
)

KEY_ROW = 3
HEADER_ROW = 4
EDIT_FILL = PatternFill("solid", fgColor="FFF2CC")
READONLY_FONT = Font(color="595959")
BASIS_BY_LABEL = {label: code for code, label in BASIS_LABELS.items()}
OVERTIME_BUCKETS = ("weekday", "rest_day", "public_holiday")
DEFAULT_REASON = "Excel оролт"


class TemplateError(ValueError):
    """Unreadable or foreign workbook; message is user-facing."""


@dataclass(frozen=True)
class TemplateColumn:
    key: str
    label: str
    group: str | None = None
    editable: bool = False
    number_format: str | None = None
    width: float = 13
    value: Callable[[Mapping[str, Any]], Any] = lambda row: None


def _identity(key: str) -> Callable[[Mapping[str, Any]], Any]:
    return lambda row: row["identity"].get(key) or ""


def _inputs(key: str) -> Callable[[Mapping[str, Any]], Any]:
    return lambda row: _num(row["inputs"].get(key)) if row["inputs"].get(key) not in (None, "") else None


def _overtime(bucket: str) -> Callable[[Mapping[str, Any]], Any]:
    return lambda row: _num((row["inputs"].get("overtime_hours") or {}).get(bucket))


def _basis(row: Mapping[str, Any]) -> str | None:
    return row["inputs"].get("advance_basis") or row["result"].get("advance_basis") or row["profile"].get("advance_basis")


def _advance_value(row: Mapping[str, Any]) -> Any:
    inputs = row["inputs"]
    basis = _basis(row)
    if basis == "WORKED-TO-DATE":
        return None
    value = inputs.get("advance_percent") if basis == "PERCENT" else inputs.get("fixed_advance")
    value = value if value not in (None, "") else row["result"].get("advance_value")
    return _num(value) if value not in (None, "") else None


def _deduction(row: Mapping[str, Any], field: str) -> Any:
    lines = row["inputs"].get("other_deductions") or []
    if not lines:
        return None
    if field == "amount":
        return _num(sum((_d(line.get("amount")) for line in lines), Decimal("0")))
    if len(lines) > 1:
        return f"{len(lines)} мөр" if field == "type" else "; ".join(str(line.get("note") or "") for line in lines if line.get("note"))
    return lines[0].get(field) or ""


def template_columns(run_type: str, with_cutoff_hours: bool) -> list[TemplateColumn]:
    """Register order; `editable` marks the cells an accountant may change."""
    identity = [
        TemplateColumn("index", "№", width=5),
        TemplateColumn("employee_id", "ID", width=7, value=lambda row: row["employee_id"]),
        TemplateColumn("last_name", "Овог", width=14, value=lambda row: split_name(row["identity"])[0]),
        TemplateColumn("first_name", "Нэр", width=14, value=lambda row: split_name(row["identity"])[1]),
        TemplateColumn("rd", "РД", width=12, value=_identity("rd")),
        TemplateColumn("job_title", "Албан тушаал", width=18, value=_identity("job_title")),
        TemplateColumn("base_salary", "Үндсэн цалин", number_format=MONEY, value=lambda row: _num(row["profile"].get("base_salary"))),
    ]
    advance = [
        TemplateColumn("salary_type", "Цалингийн төрөл", value=lambda row: SALARY_TYPE_LABELS.get(row["profile"].get("salary_type"), row["profile"].get("salary_type") or "")),
        TemplateColumn("advance_basis", "Суурь", "Урьдчилгааны тооцоо", editable=True, width=18, value=lambda row: BASIS_LABELS.get(_basis(row) or "", _basis(row) or "")),
        TemplateColumn("advance_value", "Хувь / дүн", "Урьдчилгааны тооцоо", editable=True, number_format=MONEY, value=_advance_value),
        *([TemplateColumn("worked_to_date_hours", "Таслах өдөр хүртэл цаг", "Урьдчилгааны тооцоо", editable=True, number_format=HOURS, value=lambda row: _num(row["inputs"].get("worked_to_date_hours")) if _basis(row) == "WORKED-TO-DATE" else None)] if with_cutoff_hours else []),
    ] if run_type == "advance" else []
    month = [
        TemplateColumn("planned_days", "Өдөр", "Ажиллах", number_format="0", width=7, value=lambda row: _num(row["result"].get("planned_days"))),
        TemplateColumn("planned_hours", "Цаг", "Ажиллах", number_format=HOURS, width=8, value=lambda row: _num(row["result"].get("planned_hours"))),
        TemplateColumn("worked_normal_hours", "Ажилласан цаг", editable=True, number_format=HOURS, value=_inputs("worked_normal_hours")),
        TemplateColumn("worked_days", "Ажилласан өдөр", editable=True, number_format=HOURS, value=_inputs("worked_days")),
        TemplateColumn("overtime_weekday", "Ажлын өдөр (И)", "Илүү цаг", editable=True, number_format=HOURS, value=_overtime("weekday")),
        TemplateColumn("overtime_rest_day", "Амралтын өдөр (А)", "Илүү цаг", editable=True, number_format=HOURS, value=_overtime("rest_day")),
        TemplateColumn("overtime_public_holiday", "Баярын өдөр (Б)", "Илүү цаг", editable=True, number_format=HOURS, value=_overtime("public_holiday")),
        TemplateColumn("leave_pay", "Ээлжийн амралтын мөнгө", editable=True, number_format=MONEY, value=_inputs("leave_pay")),
        TemplateColumn("bonus", "Урамшуулал", editable=True, number_format=MONEY, value=_inputs("bonus")),
    ]
    deductions = [
        TemplateColumn("deduction_type", "Төрөл", "Бусад суутгал", editable=True, width=14, value=lambda row: _deduction(row, "type")),
        TemplateColumn("deduction_amount", "Дүн", "Бусад суутгал", editable=True, number_format=MONEY, value=lambda row: _deduction(row, "amount")),
        TemplateColumn("deduction_note", "Тайлбар", "Бусад суутгал", editable=True, width=18, value=lambda row: _deduction(row, "note")),
    ] if run_type == "final" else []
    reason = [TemplateColumn("reason", "Засварын шалтгаан", editable=True, width=24)]
    return [*identity, *advance, *month, *deductions, *reason]


def _group_spans(columns: Sequence[TemplateColumn]) -> dict[tuple[int, int], str]:
    spans: dict[tuple[int, int], str] = {}
    for index, column in enumerate(columns, start=1):
        if not column.group:
            continue
        start = next((span for span, label in spans.items() if label == column.group and span[1] == index - 1), None)
        if start:
            spans.pop(start)
            spans[(start[0], index)] = column.group
        else:
            spans[(index, index)] = column.group
    return spans


def build_input_template(
    *, run_type: str, rows: Sequence[Mapping[str, Any]], company: str | None, year: int, month: int, pay_date_label: str,
) -> bytes:
    """Rows are ``{"employee_id", "identity", "profile", "inputs", "result"}``."""
    with_cutoff_hours = run_type == "advance" and any(_basis(row) == "WORKED-TO-DATE" for row in rows)
    columns = template_columns(run_type, with_cutoff_hours)
    width = len(columns)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Оролт"
    kind = "урьдчилгаа цалингийн" if run_type == "advance" else "цалингийн"
    sheet.cell(1, 1, company_title(company, f"{year} оны {month:02d} сарын {kind} оролтын загвар ({pay_date_label})")).font = Font(bold=True, size=14)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    sheet.cell(2, 1, "Зөвхөн шар нүдийг засна. Хоосон эсвэл өөрчлөгдөөгүй нүд системийн утгаа хадгална; тэглэх бол 0 бичнэ. Мөр, багана нэмэх, устгахгүй.").font = Font(italic=True, color="595959")
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    for index, column in enumerate(columns, start=1):
        sheet.cell(KEY_ROW, index, column.key)
    sheet.row_dimensions[KEY_ROW].hidden = True
    _grouped_header(sheet, HEADER_ROW, [column.label for column in columns], _group_spans(columns))
    for index, column in enumerate(columns, start=1):
        if column.editable:
            for header_row in (HEADER_ROW, HEADER_ROW + 1):
                sheet.cell(header_row, index).fill = EDIT_FILL
    current = HEADER_ROW + 2
    counter = 0
    first_data_row = current
    for department, department_rows in _by_department(rows).items():
        sheet.cell(current, 1, department).font = Font(bold=True)
        sheet.merge_cells(start_row=current, start_column=1, end_row=current, end_column=width)
        sheet.cell(current, 1).fill = GROUP_FILL
        current += 1
        for row in department_rows:
            counter += 1
            for index, column in enumerate(columns, start=1):
                value = counter if column.key == "index" else column.value(row)
                cell = sheet.cell(current, index, value if value != "" else None)
                cell.border = BORDER
                if column.number_format:
                    cell.number_format = column.number_format
                if column.editable:
                    cell.fill = EDIT_FILL
                else:
                    cell.font = READONLY_FONT
            current += 1
    last_data_row = max(first_data_row, current - 1)
    if run_type == "advance":
        basis_index = next(index for index, column in enumerate(columns, start=1) if column.key == "advance_basis")
        validation = DataValidation(type="list", formula1='"' + ",".join(BASIS_LABELS.values()) + '"', allow_blank=True)
        validation.add(f"{get_column_letter(basis_index)}{first_data_row}:{get_column_letter(basis_index)}{last_data_row}")
        sheet.add_data_validation(validation)
    for index, column in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = column.width
    sheet.freeze_panes = f"E{HEADER_ROW + 2}"
    notes = workbook.create_sheet("Заавар")
    for line_index, line in enumerate((
        "Оролтын загвар — хэрхэн ашиглах вэ",
        "1. Энэ загвар системийн бодолтын хүснэгттэй ижил дараалалтай, одоогийн утгуудаар бөглөгдсөн.",
        "2. Шар өнгөтэй багануудыг л засна: ажилласан цаг/өдөр, илүү цаг (И/А/Б), ээлжийн амралтын мөнгө, урамшуулал"
        + (", бусад суутгал." if run_type == "final" else ", урьдчилгааны суурь, хувь/дүн."),
        "3. Хоосон эсвэл өөрчлөгдөөгүй нүд системийн утгаа хадгална. Утгыг тэглэх бол 0 бичнэ.",
        "4. Засварын шалтгаан заавал биш; хоосон бол «Excel оролт» гэж бүртгэгдэнэ.",
        "5. ID багана, 3-р (нуусан) мөрийг өөрчлөхгүй. Хэлтсийн мөрүүдийг устгах шаардлагагүй.",
        "6. Хадгалаад «Import (Excel)» товчоор оруулна. Засагдсан мөрүүдийн батлалт цуцлагдаж дахин бодогдоно.",
    ), start=1):
        notes.cell(line_index, 1, line).font = Font(bold=line_index == 1, size=13 if line_index == 1 else 11)
    notes.column_dimensions["A"].width = 120
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _decimal(value: Any, *, label: str, where: str) -> Decimal | None:
    text = _text(value)
    if text is None:
        return None
    try:
        return Decimal(text.replace(",", "").replace("%", "").replace("₮", "").strip())
    except InvalidOperation as exc:
        raise TemplateError(f"{where}: «{label}» баганын «{text}» утга тоо биш байна.") from exc


def _same(new: Decimal | None, old: Any) -> bool:
    return new is None or (old not in (None, "") and new == _d(old))


def parse_input_template(content: bytes, rows: Mapping[int, Mapping[str, Any]], run_type: str) -> list[tuple[int, dict[str, Any]]]:
    """Return ``(employee_id, RowInput payload)`` for rows whose cells changed.

    ``rows`` maps employee id → the same row mapping the template was built
    from. Also reads the earlier flat template (keys in row 1, labels in row 2).
    """
    try:
        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
        values = list(sheet.iter_rows(values_only=True))
    except Exception as exc:  # noqa: BLE001 - any parser failure is a bad upload
        raise TemplateError("Excel загварыг уншиж чадсангүй.") from exc
    key_index = next((index for index, row in enumerate(values[:10]) if "employee_id" in [_text(cell) for cell in row]), None)
    if key_index is None:
        raise TemplateError("Системээс татсан оролтын загвар ашиглана уу (ID багана олдсонгүй).")
    keys = [_text(cell) or "" for cell in values[key_index]]
    named = [key for key in keys if key]
    if len(named) != len(set(named)):
        raise TemplateError("Загварын багана давхардсан байна. Шинэ загвар татаж ашиглана уу.")
    labels = {column.key: column.label for column in template_columns(run_type, True)}
    changes: list[tuple[int, dict[str, Any]]] = []
    seen: set[int] = set()
    started = False
    for offset, raw in enumerate(values[key_index + 1:], start=key_index + 2):
        record = {key: raw[index] if index < len(raw) else None for index, key in enumerate(keys) if key}
        employee_text = _text(record.get("employee_id"))
        if employee_text is None:
            continue  # department separator or blank line
        try:
            employee_id = int(Decimal(employee_text))
        except InvalidOperation as exc:
            if not started:
                continue  # header label rows
            raise TemplateError(f"{offset}-р мөрийн ажилтны ID буруу байна.") from exc
        started = True
        if employee_id in seen:
            raise TemplateError("Нэг ажилтныг файлд давхар оруулсан байна.")
        seen.add(employee_id)
        row = rows.get(employee_id)
        if row is None:
            raise TemplateError(f"{offset}-р мөр: ID {employee_id} ажилтан энэ бодолтод байхгүй.")
        where = f"{offset}-р мөр ({row['identity'].get('name') or employee_id})"
        inputs = row["inputs"]

        def number(key: str) -> Decimal | None:
            return _decimal(record.get(key), label=labels.get(key, key), where=where) if key in record else None

        provided: dict[str, Any] = {}
        hour_keys = ("worked_normal_hours", "worked_days", "leave_pay", "bonus") + (("worked_to_date_hours",) if run_type == "advance" else ())
        for key in hour_keys:
            value = number(key)
            if not _same(value, inputs.get(key)):
                provided[key] = str(value)
        overtime_now = inputs.get("overtime_hours") or {}
        overtime = {bucket: number(f"overtime_{bucket}") for bucket in OVERTIME_BUCKETS}
        if any(not _same(value, overtime_now.get(bucket, "0")) for bucket, value in overtime.items()):
            provided["overtime_hours"] = {**overtime_now, **{bucket: str(value) for bucket, value in overtime.items() if value is not None}}
        if run_type == "advance" and ("advance_basis" in record or "advance_value" in record):
            basis_text = _text(record.get("advance_basis"))
            basis = BASIS_BY_LABEL.get(basis_text or "", basis_text) if basis_text else _basis(row)
            if basis not in BASIS_LABELS:
                raise TemplateError(f"{where}: урьдчилгааны суурь «{basis_text}» буруу. {', '.join(BASIS_LABELS.values())} -аас сонгоно уу.")
            value = number("advance_value")
            current_value = _advance_value(row)
            if basis != _basis(row) or not _same(value, current_value):
                provided["advance_basis"] = basis
                effective = value if value is not None else (_d(current_value) if current_value is not None else None)
                if basis == "PERCENT" and effective:
                    provided["advance_percent"] = str(effective)
                elif basis == "FIXED" and effective:
                    provided["fixed_advance"] = str(effective)
        if run_type == "final" and any(key in record for key in ("deduction_type", "deduction_amount", "deduction_note")):
            cells = (_text(record.get("deduction_type")), number("deduction_amount"), _text(record.get("deduction_note")))
            current = (_text(_deduction(row, "type")), _d(_deduction(row, "amount")) if _deduction(row, "amount") is not None else None, _text(_deduction(row, "note")))
            if cells[1] is not None and cells != current:
                provided["other_deductions"] = [{"type": cells[0] or "Бусад", "amount": str(cells[1]), "note": cells[2] or ""}] if cells[1] > 0 else []
        if provided:
            provided["reason"] = (_text(record.get("reason")) or DEFAULT_REASON)[:1000]
            changes.append((employee_id, provided))
    return changes
