"""Organization-scoped API for the run-based monthly payroll design."""

from __future__ import annotations

import calendar
import base64
import io
import json
import statistics
from dataclasses import asdict
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, ROUND_HALF_UP
from typing import Any, Literal

from fastapi import APIRouter, Depends, File, HTTPException, Response, UploadFile, status
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from pydantic import BaseModel, Field
from sqlalchemy import or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.core.enterprise_deps import ActorContext, get_actor
from app.erp.service import require_capability
from app.models.models import (
    AttendanceLog, Department, Employee, EmployeeBankAccount, EmployeeDetails, HolidayRecord, MonthlyPayrollArchive, MonthlyPayrollCalendarDay, MonthlyPayrollCompanySettings,
    MonthlyPayrollMonth, MonthlyPayrollProfile, MonthlyPayrollRowAudit,
    MonthlyPayrollRuleSet, MonthlyPayrollRun, MonthlyPayrollRunRow, ERPAccount,
    MonthlyPayrollSalaryHistory, Organization, Schedule, TimeOff, WorkTimeEntry,
)
from app.services.enterprise_events import record_change
from app.services.secret_box import decrypt_secret, encrypt_secret
from .inputs import payment_days
from .monthly_exports import build_run_workbook
from .monthly_engine import (
    AdvanceBasis, AllowanceBasis, CalendarDayType, PayrollProfile, PayrollRunType, PayrollRules, apply_computed_overrides,
    SalarySegment, amount, calculate_monthly_run, classify_work_hours, default_2026_rules, month_calendar,
    overtime_day_lines, whole_tugrik,
)


router = APIRouter()
# «Урьдчилгаа бодоогүй» (advance_not_calculated) is a warning by design: the
# accountant may knowingly settle a worker whose advance run was never made.
BLOCKING_ROW_WARNINGS = {"negative_final_pay", "profile_missing", "salary_history_missing_or_incomplete", "row_flagged", "advance_changed", "advance_not_due", "advance_not_positive"}
TIME_INPUT_KEYS = ("worked_normal_hours", "worked_days", "worked_to_date_hours", "worked_days_to_date", "projected_remaining_hours", "elapsed_planned_days", "overtime_hours", "day_lines", "missing_dates", "approved_leave_days", "time_source")


class MonthInput(BaseModel):
    year: int = Field(ge=2000, le=2100)
    month: int = Field(ge=1, le=12)


class RunInput(BaseModel):
    run_type: Literal["advance", "final"]
    pay_date: date
    cutoff_date: date | None = None
    department_id: int | None = None
    employee_ids: list[int] | None = Field(default=None, min_length=1, max_length=500)
    note: str | None = Field(default=None, max_length=2000)


class RowInput(BaseModel):
    worked_normal_hours: Decimal | None = Field(default=None, ge=0)
    worked_days: Decimal | None = Field(default=None, ge=0, le=31)
    worked_to_date_hours: Decimal = Field(default=Decimal("0"), ge=0)
    elapsed_planned_days: int = Field(default=0, ge=0)
    overtime_hours: dict[str, Decimal] = Field(default_factory=dict)
    leave_pay: Decimal = Field(default=Decimal("0"), ge=0)
    bonus: Decimal = Field(default=Decimal("0"), ge=0)
    other_deductions: list[dict[str, Any]] = Field(default_factory=list)
    fixed_advance: Decimal | None = Field(default=None, gt=0)
    advance_percent: Decimal | None = Field(default=None, gt=0, le=100)
    advance_basis: Literal["FIXED", "PERCENT", "WORKED-TO-DATE"] | None = None
    reason: str | None = Field(default=None, max_length=1000)


class UnlockInput(BaseModel):
    reason: str = Field(min_length=1, max_length=1000)


class FlagInput(BaseModel):
    reason: str = Field(min_length=1, max_length=1000)


class ComputedOverrideInput(BaseModel):
    field: Literal["gross", "employee_shi", "employer_shi", "pit", "advance", "other_deductions"]
    value: Decimal = Field(ge=0, le=Decimal("999999999999999"))
    reason: str = Field(min_length=1, max_length=1000)


class ComputedOverrideRevertInput(BaseModel):
    field: Literal["gross", "employee_shi", "employer_shi", "pit", "advance", "other_deductions"]
    reason: str = Field(min_length=1, max_length=1000)


class CalendarDayInput(BaseModel):
    day_type: Literal["working", "weekly_rest", "public_holiday"]
    holiday_name: str | None = Field(default=None, max_length=240)


class CompanyMonthlySettingsInput(BaseModel):
    legal_company_name: str | None = Field(default=None, max_length=240)
    daily_norm_hours: Decimal = Field(gt=0, le=24)
    employer_injury_rate: Decimal = Field(ge=Decimal("0.005"), le=Decimal("0.025"))
    weekday_overtime_multiplier: Decimal = Field(ge=Decimal("1.5"), le=Decimal("10"))
    rest_day_overtime_multiplier: Decimal = Field(ge=Decimal("1.5"), le=Decimal("10"))
    public_holiday_overtime_multiplier: Decimal = Field(ge=Decimal("2"), le=Decimal("10"))
    default_advance_basis: Literal["FIXED", "PERCENT", "WORKED-TO-DATE"]
    default_advance_percent: Decimal = Field(gt=0, le=100)
    deduction_types: list[str] = Field(min_length=1, max_length=32)
    salary_expense_account_id: int | None = Field(default=None, gt=0)
    employer_shi_account_id: int | None = Field(default=None, gt=0)
    advance_clearing_account_id: int | None = Field(default=None, gt=0)


class MonthlyRuleSetInput(BaseModel):
    valid_from: date
    valid_to: date | None = None
    minimum_wage: Decimal = Field(gt=0, le=Decimal("999999999"))
    shi_cap_multiplier: Decimal = Field(gt=0, le=Decimal("100"))
    employee_rates: dict[str, Decimal]
    employer_rates: dict[str, Decimal]
    pit_brackets: list[dict[str, Any]] = Field(min_length=1, max_length=20)
    relief_tiers: list[dict[str, Any]] = Field(default_factory=list, max_length=20)
    overtime_multipliers: dict[str, Decimal]
    source_references: list[str] = Field(min_length=1, max_length=20)


APPROVED_RUN_STATUSES = ("approved", "paid", "closed")


def _money(value: Any) -> Decimal:
    return Decimal(str(value or 0))


async def _account_label(db: AsyncSession, organization_id: int, account_id: int | None) -> dict[str, Any] | None:
    if not account_id:
        return None
    account = await db.scalar(select(ERPAccount).where(ERPAccount.id == account_id, ERPAccount.organization_id == organization_id))
    return {"id": account_id, "code": account.code, "name": account.name} if account else {"id": account_id, "code": None, "name": None}


async def _closing_stats(db: AsyncSession, month: MonthlyPayrollMonth, organization_id: int, *, waived_run_ids: set[int] | None = None) -> dict[str, Any]:
    """Month statistics for the dashboard, closing review, «Дүн» sheet and archive."""
    waived_run_ids = waived_run_ids or set()
    runs = (await db.execute(select(MonthlyPayrollRun).where(
        MonthlyPayrollRun.month_id == month.id, MonthlyPayrollRun.organization_id == organization_id,
    ).order_by(MonthlyPayrollRun.pay_date, MonthlyPayrollRun.run_type))).scalars().all()
    run_rows: dict[int, list[MonthlyPayrollRunRow]] = {}
    for run in runs:
        run_rows[run.id] = (await db.execute(select(MonthlyPayrollRunRow).where(MonthlyPayrollRunRow.run_id == run.id))).scalars().all()
    first = date(month.year, month.month, 1)
    last = date(month.year, month.month, calendar.monthrange(month.year, month.month)[1])
    final_run = next((run for run in runs if run.run_type == "final"), None)
    final_rows = run_rows.get(final_run.id, []) if final_run else []
    fields = ("gross", "base_pay", "overtime_pay", "meal_commute", "employee_shi", "employer_shi", "taxable_income", "pit_before_relief", "relief", "pit", "advance", "other_deductions", "total_deductions", "net_pay")
    totals = {key: sum((_money((row.result or {}).get(key)) for row in final_rows), Decimal("0")) for key in fields}
    totals["leave_pay"] = sum((_money((row.inputs or {}).get("leave_pay")) for row in final_rows), Decimal("0"))
    totals["bonus"] = sum((_money((row.inputs or {}).get("bonus")) for row in final_rows), Decimal("0"))
    gross_values = sorted(_money((row.result or {}).get("gross")) for row in final_rows)
    net_values = [_money((row.result or {}).get("net_pay")) for row in final_rows]
    advance_runs = [run for run in runs if run.run_type == "advance"]
    counted_advance_runs = [run for run in advance_runs if run.status in APPROVED_RUN_STATUSES and run.id not in waived_run_ids]
    advance_by_run = [{
        "run_id": run.id, "pay_date": run.pay_date, "status": run.status, "waived": run.id in waived_run_ids,
        "workers": len(run_rows.get(run.id, [])),
        "total": sum((_money((row.result or {}).get("advance")) for row in run_rows.get(run.id, [])), Decimal("0")),
    } for run in advance_runs]
    advance_total = sum((item["total"] for item in advance_by_run if item["run_id"] in {run.id for run in counted_advance_runs}), Decimal("0"))
    overtime_hours: dict[str, Decimal] = {}
    overtime_amounts: dict[str, Decimal] = {}
    overtime_workers: list[dict[str, Any]] = []
    deduction_totals: dict[str, Decimal] = {}
    departments: dict[str, dict[str, Any]] = {}
    new_workers = left_workers = cap_hits = warned_rows = overridden_rows = 0
    cap = _money(month.rule_snapshot.get("minimum_wage")) * _money(month.rule_snapshot.get("shi_cap_multiplier", 10))
    for row in final_rows:
        identity, inputs, result = row.identity_snapshot or {}, row.inputs or {}, row.result or {}
        department = identity.get("department") or "Бусад"
        group = departments.setdefault(department, {key: Decimal("0") for key in ("headcount", "gross", "employee_shi", "pit", "employer_shi", "other_deductions", "advance", "net_pay", "company_cost")})
        group["headcount"] += 1
        for key in ("gross", "employee_shi", "pit", "employer_shi", "other_deductions", "advance", "net_pay"):
            group[key] += _money(result.get(key))
        group["company_cost"] += _money(result.get("gross")) + _money(result.get("employer_shi"))
        hours = {kind: _money(value) for kind, value in (inputs.get("overtime_hours") or {}).items()}
        if sum(hours.values(), Decimal("0")) > 0:
            overtime_workers.append({"employee_id": row.employee_id, "name": identity.get("name"), "department": department, "hours": sum(hours.values(), Decimal("0")), "amount": _money(result.get("overtime_pay"))})
        for kind, value in hours.items():
            overtime_hours[kind] = overtime_hours.get(kind, Decimal("0")) + value
        for kind, value in (result.get("overtime_by_bucket") or {}).items():
            overtime_amounts[kind] = overtime_amounts.get(kind, Decimal("0")) + _money(value)
        for line in inputs.get("other_deductions") or []:
            label = line.get("type") or "Төрөлгүй"
            deduction_totals[label] = deduction_totals.get(label, Decimal("0")) + _money(line.get("amount"))
        if identity.get("start_date") and first <= date.fromisoformat(identity["start_date"]) <= last:
            new_workers += 1
        if identity.get("end_date") and first <= date.fromisoformat(identity["end_date"]) <= last:
            left_workers += 1
        if cap > 0 and _money(result.get("shi_base")) >= cap and _money(result.get("gross")) > cap:
            cap_hits += 1
        if row.warnings:
            warned_rows += 1
        if result.get("computed_overrides"):
            overridden_rows += 1
    company_cost = totals["gross"] + totals["employer_shi"]
    for group in departments.values():
        group["share"] = (group["company_cost"] * 100 / company_cost).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP) if company_cost else Decimal("0")
    run_ids = [run.id for run in runs] or [-1]
    audits = (await db.execute(select(MonthlyPayrollRowAudit).where(
        MonthlyPayrollRowAudit.organization_id == organization_id, MonthlyPayrollRowAudit.run_id.in_(run_ids),
    ))).scalars().all()
    edited_rows = {audit.row_id for audit in audits if audit.field_name in {"inputs", "excel_import", "overrides_reverted"} or audit.field_name.startswith("computed:")}
    flagged_rows = {audit.row_id for audit in audits if audit.field_name == "status" and isinstance(audit.new_value, dict) and audit.new_value.get("status") == "flagged"}
    currently_flagged = {row.id for rows in run_rows.values() for row in rows if row.status == "flagged"}
    settings = await db.scalar(select(MonthlyPayrollCompanySettings).where(MonthlyPayrollCompanySettings.organization_id == organization_id))
    account_a = totals["total_deductions"] + totals["net_pay"]
    accounting = {
        "account_a": {
            "account": await _account_label(db, organization_id, settings.salary_expense_account_id if settings else None),
            "total": account_a, "equals_gross": account_a == totals["gross"],
            "lines": {"advance": totals["advance"], "pit": totals["pit"], "employee_shi": totals["employee_shi"], "other_deductions": deduction_totals, "net_pay": totals["net_pay"]},
        },
        "account_b": {"account": await _account_label(db, organization_id, settings.employer_shi_account_id if settings else None), "total": totals["employer_shi"]},
        "account_c": {"account": await _account_label(db, organization_id, settings.advance_clearing_account_id if settings else None), "total": advance_total} if settings and settings.advance_clearing_account_id else None,
        "advance_reconciliation": {"pulled_into_final": totals["advance"], "approved_advance_runs": advance_total, "matches": totals["advance"] == advance_total},
        "by_department": {name: {"account_a": group["gross"], "account_b": group["employer_shi"]} for name, group in departments.items()},
    }
    previous_month = await db.scalar(select(MonthlyPayrollMonth).where(
        MonthlyPayrollMonth.organization_id == organization_id, MonthlyPayrollMonth.status == "closed",
        (MonthlyPayrollMonth.year * 100 + MonthlyPayrollMonth.month) < month.year * 100 + month.month,
    ).order_by(MonthlyPayrollMonth.year.desc(), MonthlyPayrollMonth.month.desc()).limit(1))
    comparison = None
    if previous_month:
        archive = await db.scalar(select(MonthlyPayrollArchive).where(MonthlyPayrollArchive.month_id == previous_month.id).order_by(MonthlyPayrollArchive.version.desc()).limit(1))
        previous = (archive.snapshot or {}).get("closing_stats", {}) if archive else {}
        current_values = {"gross": totals["gross"], "company_cost": company_cost, "headcount": Decimal(len(final_rows)), "overtime_amount": sum(overtime_amounts.values(), Decimal("0"))}
        previous_values = {
            "gross": _money((previous.get("totals") or {}).get("gross")), "company_cost": _money((previous.get("totals") or {}).get("company_cost")),
            "headcount": _money((previous.get("headcount") or {}).get("on_register")),
            "overtime_amount": sum((_money(value) for value in ((previous.get("overtime") or {}).get("amounts") or {}).values()), Decimal("0")),
        }
        comparison = {"month": f"{previous_month.year:04d}-{previous_month.month:02d}", "metrics": {
            key: {"current": current_values[key], "previous": previous_values[key], "change": current_values[key] - previous_values[key],
                  "change_pct": ((current_values[key] - previous_values[key]) * 100 / previous_values[key]).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP) if previous_values[key] else None}
            for key in current_values
        }}
    return _json_value({
        "headcount": {
            "on_register": len(final_rows), "paid": len(final_rows) if final_run and final_run.paid_at else 0,
            "new": new_workers, "left": left_workers, "with_extra_work": len(overtime_workers), "manually_edited": len(edited_rows),
        },
        "totals": {**totals, "advance_total": advance_total, "company_cost": company_cost, "account_a": account_a, "account_b": totals["employer_shi"]},
        "averages": {
            "average_gross": (sum(gross_values, Decimal("0")) / len(gross_values)).quantize(Decimal("1"), rounding=ROUND_HALF_UP) if gross_values else 0,
            "median_gross": Decimal(str(statistics.median(gross_values))).quantize(Decimal("1"), rounding=ROUND_HALF_UP) if gross_values else 0,
            "average_take_home": (sum(net_values, Decimal("0")) / len(net_values)).quantize(Decimal("1"), rounding=ROUND_HALF_UP) if net_values else 0,
            "highest_gross": gross_values[-1] if gross_values else 0, "lowest_gross": gross_values[0] if gross_values else 0,
        },
        "overtime": {"hours": overtime_hours, "amounts": overtime_amounts, "workers": len(overtime_workers), "top": sorted(overtime_workers, key=lambda item: item["amount"], reverse=True)[:5]},
        "advance_by_run": advance_by_run,
        "other_deductions_by_type": deduction_totals,
        "by_department": departments,
        "accounting": accounting,
        "comparison": comparison,
        "quality": {
            "manual_rows": len(edited_rows), "computed_overrides": overridden_rows, "shi_cap_hits": cap_hits, "rows_with_warnings": warned_rows,
            "workers_without_advance": sum("advance_not_calculated" in (row.warnings or []) for row in final_rows),
            "flagged_then_resolved": len(flagged_rows - currently_flagged),
        },
        "meta": {
            "rule_set_id": month.rule_set_id, "rule_version": (month.rule_snapshot or {}).get("version"),
            "approvers": sorted({run.approved_by_account_id for run in runs if run.approved_by_account_id}),
            "calendar_working_days": sum(value == CalendarDayType.WORKING.value for value in (month.calendar_snapshot or {}).values()),
        },
        "run_statuses": {str(run.id): run.status for run in runs},
    })


async def _month_close_issues(db: AsyncSession, actor: ActorContext, month: MonthlyPayrollMonth, runs: list[MonthlyPayrollRun], *, waived_run_ids: set[int] | None = None) -> list[str]:
    waived_run_ids = waived_run_ids or set()
    issues: list[str] = []
    counted = [run for run in runs if run.id not in waived_run_ids]
    finals = [run for run in counted if run.run_type == "final"]
    if len(finals) != 1:
        issues.append("final_run_required")
    if any(run.status not in APPROVED_RUN_STATUSES for run in counted):
        issues.append("all_runs_must_be_approved")
    all_rows: dict[int, list[MonthlyPayrollRunRow]] = {}
    for run in counted:
        rows = list((await db.execute(select(MonthlyPayrollRunRow).where(
            MonthlyPayrollRunRow.run_id == run.id,
            MonthlyPayrollRunRow.organization_id == actor.organization_id,
        ))).scalars().all())
        all_rows[run.id] = rows
        if any(row.status != "approved" for row in rows):
            issues.append("all_rows_must_be_approved")
    if not finals:
        return sorted(set(issues))
    final_rows = all_rows.get(finals[0].id, [])
    for row in final_rows:
        result = row.result or {}
        if not row.profile_snapshot.get("complete", True):
            issues.append("profile_incomplete")
        if _money(result.get("net_pay")) < 0:
            issues.append("negative_final_pay")
        if _money(result.get("total_deductions")) + _money(result.get("net_pay")) != _money(result.get("gross")):
            issues.append("row_payroll_equation_mismatch")
        if "advance_changed" in (row.warnings or []):
            issues.append("advance_changed")
    final_advance = sum((_money((row.result or {}).get("advance")) for row in final_rows), Decimal("0"))
    approved_advance = sum((_money((row.result or {}).get("advance")) for run in counted if run.run_type == "advance" and run.status in APPROVED_RUN_STATUSES for row in all_rows.get(run.id, [])), Decimal("0"))
    if final_advance != approved_advance:
        issues.append("advance_reconciliation_mismatch")
    return sorted(set(issues))


def _json_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, dict):
        return {str(key): _json_value(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_value(item) for item in value]
    return value


def _plain_value(value: Any) -> str:
    """Form value without Numeric padding: 8.0000 -> "8", 0.0050 -> "0.005"."""
    return format(value.normalize(), "f") if isinstance(value, Decimal) else str(value)


def _rules(snapshot: dict[str, Any]) -> PayrollRules:
    return PayrollRules(
        minimum_wage=amount(snapshot["minimum_wage"]),
        shi_cap_multiplier=amount(snapshot.get("shi_cap_multiplier", 10)),
        employee_shi_rates={key: amount(value) for key, value in snapshot["employee_rates"].items()},
        employer_shi_rates={key: amount(value) for key, value in snapshot["employer_rates"].items()},
        pit_brackets=tuple((amount(row["lower"]), amount(row["upper"]) if row.get("upper") is not None else None, amount(row["rate"]), amount(row.get("base_tax", 0))) for row in snapshot["pit_brackets"]),
        relief_tiers=tuple((amount(row["lower"]), amount(row["upper"]) if row.get("upper") is not None else None, amount(row["amount"])) for row in snapshot["relief_tiers"]),
        overtime_multipliers={key: amount(value) for key, value in snapshot["overtime_multipliers"].items()},
    )


def _rule_snapshot(row: MonthlyPayrollRuleSet) -> dict[str, Any]:
    return _json_value({
        "id": row.id, "version": row.version, "valid_from": row.valid_from,
        "minimum_wage": row.minimum_wage, "shi_cap_multiplier": row.shi_cap_multiplier,
        "employee_rates": row.employee_rates, "employer_rates": row.employer_rates,
        "pit_brackets": row.pit_brackets, "relief_tiers": row.relief_tiers,
        "overtime_multipliers": row.overtime_multipliers, "source_references": row.source_references,
    })


def _rule_validation_issues(data: MonthlyRuleSetInput) -> list[str]:
    issues: list[str] = []
    if data.valid_to is not None and data.valid_to < data.valid_from:
        issues.append("valid_to_before_valid_from")
    if any(not source.strip() for source in data.source_references):
        issues.append("source_reference_required")
    if not data.employee_rates or not data.employer_rates:
        issues.append("shi_rates_required")
    if any(rate < 0 or rate > 1 for rate in (*data.employee_rates.values(), *data.employer_rates.values())):
        issues.append("shi_rate_out_of_range")
    required_overtime = {"weekday", "rest_day", "public_holiday"}
    if set(data.overtime_multipliers) != required_overtime or any(value <= 0 or value > 10 for value in data.overtime_multipliers.values()):
        issues.append("overtime_buckets_invalid")
    for field_name, brackets, has_rate in (("pit", data.pit_brackets, True), ("relief", data.relief_tiers, False)):
        previous_upper: Decimal | None = None
        for index, bracket in enumerate(brackets):
            try:
                lower = amount(bracket["lower"])
                upper = amount(bracket["upper"]) if bracket.get("upper") is not None else None
                value = amount(bracket["rate"] if has_rate else bracket["amount"])
                if lower < 0 or (upper is not None and upper <= lower) or value < 0:
                    raise ValueError
                if has_rate and (value > 1 or amount(bracket.get("base_tax", 0)) < 0):
                    raise ValueError
                if index == 0 and lower != 0:
                    issues.append(f"{field_name}_must_start_at_zero")
                if index and (previous_upper is None or lower != previous_upper):
                    issues.append(f"{field_name}_tiers_must_be_contiguous")
                if index < len(brackets) - 1 and upper is None:
                    issues.append(f"{field_name}_only_last_tier_can_be_open_ended")
                previous_upper = upper
            except (KeyError, TypeError, ValueError):
                issues.append(f"{field_name}_tier_{index + 1}_invalid")
    if data.pit_brackets and data.pit_brackets[-1].get("upper") is not None:
        issues.append("pit_last_tier_must_be_open_ended")
    return sorted(set(issues))


def _rule_input_from_row(row: MonthlyPayrollRuleSet) -> dict[str, Any]:
    return {
        "id": row.id, "version": row.version, "status": row.status,
        "valid_from": row.valid_from.isoformat(), "valid_to": row.valid_to.isoformat() if row.valid_to else None,
        "minimum_wage": str(row.minimum_wage), "shi_cap_multiplier": str(row.shi_cap_multiplier),
        "employee_rates": row.employee_rates, "employer_rates": row.employer_rates,
        "pit_brackets": row.pit_brackets, "relief_tiers": row.relief_tiers,
        "overtime_multipliers": row.overtime_multipliers, "source_references": row.source_references,
        "validation_issues": row.validation_issues if hasattr(row, "validation_issues") else [],
    }


def _rule_input_default() -> dict[str, Any]:
    rules = default_2026_rules()
    return {
        "valid_from": date(2026, 1, 1), "valid_to": None,
        "minimum_wage": rules.minimum_wage, "shi_cap_multiplier": rules.shi_cap_multiplier,
        "employee_rates": rules.employee_shi_rates, "employer_rates": rules.employer_shi_rates,
        "pit_brackets": [{"lower": lower, "upper": upper, "rate": rate, "base_tax": base} for lower, upper, rate, base in rules.pit_brackets],
        "relief_tiers": [{"lower": lower, "upper": upper, "amount": value} for lower, upper, value in rules.relief_tiers],
        "overtime_multipliers": rules.overtime_multipliers,
        "source_references": [
            "https://legalinfo.mn/mn/detail?lawId=14410",
            "https://legalinfo.mn/mn/detail?lawId=16760148379551",
        ],
    }


async def ensure_default_rule_set(db: AsyncSession, organization_id: int) -> None:
    """Publish the seeded 2026 rules for an organization that has no rule set yet.

    The migration seeds organizations that existed at deploy time; this covers
    organizations created later. An admin's own versions are never touched.
    """
    existing = await db.scalar(select(MonthlyPayrollRuleSet.id).where(MonthlyPayrollRuleSet.organization_id == organization_id).limit(1))
    if existing:
        return
    defaults = _json_value(_rule_input_default())
    db.add(MonthlyPayrollRuleSet(
        organization_id=organization_id, version=1, status="published",
        valid_from=date(2026, 1, 1), valid_to=None,
        minimum_wage=Decimal(defaults["minimum_wage"]), shi_cap_multiplier=Decimal(defaults["shi_cap_multiplier"]),
        employee_rates=defaults["employee_rates"], employer_rates=defaults["employer_rates"],
        pit_brackets=defaults["pit_brackets"], relief_tiers=defaults["relief_tiers"],
        overtime_multipliers=defaults["overtime_multipliers"], source_references=defaults["source_references"],
    ))
    await db.flush()


@router.get("/rule-sets")
async def list_monthly_rule_sets(db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "view")
    rows = (await db.execute(select(MonthlyPayrollRuleSet).where(
        MonthlyPayrollRuleSet.organization_id == actor.organization_id,
    ).order_by(MonthlyPayrollRuleSet.valid_from.desc(), MonthlyPayrollRuleSet.version.desc()))).scalars().all()
    return [_rule_input_from_row(row) for row in rows]


@router.get("/rule-sets/template")
async def get_monthly_rule_template(db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "view")
    return _json_value(_rule_input_default())


@router.post("/rule-sets", status_code=status.HTTP_201_CREATED)
async def create_monthly_rule_draft(data: MonthlyRuleSetInput, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "administer")
    values = data.model_dump()
    if not values["pit_brackets"]:
        values["pit_brackets"] = _rule_input_default()["pit_brackets"]
    current_version = await db.scalar(select(MonthlyPayrollRuleSet.version).where(
        MonthlyPayrollRuleSet.organization_id == actor.organization_id,
    ).order_by(MonthlyPayrollRuleSet.version.desc()).limit(1)) or 0
    row = MonthlyPayrollRuleSet(
        organization_id=actor.organization_id, version=current_version + 1, status="draft",
        valid_from=data.valid_from, valid_to=data.valid_to, minimum_wage=data.minimum_wage,
        shi_cap_multiplier=data.shi_cap_multiplier,
        employee_rates=_json_value(data.employee_rates), employer_rates=_json_value(data.employer_rates),
        pit_brackets=_json_value(data.pit_brackets), relief_tiers=_json_value(data.relief_tiers),
        overtime_multipliers=_json_value(data.overtime_multipliers),
        source_references=[source.strip() for source in data.source_references],
    )
    db.add(row)
    await db.flush()
    await record_change(db, actor=actor, topic="payroll", aggregate_type="monthly_payroll_rule_set", aggregate_id=row.id, operation="drafted", after=_rule_input_from_row(row))
    await db.commit()
    await db.refresh(row)
    return _rule_input_from_row(row)


@router.put("/rule-sets/{rule_id}")
async def update_monthly_rule_draft(rule_id: int, data: MonthlyRuleSetInput, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "administer")
    row = await db.scalar(select(MonthlyPayrollRuleSet).where(
        MonthlyPayrollRuleSet.id == rule_id,
        MonthlyPayrollRuleSet.organization_id == actor.organization_id,
    ).with_for_update())
    if row is None:
        raise HTTPException(status_code=404, detail="Дүрмийн хувилбар олдсонгүй.")
    if row.status not in {"draft", "validated"}:
        raise HTTPException(status_code=409, detail="Нийтэлсэн дүрмийг засах боломжгүй. Шинэ хувилбар ноороглоно уу.")
    before = _rule_input_from_row(row)
    row.status = "draft"
    row.valid_from, row.valid_to = data.valid_from, data.valid_to
    row.minimum_wage, row.shi_cap_multiplier = data.minimum_wage, data.shi_cap_multiplier
    row.employee_rates, row.employer_rates = _json_value(data.employee_rates), _json_value(data.employer_rates)
    row.pit_brackets, row.relief_tiers = _json_value(data.pit_brackets), _json_value(data.relief_tiers)
    row.overtime_multipliers = _json_value(data.overtime_multipliers)
    row.source_references = [source.strip() for source in data.source_references]
    await record_change(db, actor=actor, topic="payroll", aggregate_type="monthly_payroll_rule_set", aggregate_id=row.id, operation="draft_updated", before=before, after=_rule_input_from_row(row))
    await db.commit()
    await db.refresh(row)
    return _rule_input_from_row(row)


@router.post("/rule-sets/{rule_id}/validate")
async def validate_monthly_rule_draft(rule_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "administer")
    row = await db.scalar(select(MonthlyPayrollRuleSet).where(
        MonthlyPayrollRuleSet.id == rule_id, MonthlyPayrollRuleSet.organization_id == actor.organization_id,
    ).with_for_update())
    if row is None:
        raise HTTPException(status_code=404, detail="Дүрмийн хувилбар олдсонгүй.")
    if row.status not in {"draft", "validated"}:
        raise HTTPException(status_code=409, detail="Зөвхөн ноорог дүрмийг шалгана.")
    payload = MonthlyRuleSetInput(
        valid_from=row.valid_from, valid_to=row.valid_to, minimum_wage=row.minimum_wage,
        shi_cap_multiplier=row.shi_cap_multiplier, employee_rates=row.employee_rates,
        employer_rates=row.employer_rates, pit_brackets=row.pit_brackets, relief_tiers=row.relief_tiers,
        overtime_multipliers=row.overtime_multipliers, source_references=row.source_references,
    )
    issues = _rule_validation_issues(payload)
    if not issues:
        overlaps = (await db.execute(select(MonthlyPayrollRuleSet).where(
            MonthlyPayrollRuleSet.organization_id == actor.organization_id,
            MonthlyPayrollRuleSet.status == "published", MonthlyPayrollRuleSet.id != row.id,
        ))).scalars().all()
        for current in overlaps:
            candidate_end = payload.valid_to or date.max
            current_end = current.valid_to or date.max
            if payload.valid_from <= current_end and current.valid_from <= candidate_end:
                # A later-dated version supersedes the earlier rule's open
                # tail when published. A rule cannot be inserted before or
                # at the start of an already published successor.
                if current.valid_from >= payload.valid_from:
                    issues.append(f"published_effective_period_overlaps_version_{current.version}")
    before = row.status
    row.status = "validated" if not issues else "draft"
    await record_change(db, actor=actor, topic="payroll", aggregate_type="monthly_payroll_rule_set", aggregate_id=row.id, operation="validated", before={"status": before}, after={"status": row.status, "issues": issues})
    await db.commit()
    await db.refresh(row)
    return {**_rule_input_from_row(row), "validation_issues": issues}


@router.post("/rule-sets/{rule_id}/publish")
async def publish_monthly_rule_draft(rule_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "administer")
    row = await db.scalar(select(MonthlyPayrollRuleSet).where(
        MonthlyPayrollRuleSet.id == rule_id, MonthlyPayrollRuleSet.organization_id == actor.organization_id,
    ).with_for_update())
    if row is None:
        raise HTTPException(status_code=404, detail="Дүрмийн хувилбар олдсонгүй.")
    if row.status != "validated":
        raise HTTPException(status_code=409, detail="Нийтлэхийн өмнө дүрмийг шалгана уу.")
    before = _rule_input_from_row(row)
    published = (await db.execute(select(MonthlyPayrollRuleSet).where(
        MonthlyPayrollRuleSet.organization_id == actor.organization_id,
        MonthlyPayrollRuleSet.status == "published", MonthlyPayrollRuleSet.id != row.id,
    ).with_for_update())).scalars().all()
    for previous in published:
        previous_end = previous.valid_to or date.max
        if previous.valid_from < row.valid_from <= previous_end:
            old_valid_to = previous.valid_to
            previous.valid_to = row.valid_from - timedelta(days=1)
            await record_change(
                db, actor=actor, topic="payroll", aggregate_type="monthly_payroll_rule_set",
                aggregate_id=previous.id, operation="effective_period_superseded",
                before={"valid_to": old_valid_to.isoformat() if old_valid_to else None},
                after={"valid_to": previous.valid_to.isoformat(), "successor_version": row.version},
            )
    row.status = "published"
    await record_change(db, actor=actor, topic="payroll", aggregate_type="monthly_payroll_rule_set", aggregate_id=row.id, operation="published", before=before, after=_rule_input_from_row(row))
    await db.commit()
    await db.refresh(row)
    return _rule_input_from_row(row)


@router.get("/settings")
async def get_monthly_settings(db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "view")
    row = await db.scalar(select(MonthlyPayrollCompanySettings).where(MonthlyPayrollCompanySettings.organization_id == actor.organization_id))
    if row is None:
        return {"legal_company_name": None, "daily_norm_hours": "8", "employer_injury_rate": "0.005", "weekday_overtime_multiplier": "1.5", "rest_day_overtime_multiplier": "1.5", "public_holiday_overtime_multiplier": "2", "default_advance_basis": "FIXED", "default_advance_percent": "40", "deduction_types": ["Торгууль / сахилгын шийтгэл", "Хохирол / ажилтнаас авах авлага", "Бусад"], "salary_expense_account_id": None, "employer_shi_account_id": None, "advance_clearing_account_id": None}
    fields = ("legal_company_name", "daily_norm_hours", "employer_injury_rate", "weekday_overtime_multiplier", "rest_day_overtime_multiplier", "public_holiday_overtime_multiplier", "default_advance_basis", "default_advance_percent", "deduction_types", "salary_expense_account_id", "employer_shi_account_id", "advance_clearing_account_id")
    account_fields = {"salary_expense_account_id", "employer_shi_account_id", "advance_clearing_account_id"}
    return {key: _plain_value(getattr(row, key)) if key not in {"deduction_types", "legal_company_name", *account_fields} and getattr(row, key) is not None else getattr(row, key) for key in fields}


@router.put("/settings")
async def save_monthly_settings(data: CompanyMonthlySettingsInput, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "administer")
    row = await db.scalar(select(MonthlyPayrollCompanySettings).where(MonthlyPayrollCompanySettings.organization_id == actor.organization_id).with_for_update())
    if row is None:
        row = MonthlyPayrollCompanySettings(organization_id=actor.organization_id)
        db.add(row)
    account_ids = {value for value in (data.salary_expense_account_id, data.employer_shi_account_id, data.advance_clearing_account_id) if value is not None}
    if account_ids:
        valid_ids = set((await db.execute(select(ERPAccount.id).where(ERPAccount.organization_id == actor.organization_id, ERPAccount.id.in_(account_ids), ERPAccount.is_active.is_(True), ERPAccount.is_group.is_(False)))).scalars().all())
        if valid_ids != account_ids:
            raise HTTPException(status_code=422, detail="Сонгосон данс танай байгууллагын идэвхтэй данс биш байна.")
    before = {key: getattr(row, key) for key in ("legal_company_name", "daily_norm_hours", "employer_injury_rate", "weekday_overtime_multiplier", "rest_day_overtime_multiplier", "public_holiday_overtime_multiplier", "default_advance_basis", "default_advance_percent", "deduction_types", "salary_expense_account_id", "employer_shi_account_id", "advance_clearing_account_id")}
    for key, value in data.model_dump().items():
        setattr(row, key, _json_value(value))
    await db.flush()
    await record_change(db, actor=actor, topic="payroll", aggregate_type="monthly_payroll_company_settings", aggregate_id=row.id, operation="updated", before=_json_value(before), after=_json_value({key: getattr(row, key) for key in before}))
    await db.commit()
    return await get_monthly_settings(db, actor)


@router.get("/calendar/{year}/{month_num}")
async def get_monthly_calendar(year: int, month_num: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "view")
    if month_num < 1 or month_num > 12:
        raise HTTPException(status_code=422, detail="Сарыг 1-12 хооронд сонгоно уу.")
    first, last = date(year, month_num, 1), date(year, month_num, calendar.monthrange(year, month_num)[1])
    rows = (await db.execute(select(MonthlyPayrollCalendarDay).where(
        MonthlyPayrollCalendarDay.organization_id == actor.organization_id,
        MonthlyPayrollCalendarDay.calendar_date.between(first, last),
    ))).scalars().all()
    overrides = {row.calendar_date: row.day_type for row in rows}
    days = month_calendar(year, month_num, overrides=overrides)
    names = {row.calendar_date: row.holiday_name for row in rows}
    return [{"date": day.isoformat(), "day_type": day_type.value, "holiday_name": names.get(day)} for day, day_type in days.items()]


@router.put("/calendar/{calendar_date}")
async def set_monthly_calendar_day(calendar_date: date, data: CalendarDayInput, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "administer")
    month = await db.scalar(select(MonthlyPayrollMonth).where(
        MonthlyPayrollMonth.organization_id == actor.organization_id,
        MonthlyPayrollMonth.year == calendar_date.year, MonthlyPayrollMonth.month == calendar_date.month,
    ))
    if month:
        raise HTTPException(status_code=409, detail="Сарын хуанли бодолтод хадгалагдсан тул өөрчлөх боломжгүй.")
    row = await db.scalar(select(MonthlyPayrollCalendarDay).where(
        MonthlyPayrollCalendarDay.organization_id == actor.organization_id,
        MonthlyPayrollCalendarDay.calendar_date == calendar_date,
    ).with_for_update())
    if row is None:
        row = MonthlyPayrollCalendarDay(organization_id=actor.organization_id, calendar_date=calendar_date, day_type=data.day_type, holiday_name=data.holiday_name, updated_by_account_id=actor.account_id)
        db.add(row)
    else:
        row.day_type = data.day_type
        row.holiday_name = data.holiday_name
        row.updated_by_account_id = actor.account_id
    await db.commit()
    return {"date": calendar_date.isoformat(), "day_type": data.day_type, "holiday_name": data.holiday_name}


def _month_out(month: MonthlyPayrollMonth, runs: list[MonthlyPayrollRun]) -> dict[str, Any]:
    return {
        "id": month.id, "year": month.year, "month": month.month, "status": month.status,
        "rule_set_id": month.rule_set_id, "rule_snapshot": month.rule_snapshot,
        "calendar_snapshot": month.calendar_snapshot,
        "runs": [_run_out(run) for run in runs],
    }


def _run_out(run: MonthlyPayrollRun) -> dict[str, Any]:
    return {
        "id": run.id, "month_id": run.month_id, "run_type": run.run_type,
        "pay_date": run.pay_date.isoformat(), "cutoff_date": run.cutoff_date.isoformat() if run.cutoff_date else None,
        "department_id": run.department_id, "status": run.status, "note": run.note,
        "advance_snapshot": run.advance_snapshot or {},
        "approved_at": run.approved_at.isoformat() if run.approved_at else None,
        "paid_at": run.paid_at.isoformat() if run.paid_at else None,
        "approved_by_account_id": run.approved_by_account_id,
        "paid_by_account_id": run.paid_by_account_id,
        "created_by_account_id": run.created_by_account_id,
    }


def _row_out(row: MonthlyPayrollRunRow, *, include_payout_snapshot: bool = False) -> dict[str, Any]:
    result = {
        "id": row.id, "employee_id": row.employee_id, "status": row.status,
        "identity": row.identity_snapshot or {}, "profile": row.profile_snapshot or {},
        "inputs": row.inputs or {}, "result": row.result or {}, "warnings": row.warnings or [],
        "approved_at": row.approved_at.isoformat() if row.approved_at else None,
        "approved_by_account_id": row.approved_by_account_id,
    }
    if include_payout_snapshot:
        result["payout_snapshot_ciphertext"] = row.payout_snapshot_ciphertext
    return result


async def _month(db: AsyncSession, actor: ActorContext, month_id: int, *, lock: bool = False) -> MonthlyPayrollMonth:
    query = select(MonthlyPayrollMonth).where(
        MonthlyPayrollMonth.id == month_id,
        MonthlyPayrollMonth.organization_id == actor.organization_id,
    )
    if lock:
        query = query.with_for_update()
    row = await db.scalar(query)
    if row is None:
        raise HTTPException(status_code=404, detail="Payroll month not found")
    return row


async def _run(db: AsyncSession, actor: ActorContext, run_id: int, *, lock: bool = False) -> MonthlyPayrollRun:
    query = select(MonthlyPayrollRun).where(
        MonthlyPayrollRun.id == run_id,
        MonthlyPayrollRun.organization_id == actor.organization_id,
    )
    if lock:
        query = query.with_for_update()
    row = await db.scalar(query)
    if row is None:
        raise HTTPException(status_code=404, detail="Payroll run not found")
    if lock:
        # Every write locks its run: a closed month is read-only for all runs,
        # including an advance run that was waived at close.
        month_status = await db.scalar(select(MonthlyPayrollMonth.status).where(MonthlyPayrollMonth.id == row.month_id))
        if month_status != "open":
            raise HTTPException(status_code=409, detail={"code": "monthly_payroll_month_closed", "message": "Хаагдсан сарын бодолтыг өөрчлөх боломжгүй."})
    return row


async def _profiles_for_month(db: AsyncSession, actor: ActorContext, month: MonthlyPayrollMonth, department_id: int | None):
    first = date(month.year, month.month, 1)
    last = date(month.year, month.month, calendar.monthrange(month.year, month.month)[1])
    query = select(Employee, EmployeeDetails, MonthlyPayrollProfile).outerjoin(
        MonthlyPayrollProfile,
        (MonthlyPayrollProfile.employee_id == Employee.id) & (MonthlyPayrollProfile.organization_id == actor.organization_id),
    ).outerjoin(
        EmployeeDetails,
        (EmployeeDetails.employee_id == Employee.id) & (EmployeeDetails.organization_id == actor.organization_id),
    ).where(
        Employee.organization_id == actor.organization_id,
        Employee.deleted_at.is_(None),
        or_(Employee.is_active.is_(True), EmployeeDetails.end_date.is_not(None)),
    )
    if department_id is not None:
        query = query.where(EmployeeDetails.department_id == department_id)
    workers = (await db.execute(query.order_by(Employee.name))).all()
    output = []
    for employee, details, profile in workers:
        # Include recently departed and not-yet-configured workers whenever
        # their employment overlapped this earning month. They must appear as
        # actionable payroll validation rows rather than disappear silently.
        if details and ((details.start_date and details.start_date > last) or (details.end_date and details.end_date < first)):
            continue
        if profile is None:
            output.append((employee, details, None, None, []))
            continue
        history = (await db.execute(select(MonthlyPayrollSalaryHistory).where(
            MonthlyPayrollSalaryHistory.profile_id == profile.id,
            MonthlyPayrollSalaryHistory.valid_from <= last,
        ).order_by(MonthlyPayrollSalaryHistory.valid_from, MonthlyPayrollSalaryHistory.id))).scalars().all()
        active_history = [item for item in history if item.valid_from <= last]
        current_salary = active_history[-1] if active_history else None
        # Segments cover only the employment window inside this month, so a
        # mid-month hire or leaver is paid for their own days only.
        employed_from = max(first, details.start_date) if details and details.start_date else first
        employed_to = min(last, details.end_date) if details and details.end_date else last
        segments = []
        for position, item in enumerate(active_history):
            # The earliest salary has no predecessor, so it covers the month
            # from employment start even when HR dated it mid-month.
            segment_start = employed_from if position == 0 else max(employed_from, item.valid_from)
            later = [entry.valid_from for entry in active_history if entry.valid_from > item.valid_from]
            segment_end = min(employed_to, (min(later) - date.resolution) if later else last)
            if segment_end < segment_start:
                continue
            segments.append({"monthly_salary": str(item.monthly_salary), "valid_from": segment_start.isoformat(), "valid_to": segment_end.isoformat()})
        covered_from = min((date.fromisoformat(item["valid_from"]) for item in segments), default=None)
        output.append((employee, details, profile, current_salary if covered_from is not None and covered_from <= employed_from else None, segments))
    return output


async def _attendance_inputs(db: AsyncSession, organization_id: int, employee_id: int, month: MonthlyPayrollMonth, cutoff: date | None, daily_norm_hours: Decimal, overrides: dict[date, str]) -> dict[str, Any]:
    first = date(month.year, month.month, 1)
    last = date(month.year, month.month, calendar.monthrange(month.year, month.month)[1])
    end = min(last, cutoff) if cutoff else last
    # Unconfirmed logs are included: worktime sync never confirms them, and a
    # draft run is reviewed by the accountant anyway. Confirmed logs win below.
    logs = (await db.execute(select(AttendanceLog).where(
        AttendanceLog.organization_id == organization_id, AttendanceLog.employee_id == employee_id,
        AttendanceLog.attendance_date >= first, AttendanceLog.attendance_date <= end,
    ).order_by(AttendanceLog.attendance_date))).scalars().all()
    from zoneinfo import ZoneInfo
    local_zone = ZoneInfo("Asia/Ulaanbaatar")
    log_by_day = {log.attendance_date: log for log in logs}
    holiday_days = set((await db.execute(select(HolidayRecord.holiday_date).where(
        HolidayRecord.organization_id == organization_id, HolidayRecord.is_active.is_(True),
        HolidayRecord.holiday_date >= first, HolidayRecord.holiday_date <= end,
    ))).scalars().all())
    calendar_overrides = dict(overrides)
    for holiday_day in holiday_days:
        calendar_overrides.setdefault(holiday_day, CalendarDayType.PUBLIC_HOLIDAY.value)
    calendar_days = month_calendar(month.year, month.month, overrides=calendar_overrides)
    leave_rows = (await db.execute(select(TimeOff).where(
        TimeOff.organization_id == organization_id, TimeOff.employee_id == employee_id,
        TimeOff.status == "approved", TimeOff.starts_on <= end, TimeOff.ends_on >= first,
    ))).scalars().all()
    approved_leave = [{"id": row.id, "start": row.starts_on, "end": row.ends_on, "type": row.time_off_type, "minutes": row.partial_day_minutes} for row in leave_rows]
    # Worker time entries stay "pending" unless someone rejects them, so every
    # finished, non-rejected work interval is payroll evidence.
    work_rows = (await db.execute(select(WorkTimeEntry).join(Employee, Employee.id == WorkTimeEntry.employee_id).where(
        Employee.organization_id == organization_id, WorkTimeEntry.employee_id == employee_id,
        WorkTimeEntry.entry_type == "work", WorkTimeEntry.approval_status != "rejected",
        WorkTimeEntry.ended_at.is_not(None),
        WorkTimeEntry.started_at >= datetime.combine(first, datetime.min.time(), local_zone),
        WorkTimeEntry.started_at < datetime.combine(end + timedelta(days=1), datetime.min.time(), local_zone),
    ))).scalars().all()
    approved_work: dict[date, Decimal] = {}
    for entry in work_rows:
        day = entry.local_work_date or entry.started_at.astimezone(local_zone).date()
        if first <= day <= end:
            minutes = Decimal(str(max(0, (entry.ended_at - entry.started_at).total_seconds() / 60)))
            approved_work[day] = approved_work.get(day, Decimal("0")) + minutes
    details = await db.scalar(select(EmployeeDetails).where(EmployeeDetails.organization_id == organization_id, EmployeeDetails.employee_id == employee_id))
    schedule = await db.scalar(select(Schedule).where(Schedule.employee_id == employee_id))
    weekdays = tuple(dict.fromkeys(max(0, min(6, int(day) - 1)) for day in (schedule.weekdays if schedule and schedule.weekdays else [1, 2, 3, 4, 5]))) or (0, 1, 2, 3, 4)
    normalized = payment_days(
        first, end, employment_start=details.start_date if details else None,
        employment_end=details.end_date if details else None, holidays=holiday_days,
        attendance={day: {"id": row.id, "status": row.status, "minutes": row.worked_minutes} for day, row in log_by_day.items()},
        leaves=approved_leave, approved_work=approved_work, validate_attendance=True,
        workweek=weekdays, hours_per_day=float(daily_norm_hours),
    )
    leave_fraction = {date.fromisoformat(item["date"]): Decimal(item["fraction"]) for item in normalized["days"]}
    normal_hours = Decimal("0")
    overtime = {"weekday": Decimal("0"), "rest_day": Decimal("0"), "public_holiday": Decimal("0")}
    day_minutes: dict[date, Decimal] = {}
    day_sources: dict[date, str] = {}
    worked_statuses = {"present", "remote", "late"}
    # Priority per day: an HR-confirmed log, then the worker's own time
    # entries, then an unconfirmed (worktime-synced) log.
    for day, log in log_by_day.items():
        if not log.confirmed_at:
            continue
        if log.status == "half_day":
            # Minutes here are final; skip the 0.5 payable-day fraction.
            day_minutes[day] = Decimal(log.worked_minutes) if log.worked_minutes > 0 else daily_norm_hours * 30
            day_sources[day] = "attendance"
            leave_fraction[day] = Decimal("1")
            continue
        if log.status not in worked_statuses:
            day_minutes[day] = Decimal("0")
            continue
        if log.worked_minutes > 0:
            day_minutes[day], day_sources[day] = Decimal(log.worked_minutes), log.source
        elif day in approved_work:
            day_minutes[day], day_sources[day] = approved_work[day], "worktime"
        elif calendar_days[day] is CalendarDayType.WORKING:
            # HR marked the day worked without clock times: one norm day.
            day_minutes[day], day_sources[day] = daily_norm_hours * 60, "attendance"
    for day, minutes in approved_work.items():
        if day not in day_minutes:
            day_minutes[day], day_sources[day] = minutes, "worktime"
    for day, log in log_by_day.items():
        if day not in day_minutes and log.status in worked_statuses and log.worked_minutes > 0:
            day_minutes[day], day_sources[day] = Decimal(log.worked_minutes), log.source
    day_minutes = {day: minutes for day, minutes in day_minutes.items() if minutes > 0}
    lines = []
    for work_day, minutes in sorted(day_minutes.items()):
        total_hours = (minutes / Decimal(60) * leave_fraction.get(work_day, Decimal("1"))).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        normal, buckets = classify_work_hours(calendar_days[work_day], total_hours, daily_norm_hours)
        normal_hours += normal
        for key, value in buckets.items():
            overtime[key] += value
        lines.append({"date": work_day.isoformat(), "day_type": calendar_days[work_day].value, "hours": str(total_hours), "normal_hours": str(normal), "overtime_hours": {key: str(value) for key, value in buckets.items() if value > 0}, "source": day_sources[work_day]})
    planned_days_to_cutoff = sum(day <= end and day_type is CalendarDayType.WORKING for day, day_type in calendar_days.items())
    worked_hours, remaining_hours = normal_hours, Decimal("0")
    # A day counts once it has any recorded work, rest days and holidays too.
    worked_days = len(lines)
    if cutoff:
        # Advance runs project the whole month: hours worked through the
        # cut-off plus the planned hours still ahead. Without any time data
        # the worker is assumed to work the planned month.
        employed_from = max(first, details.start_date) if details and details.start_date else first
        employed_to = min(last, details.end_date) if details and details.end_date else last
        remaining_from = max(employed_from, end + timedelta(days=1)) if lines else employed_from
        remaining_days = sum(remaining_from <= day <= employed_to and day_type is CalendarDayType.WORKING for day, day_type in calendar_days.items())
        remaining_hours = daily_norm_hours * remaining_days
        worked_hours = normal_hours + remaining_hours
        worked_days += remaining_days
    return {"worked_normal_hours": _plain_value(worked_hours) if cutoff else str(worked_hours), "worked_days": str(worked_days), "worked_to_date_hours": str(normal_hours), "worked_days_to_date": str(len(lines)), "projected_remaining_hours": _plain_value(remaining_hours), "elapsed_planned_days": planned_days_to_cutoff, "overtime_hours": {key: str(value) for key, value in overtime.items()}, "day_lines": lines, "missing_dates": normalized["missing_dates"], "approved_leave_days": normalized["days"], "time_source": "confirmed_hr_attendance_or_approved_worktime" if lines else "manual"}


async def _approved_advances(db: AsyncSession, run: MonthlyPayrollRun, employee_id: int, *, through_date: date | None = None) -> tuple[list[str], list[str]]:
    runs = (await db.execute(select(MonthlyPayrollRun).where(
        MonthlyPayrollRun.month_id == run.month_id,
        MonthlyPayrollRun.organization_id == run.organization_id,
        MonthlyPayrollRun.run_type == "advance",
        MonthlyPayrollRun.id != run.id,
        MonthlyPayrollRun.status.in_(("approved", "paid", "closed")),
        MonthlyPayrollRun.pay_date <= (through_date or run.pay_date),
    ).order_by(MonthlyPayrollRun.pay_date, MonthlyPayrollRun.id))).scalars().all()
    values, ids = [], []
    for advance_run in runs:
        row = await db.scalar(select(MonthlyPayrollRunRow).where(
            MonthlyPayrollRunRow.run_id == advance_run.id,
            MonthlyPayrollRunRow.employee_id == employee_id,
        ))
        if row and row.result:
            values.append(str(row.result.get("advance", "0")))
            ids.append(str(advance_run.id))
    return values, ids


async def _calculate_row(db: AsyncSession, month: MonthlyPayrollMonth, run: MonthlyPayrollRun, row: MonthlyPayrollRunRow) -> dict[str, Any]:
    profile_data = row.profile_snapshot
    inputs = dict(row.inputs or {})
    # Blank cells from the register or an import mean zero, never a crash.
    for key in ("worked_normal_hours", "worked_to_date_hours", "leave_pay", "bonus"):
        if inputs.get(key) in (None, ""):
            inputs[key] = "0"
    inputs["overtime_hours"] = {bucket: hours for bucket, hours in (inputs.get("overtime_hours") or {}).items() if hours not in (None, "")}
    if not profile_data.get("complete", True):
        row.result = {
            "run_type": run.run_type, "gross": "0", "advance": "0",
            "employee_shi": "0", "employer_shi": "0", "pit": "0",
            "net_pay": "0", "total_deductions": "0",
        }
        row.warnings = list(profile_data.get("validation_issues") or ["profile_incomplete"])
        if inputs.get("_hr_pending"):
            row.warnings.append("hr_changed")
        return row.result
    profile = PayrollProfile(
        base_salary=amount(profile_data["base_salary"]), salary_type=profile_data["salary_type"],
        meal_allowance=amount(profile_data.get("meal_allowance", 0)), commute_allowance=amount(profile_data.get("commute_allowance", 0)),
        # Snapshots taken before daily allowances keep their monthly amounts.
        allowance_basis=AllowanceBasis(profile_data.get("allowance_basis") or "MONTHLY"),
        payment_frequency=profile_data["payment_frequency"], pay_days=tuple(profile_data["pay_days"]),
        advance_basis=AdvanceBasis(profile_data["advance_basis"]), advance_amount=amount(profile_data.get("advance_amount", 0)),
        advance_percent=amount(profile_data.get("advance_percent", 40)), daily_norm_hours=amount(profile_data.get("daily_norm_hours", 8)),
        insured_type=profile_data.get("insured_type", "01001"), tax_relief_eligible=bool(profile_data.get("tax_relief_eligible", False)),
    )
    advance_not_due = False
    advance_reference: dict[str, Any] = {}
    if run.run_type == "advance":
        basis = AdvanceBasis(inputs.get("advance_basis") or profile.advance_basis.value)
        adjustments = {**asdict(profile), "advance_basis": basis}
        advance_days = list(profile.pay_days[:-1]) if profile.payment_frequency != "MONTHLY" else []
        slot = next((index for index, day in enumerate(advance_days) if min(day, calendar.monthrange(month.year, month.month)[1]) == run.pay_date.day), None)
        # Pay days are clamped to month end here, so the engine's raw-day
        # check is not used; one-off advances are allowed off schedule.
        advance_not_due = slot is None and not inputs.get("one_off_advance")
        slot_values = profile_data.get("advance_values") or []
        slot_value = amount(slot_values[slot]) if slot is not None and slot < len(slot_values) and slot_values[slot] not in (None, "") else None
        if slot_value is not None and basis is AdvanceBasis.FIXED:
            adjustments["advance_amount"] = slot_value
        elif slot_value is not None and basis is AdvanceBasis.PERCENT:
            adjustments["advance_percent"] = slot_value
        if inputs.get("fixed_advance") is not None:
            adjustments["advance_amount"] = amount(inputs["fixed_advance"])
        if inputs.get("advance_percent") is not None:
            adjustments["advance_percent"] = amount(inputs["advance_percent"])
        profile = PayrollProfile(**adjustments)
        advance_reference = {
            "advance_basis": basis.value,
            "advance_value": str(profile.advance_amount if basis is AdvanceBasis.FIXED else profile.advance_percent if basis is AdvanceBasis.PERCENT else ""),
            "scheduled_slot": slot,
        }
    rule_snapshot = month.rule_snapshot
    rules = _rules(rule_snapshot)
    calendar_snapshot = month.calendar_snapshot
    year, month_num = month.year, month.month
    overrides = {date.fromisoformat(key): value for key, value in calendar_snapshot.items()}
    days = month_calendar(year, month_num, overrides=overrides)
    planned_days = sum(value is CalendarDayType.WORKING for value in days.values())
    planned_hours = sum((amount(profile.daily_norm_hours) for value in days.values() if value is CalendarDayType.WORKING), Decimal("0"))
    segments_raw = profile_data.get("salary_segments") or []
    segment_objects = []
    employed_days = planned_days
    if segments_raw:
        raw_segments = []
        total_segment_days = 0
        for segment in segments_raw:
            start, end = date.fromisoformat(segment["valid_from"]), date.fromisoformat(segment["valid_to"])
            segment_days = sum(start <= day <= end and day_type is CalendarDayType.WORKING for day, day_type in days.items())
            if segment_days:
                raw_segments.append((segment, segment_days))
                total_segment_days += segment_days
        day_lines = inputs.get("day_lines") or []
        normal_by_segment = [Decimal("0") for _ in raw_segments]
        overtime_by_segment = [{key: Decimal("0") for key in (inputs.get("overtime_hours") or {})} for _ in raw_segments]
        for line in day_lines:
            line_day = date.fromisoformat(line["date"])
            index = next((i for i, (segment, _) in enumerate(raw_segments) if date.fromisoformat(segment["valid_from"]) <= line_day <= date.fromisoformat(segment["valid_to"])), None)
            if index is None:
                continue
            normal_by_segment[index] += amount(line.get("normal_hours", 0))
            for bucket, hours in (line.get("overtime_hours") or {}).items():
                overtime_by_segment[index][bucket] = overtime_by_segment[index].get(bucket, Decimal("0")) + amount(hours)
        has_segment_lines = any(value > 0 for value in normal_by_segment) or any(any(value > 0 for value in group.values()) for group in overtime_by_segment)
        for index, (segment, segment_days) in enumerate(raw_segments):
            fraction = Decimal(segment_days) / Decimal(max(1, total_segment_days))
            # When an accountant edits aggregate hours, retain the date-based
            # allocation if it still reconciles; otherwise scale the source
            # lines proportionally and fall back to planned-day weights.
            worked_hours = normal_by_segment[index]
            overtime_values = overtime_by_segment[index]
            if has_segment_lines:
                normal_total = sum(normal_by_segment, Decimal("0"))
                if normal_total > 0:
                    worked_hours = amount(inputs.get("worked_normal_hours", 0)) * normal_by_segment[index] / normal_total
                else:
                    worked_hours = amount(inputs.get("worked_normal_hours", 0)) * fraction
                for bucket, total in (inputs.get("overtime_hours") or {}).items():
                    source_total = sum((group.get(bucket, Decimal("0")) for group in overtime_by_segment), Decimal("0"))
                    overtime_values[bucket] = amount(total) * overtime_by_segment[index].get(bucket, Decimal("0")) / source_total if source_total > 0 else amount(total) * fraction
            else:
                worked_hours = amount(inputs.get("worked_normal_hours", 0)) * fraction
                overtime_values = {key: amount(value) * fraction for key, value in (inputs.get("overtime_hours") or {}).items()}
            segment_objects.append(SalarySegment(
                amount(segment["monthly_salary"]), segment_days,
                worked_hours, overtime_values,
            ))
        # Working days outside employment earn nothing but still count in the
        # month, so FIXED pay is prorated by days employed.
        employed_days = min(planned_days, total_segment_days)
        if total_segment_days < planned_days:
            segment_objects.append(SalarySegment(Decimal("0"), planned_days - total_segment_days, Decimal("0"), {}))
    if inputs.get("worked_days") in (None, ""):
        # Rows captured before worked days were tracked: whole norm days.
        inputs["worked_days"] = _plain_value((amount(inputs["worked_normal_hours"]) / profile.daily_norm_hours).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
    advance_snapshot = (run.advance_snapshot or {}).get(str(row.employee_id), {}) if run.run_type == "final" else {}
    approved_advances = advance_snapshot.get("amounts", [])
    advance_ids = advance_snapshot.get("run_ids", [])
    prior_advances, _ = await _approved_advances(db, run, row.employee_id) if run.run_type == "advance" else ([], [])
    deductions = inputs.get("other_deductions", [])
    for hours in (inputs.get("overtime_hours") or {}).values():
        if amount(hours) < 0:
            raise HTTPException(status_code=422, detail={"code": "payroll_hours_nonnegative", "message": "Илүү цаг сөрөг байж болохгүй."})
    for line in deductions:
        if amount(line.get("amount", 0)) <= 0:
            raise HTTPException(status_code=422, detail={"code": "payroll_deduction_positive", "message": "Бусад суутгалын дүн 0-ээс их байх ёстой."})
    result = calculate_monthly_run(
        run.run_type, profile, rules=rules, planned_days=planned_days, planned_hours=planned_hours,
        worked_normal_hours=inputs.get("worked_normal_hours") or 0,
        overtime_hours=inputs.get("overtime_hours") or {}, leave_pay=inputs.get("leave_pay") or 0,
        bonus=inputs.get("bonus") or 0, approved_advances=approved_advances,
        prior_approved_advances=prior_advances,
        other_deductions=[item["amount"] for item in deductions],
        worked_to_date_hours=inputs.get("worked_to_date_hours") or 0,
        elapsed_planned_days=int(inputs.get("elapsed_planned_days") or 0), salary_segments=segment_objects or None,
        worked_days=inputs["worked_days"], allowance_planned_days=employed_days,
        worked_days_to_date=inputs.get("worked_days_to_date"),
    )
    result_data = _json_value(asdict(result))
    result_data["advance_run_ids"] = advance_ids
    result_data["advance_lines"] = [
        {"run_id": run_id, "amount": value, "pay_date": pay_date}
        for run_id, value, pay_date in zip(advance_ids, approved_advances, advance_snapshot.get("pay_dates") or [None] * len(advance_ids))
    ]
    result_data["planned_days"] = planned_days
    result_data["planned_hours"] = str(planned_hours)
    result_data["hourly_rate"] = str((profile.base_salary / planned_hours).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
    result_data.update(advance_reference)
    if run.run_type == "final":
        result_data["overtime_lines"] = overtime_day_lines(
            day_lines=inputs.get("day_lines") or [], aggregate_hours=inputs.get("overtime_hours") or {},
            bucket_totals=result.overtime_by_bucket, salary_segments=profile_data.get("salary_segments") or [],
            base_salary=profile.base_salary, planned_hours=planned_hours, multipliers=rules.overtime_multipliers,
        )
    warnings = []
    computed_overrides = inputs.get("_computed_overrides") or {}
    if computed_overrides:
        result_data = apply_computed_overrides(
            result_data, computed_overrides, rules=rules,
            tax_relief_eligible=profile.tax_relief_eligible,
        )
        warnings.append("computed_cell_overridden")
    if run.run_type == "final":
        due_advances = list(profile.pay_days[:-1]) if profile.payment_frequency != "MONTHLY" else []
        month_end = date(year, month_num, calendar.monthrange(year, month_num)[1])
        current_rows = (await db.execute(select(MonthlyPayrollRun.id, MonthlyPayrollRun.pay_date, MonthlyPayrollRunRow.result).join(
            MonthlyPayrollRunRow, MonthlyPayrollRunRow.run_id == MonthlyPayrollRun.id,
        ).where(
            MonthlyPayrollRun.month_id == run.month_id,
            MonthlyPayrollRun.organization_id == run.organization_id,
            MonthlyPayrollRunRow.organization_id == run.organization_id,
            MonthlyPayrollRunRow.employee_id == row.employee_id,
            MonthlyPayrollRun.run_type == "advance",
            MonthlyPayrollRun.status.in_(("approved", "paid", "closed")),
            MonthlyPayrollRun.pay_date <= month_end,
        ).order_by(MonthlyPayrollRun.pay_date, MonthlyPayrollRun.id))).all()
        current_ids = [str(item.id) for item in current_rows]
        current_amounts = [str((item.result or {}).get("advance", "0")) for item in current_rows]
        if current_ids != advance_ids or current_amounts != approved_advances:
            warnings.append("advance_changed")
        scheduled_dates = {date(year, month_num, min(day, calendar.monthrange(year, month_num)[1])) for day in due_advances}
        approved_dates = {item.pay_date for item in current_rows if item.pay_date in scheduled_dates}
        if scheduled_dates - approved_dates:
            warnings.append("advance_not_calculated")
    if amount(result_data.get("net_pay", 0)) < 0:
        warnings.append("negative_final_pay")
    if advance_not_due:
        warnings.append("advance_not_due")
    if run.run_type == "advance":
        # Reference only: a full month on the HR profile, so each salary
        # segment is credited with all of its planned hours (plan §6.1).
        full_month_segments = [
            SalarySegment(segment.monthly_salary, segment.planned_days, amount(profile.daily_norm_hours) * segment.planned_days, {})
            for segment in segment_objects
        ]
        estimated = calculate_monthly_run(
            PayrollRunType.FINAL, profile, rules=rules, planned_days=planned_days, planned_hours=planned_hours,
            worked_normal_hours=planned_hours, salary_segments=full_month_segments or None,
            worked_days=employed_days, allowance_planned_days=employed_days,
        )
        result_data["estimated_net"] = str(estimated.net_pay)
        # The advance table shows the month as the final run will see it:
        # projected hours, overtime so far, leave pay, meal/commute and bonus,
        # with НДШ and ХХОАТ, so Суутгалын дүн = this month's advances + НДШ +
        # ХХОАТ. Kept under "projection" so totals never count it as payroll.
        projection = calculate_monthly_run(
            PayrollRunType.FINAL, profile, rules=rules, planned_days=planned_days, planned_hours=planned_hours,
            worked_normal_hours=inputs.get("worked_normal_hours") or 0,
            overtime_hours=inputs.get("overtime_hours") or {}, leave_pay=inputs.get("leave_pay") or 0,
            bonus=inputs.get("bonus") or 0, approved_advances=[*prior_advances, result_data["advance"]],
            salary_segments=segment_objects or None,
            worked_days=inputs["worked_days"], allowance_planned_days=employed_days,
        )
        projection_data = _json_value(asdict(projection))
        projection_data.pop("run_type", None)
        # Advances are payments against net pay, not deductions: Суутгалын дүн
        # here is НДШ + ХХОАТ + бусад, so Суутгалын дүн + Сүүл цалин = олговол
        # зохих (company salary expense). Advances are shown on their own.
        projection_data["total_deductions"] = str(whole_tugrik(projection.employee_shi + projection.pit + projection.other_deductions))
        projection_data["net_pay"] = str(whole_tugrik(projection.gross - amount(projection_data["total_deductions"])))
        projection_data["net_after_advances"] = str(projection.net_pay)
        projection_data["prior_advances"] = str(sum((amount(value) for value in prior_advances), Decimal("0")))
        projection_data["overtime_lines"] = overtime_day_lines(
            day_lines=inputs.get("day_lines") or [], aggregate_hours=inputs.get("overtime_hours") or {},
            bucket_totals=projection.overtime_by_bucket, salary_segments=profile_data.get("salary_segments") or [],
            base_salary=profile.base_salary, planned_hours=planned_hours, multipliers=rules.overtime_multipliers,
        )
        result_data["projection"] = projection_data
        if amount(result_data["advance"]) > estimated.net_pay:
            warnings.append("advance_above_estimated_net")
        if amount(result_data["advance"]) <= 0:
            warnings.append("advance_not_positive")
    if sum((amount(value) for value in inputs.get("overtime_hours", {}).values()), Decimal("0")) > 0:
        warnings.append("overtime_work")
    if profile.base_salary < rules.minimum_wage:
        warnings.append("base_below_minimum_wage")
    if amount(inputs.get("worked_normal_hours", 0)) + sum((amount(value) for value in inputs.get("overtime_hours", {}).values()), Decimal("0")) > planned_hours:
        warnings.append("worked_hours_above_planned")
    if run.run_type == "final" and profile.salary_type == "PRORATION" and amount(inputs.get("worked_normal_hours", 0)) == 0:
        warnings.append("zero_worked_hours")
    if run.run_type == "advance" and profile.advance_basis is AdvanceBasis.WORKED_TO_DATE and not inputs.get("day_lines"):
        warnings.append("worked_to_date_without_time")
    if result.shi_base and result.gross > result.shi_base:
        warnings.append("shi_cap_hit")
    if inputs.get("_hr_pending"):
        warnings.append("hr_changed")
    if any(not line.get("type") or not line.get("note") for line in deductions):
        warnings.append("deduction_details_missing")
    row.result = result_data
    row.warnings = warnings
    return result_data


async def _flag_final_advance_changes(db: AsyncSession, month_id: int, organization_id: int) -> None:
    """Mark final-run rows whose pulled advances no longer match approved advance runs.

    Called whenever an advance run enters or leaves approval, so the final
    register shows «Урьдчилгаа өөрчлөгдсөн» without waiting for a recalculation.
    """
    await db.flush()
    final = await db.scalar(select(MonthlyPayrollRun).where(MonthlyPayrollRun.month_id == month_id, MonthlyPayrollRun.organization_id == organization_id, MonthlyPayrollRun.run_type == "final"))
    if final is None:
        return
    current = await _approved_advance_snapshot(db, month_id, organization_id)
    rows = (await db.execute(select(MonthlyPayrollRunRow).where(MonthlyPayrollRunRow.run_id == final.id, MonthlyPayrollRunRow.organization_id == organization_id))).scalars().all()
    for row in rows:
        pulled = (final.advance_snapshot or {}).get(str(row.employee_id), {})
        latest = current.get(str(row.employee_id), {})
        changed = (pulled.get("run_ids") or [], pulled.get("amounts") or []) != (latest.get("run_ids") or [], latest.get("amounts") or [])
        warnings = [warning for warning in (row.warnings or []) if warning != "advance_changed"]
        if changed:
            warnings.append("advance_changed")
        if warnings != (row.warnings or []):
            row.warnings = warnings


async def _calculate_row_or_422(db: AsyncSession, month: MonthlyPayrollMonth, run: MonthlyPayrollRun, row: MonthlyPayrollRunRow) -> dict[str, Any]:
    try:
        return await _calculate_row(db, month, run, row)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail={"code": "payroll_row_invalid", "employee_id": row.employee_id, "employee_name": (row.identity_snapshot or {}).get("name"), "message": str(exc)}) from exc


def _reset_approval(row: MonthlyPayrollRunRow) -> bool:
    """Editing an approved row returns it to draft (shown as Засварласан)."""
    if row.status != "approved":
        return False
    row.status = "draft"
    row.approved_by_account_id = None
    row.approved_at = None
    return True


@router.post("/months", status_code=status.HTTP_201_CREATED)
async def create_month(data: MonthInput, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "create")
    existing = await db.scalar(select(MonthlyPayrollMonth).where(
        MonthlyPayrollMonth.organization_id == actor.organization_id,
        MonthlyPayrollMonth.year == data.year, MonthlyPayrollMonth.month == data.month,
    ))
    if existing:
        return _month_out(existing, [])
    effective = date(data.year, data.month, 1)
    await ensure_default_rule_set(db, actor.organization_id)
    rule = await db.scalar(select(MonthlyPayrollRuleSet).where(
        MonthlyPayrollRuleSet.organization_id == actor.organization_id,
        MonthlyPayrollRuleSet.status == "published", MonthlyPayrollRuleSet.valid_from <= effective,
        (MonthlyPayrollRuleSet.valid_to.is_(None) | (MonthlyPayrollRuleSet.valid_to >= effective)),
    ).order_by(MonthlyPayrollRuleSet.version.desc()).limit(1))
    if not rule:
        raise HTTPException(status_code=409, detail="Тухайн сард хүчинтэй нийтэлсэн цалингийн дүрэм алга.")
    last = date(data.year, data.month, calendar.monthrange(data.year, data.month)[1])
    overrides_rows = (await db.execute(select(MonthlyPayrollCalendarDay).where(
        MonthlyPayrollCalendarDay.organization_id == actor.organization_id,
        MonthlyPayrollCalendarDay.calendar_date >= effective,
        MonthlyPayrollCalendarDay.calendar_date <= last,
    ))).scalars().all()
    overrides = {row.calendar_date: row.day_type for row in overrides_rows}
    built = month_calendar(data.year, data.month, overrides=overrides)
    rule_snapshot = _rule_snapshot(rule)
    settings = await db.scalar(select(MonthlyPayrollCompanySettings).where(MonthlyPayrollCompanySettings.organization_id == actor.organization_id))
    if settings:
        employer_rates = dict(rule_snapshot.get("employer_rates", {}))
        employer_rates["injury"] = str(settings.employer_injury_rate)
        rule_snapshot["employer_rates"] = employer_rates
        rule_snapshot["overtime_multipliers"] = {
            "weekday": str(settings.weekday_overtime_multiplier),
            "rest_day": str(settings.rest_day_overtime_multiplier),
            "public_holiday": str(settings.public_holiday_overtime_multiplier),
        }
    month_row = MonthlyPayrollMonth(
        organization_id=actor.organization_id, year=data.year, month=data.month,
        rule_set_id=rule.id, rule_snapshot=rule_snapshot,
        calendar_snapshot={day.isoformat(): day_type.value for day, day_type in built.items()},
        created_by_account_id=actor.account_id,
    )
    db.add(month_row)
    await db.commit()
    await db.refresh(month_row)
    return _month_out(month_row, [])


@router.get("/months")
async def list_months(db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "view")
    months = (await db.execute(select(MonthlyPayrollMonth).where(
        MonthlyPayrollMonth.organization_id == actor.organization_id,
    ).order_by(MonthlyPayrollMonth.year.desc(), MonthlyPayrollMonth.month.desc()))).scalars().all()
    result = []
    for month in months:
        runs = (await db.execute(select(MonthlyPayrollRun).where(MonthlyPayrollRun.month_id == month.id).order_by(MonthlyPayrollRun.pay_date, MonthlyPayrollRun.run_type))).scalars().all()
        result.append(_month_out(month, runs))
    return result


@router.get("/months/{month_id}")
async def get_month(month_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "view")
    month = await _month(db, actor, month_id)
    runs = (await db.execute(select(MonthlyPayrollRun).where(
        MonthlyPayrollRun.month_id == month.id, MonthlyPayrollRun.organization_id == actor.organization_id,
    ).order_by(MonthlyPayrollRun.pay_date, MonthlyPayrollRun.run_type))).scalars().all()
    return _month_out(month, runs)


def _identity_snapshot(employee: Employee, details: EmployeeDetails | None, department_name: str | None, payout_date: date) -> dict[str, Any]:
    return {
        "employee_id": employee.id, "name": employee.name, "last_name": employee.last_name, "first_name": employee.first_name,
        "rd": (employee.metadata_json or {}).get("rd"), "job_title": details.job_title if details else employee.job_title,
        "department_id": details.department_id if details else None, "department": department_name or "Бусад",
        "pay_date": payout_date.isoformat(),
        "start_date": details.start_date.isoformat() if details and details.start_date else None,
        "end_date": details.end_date.isoformat() if details and details.end_date else None,
    }


def _profile_snapshot(profile: MonthlyPayrollProfile | None, history: MonthlyPayrollSalaryHistory | None, segments: list[dict[str, Any]]) -> dict[str, Any]:
    if profile is None:
        return {"complete": False, "validation_issues": ["profile_missing"], "base_salary": "0", "salary_segments": [], "salary_type": "PRORATION", "meal_allowance": "0", "commute_allowance": "0", "allowance_basis": "FIXED", "payment_frequency": "MONTHLY", "pay_days": [25], "advance_basis": "FIXED", "advance_amount": "0", "advance_percent": "40", "advance_values": [], "daily_norm_hours": "8", "insured_type": "01001", "tax_relief_eligible": True}
    issues = ["salary_history_missing_or_incomplete"] if history is None else []
    return _json_value({
        "complete": not issues, "validation_issues": issues, "base_salary": str(history.monthly_salary if history else 0),
        "salary_segments": segments, "salary_type": profile.salary_type,
        "meal_allowance": str(profile.meal_allowance), "commute_allowance": str(profile.commute_allowance),
        "allowance_basis": profile.allowance_basis,
        "payment_frequency": profile.payment_frequency, "pay_days": list(profile.pay_days or []),
        "advance_basis": profile.advance_basis, "advance_amount": str(profile.advance_amount),
        "advance_percent": str(profile.advance_percent), "advance_values": [str(value) for value in (profile.advance_values or [])],
        "daily_norm_hours": str(profile.daily_norm_hours), "insured_type": profile.insured_type,
        "tax_relief_eligible": profile.tax_relief_eligible,
    })


async def _department_name(db: AsyncSession, organization_id: int, details: EmployeeDetails | None) -> str | None:
    if not details or not details.department_id:
        return None
    return await db.scalar(select(Department.name).where(Department.id == details.department_id, Department.organization_id == organization_id))


def _row_payout_date(month: MonthlyPayrollMonth, run: MonthlyPayrollRun, profile: MonthlyPayrollProfile | None) -> date:
    last_day = calendar.monthrange(month.year, month.month)[1]
    if run.run_type == "final" and profile is not None and profile.pay_days:
        return date(month.year, month.month, min(int(profile.pay_days[-1]), last_day))
    return run.pay_date


async def _fresh_snapshots(db: AsyncSession, actor: ActorContext, month: MonthlyPayrollMonth, run: MonthlyPayrollRun, employee: Employee, details: EmployeeDetails | None, profile: MonthlyPayrollProfile | None, history: MonthlyPayrollSalaryHistory | None, segments: list[dict[str, Any]]) -> tuple[dict[str, Any], dict[str, Any]]:
    identity = _identity_snapshot(employee, details, await _department_name(db, actor.organization_id, details), _row_payout_date(month, run, profile))
    return identity, _profile_snapshot(profile, history, segments)


async def _advance_carry_over(db: AsyncSession, organization_id: int, month_id: int, employee_id: int, attendance: dict[str, Any]) -> dict[str, Any]:
    """Accountant inputs typed in the month's latest advance row, reused by the final row.

    Leave pay and bonus carry as typed. Overtime carries only for buckets the
    accountant entered by hand and the full-month time data left empty, so
    recorded attendance always wins.
    """
    source = (await db.execute(select(MonthlyPayrollRun.id, MonthlyPayrollRunRow.inputs).join(
        MonthlyPayrollRunRow, MonthlyPayrollRunRow.run_id == MonthlyPayrollRun.id,
    ).where(
        MonthlyPayrollRun.month_id == month_id, MonthlyPayrollRun.organization_id == organization_id,
        MonthlyPayrollRunRow.organization_id == organization_id, MonthlyPayrollRun.run_type == "advance",
        MonthlyPayrollRunRow.employee_id == employee_id,
    ).order_by(MonthlyPayrollRun.pay_date.desc(), MonthlyPayrollRun.id.desc()).limit(1))).first()
    if source is None:
        return {}
    inputs = source.inputs or {}
    carried: dict[str, Any] = {key: str(inputs[key]) for key in ("leave_pay", "bonus") if amount(inputs.get(key) or 0) > 0}
    typed = inputs.get("overtime_hours") or {}
    recorded = (inputs.get("_source_snapshot") or {}).get("overtime_hours") or {}
    fresh = attendance.get("overtime_hours") or {}
    manual = {bucket: str(hours) for bucket, hours in typed.items() if amount(hours or 0) > 0 and amount(hours) != amount(recorded.get(bucket) or 0) and amount(fresh.get(bucket) or 0) == 0}
    if manual:
        carried["overtime_hours"] = {**fresh, **manual}
    if carried:
        carried["carried_from_advance_run"] = source.id
    return carried


async def _build_row(db: AsyncSession, actor: ActorContext, month: MonthlyPayrollMonth, run: MonthlyPayrollRun, employee: Employee, details: EmployeeDetails | None, profile: MonthlyPayrollProfile | None, history: MonthlyPayrollSalaryHistory | None, segments: list[dict[str, Any]], *, one_off: bool = False) -> MonthlyPayrollRunRow:
    """Snapshot one worker into a run: identity, HR profile, time and payout."""
    identity, profile_data = await _fresh_snapshots(db, actor, month, run, employee, details, profile, history, segments)
    payout_date = date.fromisoformat(identity["pay_date"])
    calendar_overrides = {date.fromisoformat(key): value for key, value in month.calendar_snapshot.items()}
    attendance = await _attendance_inputs(
        db, actor.organization_id, employee.id, month,
        run.cutoff_date if run.run_type == "advance" else None,
        amount(profile_data["daily_norm_hours"]), calendar_overrides,
    )
    attendance["_source_snapshot"] = {key: attendance.get(key) for key in TIME_INPUT_KEYS}
    payout = await db.scalar(select(EmployeeBankAccount).where(
        EmployeeBankAccount.employee_id == employee.id,
        EmployeeBankAccount.is_primary.is_(True), EmployeeBankAccount.valid_from <= payout_date,
        (EmployeeBankAccount.valid_to.is_(None) | (EmployeeBankAccount.valid_to >= payout_date)),
    ).order_by(EmployeeBankAccount.valid_from.desc(), EmployeeBankAccount.id.desc()).limit(1))
    payout_snapshot = encrypt_secret(json.dumps({
        "bank_code": payout.bank_code,
        "account_number": decrypt_secret(payout.account_number_ciphertext),
        "account_holder": decrypt_secret(payout.account_holder_ciphertext) if payout.account_holder_ciphertext else employee.name,
    }, ensure_ascii=False)) if payout else None
    carried = await _advance_carry_over(db, actor.organization_id, month.id, employee.id, attendance) if run.run_type == "final" else {}
    row = MonthlyPayrollRunRow(
        organization_id=actor.organization_id, run_id=run.id, employee_id=employee.id,
        identity_snapshot=identity, profile_snapshot=profile_data,
        payout_snapshot_ciphertext=payout_snapshot,
        inputs={**attendance, "leave_pay": "0", "bonus": "0", "other_deductions": [], **carried, **({"one_off_advance": True} if one_off else {})},
    )
    db.add(row)
    return row


async def _approved_advance_snapshot(db: AsyncSession, month_id: int, organization_id: int) -> dict[str, dict[str, list[str]]]:
    """Per-worker approved advance lines of a month, in pay-date order."""
    lines = (await db.execute(select(MonthlyPayrollRun.id, MonthlyPayrollRun.pay_date, MonthlyPayrollRunRow.employee_id, MonthlyPayrollRunRow.result).join(
        MonthlyPayrollRunRow, MonthlyPayrollRunRow.run_id == MonthlyPayrollRun.id,
    ).where(
        MonthlyPayrollRun.month_id == month_id, MonthlyPayrollRun.organization_id == organization_id,
        MonthlyPayrollRunRow.organization_id == organization_id,
        MonthlyPayrollRun.run_type == "advance", MonthlyPayrollRun.status.in_(("approved", "paid", "closed")),
    ).order_by(MonthlyPayrollRun.pay_date, MonthlyPayrollRun.id))).all()
    snapshot: dict[str, dict[str, list[str]]] = {}
    for line in lines:
        item = snapshot.setdefault(str(line.employee_id), {"amounts": [], "run_ids": [], "pay_dates": []})
        item["amounts"].append(str((line.result or {}).get("advance", "0")))
        item["run_ids"].append(str(line.id))
        item["pay_dates"].append(line.pay_date.isoformat())
    return snapshot


@router.get("/months/{month_id}/advance-dates")
async def get_advance_dates(month_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "view")
    month = await _month(db, actor, month_id)
    last = date(month.year, month.month, calendar.monthrange(month.year, month.month)[1])
    workers = await _profiles_for_month(db, actor, month, None)
    counts: dict[int, dict[str, int]] = {}
    for _, _, profile, _, _ in workers:
        if profile is None or profile.payment_frequency == "MONTHLY":
            continue
        for day in (profile.pay_days or [])[:-1]:
            item = counts.setdefault(min(int(day), last.day), {"workers": 0, "worked_to_date_workers": 0})
            item["workers"] += 1
            if profile.advance_basis == AdvanceBasis.WORKED_TO_DATE.value:
                item["worked_to_date_workers"] += 1
    existing = set((await db.execute(select(MonthlyPayrollRun.pay_date).where(
        MonthlyPayrollRun.month_id == month.id, MonthlyPayrollRun.run_type == "advance",
    ))).scalars().all())
    return [{"day": day, "date": date(month.year, month.month, day).isoformat(), **item, "run_exists": date(month.year, month.month, day) in existing} for day, item in sorted(counts.items())]


@router.post("/months/{month_id}/runs", status_code=status.HTTP_201_CREATED)
async def create_run(month_id: int, data: RunInput, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "create")
    month = await _month(db, actor, month_id, lock=True)
    if month.status != "open":
        raise HTTPException(status_code=409, detail="Хаагдсан сарын бодолт өөрчлөх боломжгүй.")
    first, last = date(month.year, month.month, 1), date(month.year, month.month, calendar.monthrange(month.year, month.month)[1])
    if not first <= data.pay_date <= last:
        raise HTTPException(status_code=422, detail="Төлбөрийн өдөр сонгосон сард багтах ёстой.")
    if data.cutoff_date and data.cutoff_date > data.pay_date:
        raise HTTPException(status_code=422, detail="Таслах өдөр төлбөрийн өдрөөс хойш байж болохгүй.")
    if data.run_type == "final" and (data.employee_ids or data.department_id):
        raise HTTPException(status_code=422, detail="Сүүл цалингийн бодолт сарын бүх ажилтныг нэг хүснэгтэд хамарна; хэлтсээр хүснэгт дотор шүүнэ үү.")
    if data.run_type == "final":
        prior_final = await db.scalar(select(MonthlyPayrollRun.id).where(
            MonthlyPayrollRun.month_id == month.id, MonthlyPayrollRun.run_type == "final",
        ))
        if prior_final:
            raise HTTPException(status_code=409, detail="Сард зөвхөн нэг сүүл цалингийн бодолт үүсгэж болно.")
    else:
        duplicate = await db.scalar(select(MonthlyPayrollRun.id).where(
            MonthlyPayrollRun.month_id == month.id, MonthlyPayrollRun.run_type == "advance",
            MonthlyPayrollRun.pay_date == data.pay_date,
        ))
        if duplicate:
            raise HTTPException(status_code=409, detail="Энэ төлбөрийн өдөр урьдчилгааны бодолт байна. «Ажилтан нэмэх»-ээр нэмнэ үү.")
    workers = await _profiles_for_month(db, actor, month, data.department_id)
    eligible = []
    for worker in workers:
        employee, _details, profile, _history, _segments = worker
        # One final register covers every worker in the earning month;
        # individual payment dates stay on the HR profiles.
        if data.run_type == "advance":
            if profile is None:
                continue
            if data.employee_ids is not None:
                if employee.id not in data.employee_ids:
                    continue
            elif profile.payment_frequency == "MONTHLY" or not any(min(int(day), last.day) == data.pay_date.day for day in (profile.pay_days or [])[:-1]):
                continue
        eligible.append(worker)
    eligible_ids = {worker[0].id for worker in eligible}
    if data.employee_ids and eligible_ids != set(data.employee_ids):
        raise HTTPException(status_code=422, detail={"code": "monthly_payroll_advance_worker_invalid", "employee_ids": sorted(set(data.employee_ids) - eligible_ids), "message": "Нэг удаагийн урьдчилгаанд цалингийн профайлтай ажилтныг сонгоно уу."})
    if not eligible:
        raise HTTPException(status_code=422, detail="Төлбөрийн өдөрт тохирох идэвхтэй ажилтан, цалингийн профайл олдсонгүй.")
    run = MonthlyPayrollRun(
        organization_id=actor.organization_id, month_id=month.id, run_type=data.run_type,
        pay_date=data.pay_date, cutoff_date=data.cutoff_date or (data.pay_date - timedelta(days=1) if data.run_type == "advance" else None),
        department_id=data.department_id, note=data.note, created_by_account_id=actor.account_id,
    )
    db.add(run)
    await db.flush()
    if data.run_type == "final":
        run.advance_snapshot = await _approved_advance_snapshot(db, month.id, actor.organization_id)
    new_rows = [
        await _build_row(db, actor, month, run, employee, details, profile, history, segments, one_off=run.run_type == "advance" and data.employee_ids is not None)
        for employee, details, profile, history, segments in eligible
    ]
    await db.flush()
    for row in new_rows:
        await _calculate_row_or_422(db, month, run, row)
    await db.commit()
    await db.refresh(run)
    return _run_out(run)


@router.get("/runs/{run_id}")
async def get_run(run_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "view_salary")
    run = await _run(db, actor, run_id)
    rows = (await db.execute(select(MonthlyPayrollRunRow).where(
        MonthlyPayrollRunRow.run_id == run.id, MonthlyPayrollRunRow.organization_id == actor.organization_id,
    ).order_by(MonthlyPayrollRunRow.id))).scalars().all()
    audits = (await db.execute(select(MonthlyPayrollRowAudit).where(
        MonthlyPayrollRowAudit.organization_id == actor.organization_id,
        MonthlyPayrollRowAudit.run_id == run.id,
    ).order_by(MonthlyPayrollRowAudit.id))).scalars().all()
    by_row: dict[int, list[dict[str, Any]]] = {}
    for audit in audits:
        by_row.setdefault(audit.row_id, []).append({"field": audit.field_name, "old": audit.old_value, "new": audit.new_value, "reason": audit.reason, "account_id": audit.account_id, "at": audit.created_at.isoformat()})
    return {**_run_out(run), "rows": [{**_row_out(row), "audit": by_row.get(row.id, [])} for row in rows]}


@router.put("/runs/{run_id}/rows/{employee_id}")
async def update_row(run_id: int, employee_id: int, data: RowInput, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "create")
    run = await _run(db, actor, run_id, lock=True)
    if run.status != "draft":
        raise HTTPException(status_code=409, detail="Баталсан бодолтын мөр засах боломжгүй.")
    row = await db.scalar(select(MonthlyPayrollRunRow).where(
        MonthlyPayrollRunRow.run_id == run.id, MonthlyPayrollRunRow.employee_id == employee_id,
        MonthlyPayrollRunRow.organization_id == actor.organization_id,
    ).with_for_update())
    if not row:
        raise HTTPException(status_code=404, detail="Бодолтын ажилтан олдсонгүй.")
    _reset_approval(row)
    changes = data.model_dump(exclude_unset=True, exclude={"reason"})
    before = dict(row.inputs or {})
    after = {**before, **_json_value(changes)}
    if data.fixed_advance is not None:
        after["fixed_advance"] = str(data.fixed_advance)
    if data.advance_percent is not None:
        after["advance_percent"] = str(data.advance_percent)
    row.inputs = after
    db.add(MonthlyPayrollRowAudit(organization_id=actor.organization_id, run_id=run.id, row_id=row.id, account_id=actor.account_id, field_name="inputs", old_value=before, new_value=after, reason=data.reason))
    month = await _month(db, actor, run.month_id)
    await _calculate_row_or_422(db, month, run, row)
    await db.commit()
    await db.refresh(row)
    return _row_out(row)


@router.post("/runs/{run_id}/rows/{employee_id}/revert-overrides")
async def revert_row_overrides(run_id: int, employee_id: int, data: UnlockInput, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "create")
    run = await _run(db, actor, run_id, lock=True)
    if run.status != "draft":
        raise HTTPException(status_code=409, detail="Зөвхөн ноорог бодолтын оролтыг буцаах боломжтой.")
    row = await db.scalar(select(MonthlyPayrollRunRow).where(MonthlyPayrollRunRow.run_id == run.id, MonthlyPayrollRunRow.employee_id == employee_id, MonthlyPayrollRunRow.organization_id == actor.organization_id).with_for_update())
    if row is None:
        raise HTTPException(status_code=404, detail="Бодолтын ажилтан олдсонгүй.")
    _reset_approval(row)
    before = dict(row.inputs or {})
    source = dict(before.get("_source_snapshot") or {})
    after = {**before, **source, "leave_pay": "0", "bonus": "0", "other_deductions": []}
    for key in ("fixed_advance", "advance_percent", "advance_basis", "other_deduction_amount", "other_deduction_type", "other_deduction_note"):
        after.pop(key, None)
    row.inputs = after
    month = await _month(db, actor, run.month_id)
    await _calculate_row(db, month, run, row)
    db.add(MonthlyPayrollRowAudit(organization_id=actor.organization_id, run_id=run.id, row_id=row.id, account_id=actor.account_id, field_name="overrides_reverted", old_value=before, new_value=after, reason=data.reason.strip()))
    await db.commit()
    await db.refresh(row)
    return _row_out(row)


@router.post("/runs/{run_id}/rows/{employee_id}/computed-overrides")
async def override_computed_cell(run_id: int, employee_id: int, data: ComputedOverrideInput, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "create")
    run = await _run(db, actor, run_id, lock=True)
    if run.status != "draft":
        raise HTTPException(status_code=409, detail="Зөвхөн ноорог бодолтын тооцсон дүнг засна.")
    row = await db.scalar(select(MonthlyPayrollRunRow).where(MonthlyPayrollRunRow.run_id == run.id, MonthlyPayrollRunRow.employee_id == employee_id, MonthlyPayrollRunRow.organization_id == actor.organization_id).with_for_update())
    if row is None:
        raise HTTPException(status_code=404, detail="Бодолтын ажилтан олдсонгүй.")
    _reset_approval(row)
    if not row.profile_snapshot.get("complete", True):
        raise HTTPException(status_code=409, detail={"code": "monthly_payroll_profile_incomplete", "employee_id": employee_id})
    before_result, before_inputs = dict(row.result or {}), dict(row.inputs or {})
    overrides = dict(before_inputs.get("_computed_overrides") or {})
    overrides[data.field] = str(data.value.quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    row.inputs = {**before_inputs, "_computed_overrides": overrides}
    month = await _month(db, actor, run.month_id)
    await _calculate_row(db, month, run, row)
    db.add(MonthlyPayrollRowAudit(organization_id=actor.organization_id, run_id=run.id, row_id=row.id, account_id=actor.account_id, field_name=f"computed:{data.field}", old_value=before_result.get(data.field), new_value=row.result.get(data.field), reason=data.reason.strip()))
    await db.commit()
    await db.refresh(row)
    return _row_out(row)


@router.post("/runs/{run_id}/rows/{employee_id}/computed-overrides/revert")
async def revert_computed_cell_override(run_id: int, employee_id: int, data: ComputedOverrideRevertInput, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "create")
    run = await _run(db, actor, run_id, lock=True)
    if run.status != "draft":
        raise HTTPException(status_code=409, detail="Зөвхөн ноорог бодолтын тооцсон дүнг буцаана.")
    row = await db.scalar(select(MonthlyPayrollRunRow).where(MonthlyPayrollRunRow.run_id == run.id, MonthlyPayrollRunRow.employee_id == employee_id, MonthlyPayrollRunRow.organization_id == actor.organization_id).with_for_update())
    if row is None:
        raise HTTPException(status_code=404, detail="Бодолтын ажилтан олдсонгүй.")
    _reset_approval(row)
    before_result, inputs = dict(row.result or {}), dict(row.inputs or {})
    overrides = dict(inputs.get("_computed_overrides") or {})
    if data.field not in overrides:
        raise HTTPException(status_code=409, detail="Энэ дүнд засвар бүртгэгдээгүй байна.")
    overrides.pop(data.field)
    if overrides:
        inputs["_computed_overrides"] = overrides
    else:
        inputs.pop("_computed_overrides", None)
    row.inputs = inputs
    month = await _month(db, actor, run.month_id)
    await _calculate_row(db, month, run, row)
    db.add(MonthlyPayrollRowAudit(organization_id=actor.organization_id, run_id=run.id, row_id=row.id, account_id=actor.account_id, field_name=f"computed:{data.field}:revert", old_value=before_result.get(data.field), new_value=row.result.get(data.field), reason=data.reason.strip()))
    await db.commit()
    await db.refresh(row)
    return _row_out(row)


@router.post("/runs/{run_id}/rows/{employee_id}/flag")
async def flag_row(run_id: int, employee_id: int, data: FlagInput, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "create")
    run = await _run(db, actor, run_id, lock=True)
    if run.status != "draft":
        raise HTTPException(status_code=409, detail="Зөвхөн ноорог бодолтын мөрийг тэмдэглэнэ.")
    row = await db.scalar(select(MonthlyPayrollRunRow).where(
        MonthlyPayrollRunRow.run_id == run.id,
        MonthlyPayrollRunRow.employee_id == employee_id,
        MonthlyPayrollRunRow.organization_id == actor.organization_id,
    ).with_for_update())
    if not row:
        raise HTTPException(status_code=404, detail="Бодолтын ажилтан олдсонгүй.")
    if row.status not in {"draft", "flagged"}:
        raise HTTPException(status_code=409, detail="Баталсан мөрийг тэмдэглэх боломжгүй.")
    before = {"status": row.status, "flag_reason": (row.inputs or {}).get("flag_reason")}
    row.status = "flagged"
    row.inputs = {**(row.inputs or {}), "flag_reason": data.reason.strip()}
    row.warnings = list(dict.fromkeys([*(row.warnings or []), "row_flagged"]))
    db.add(MonthlyPayrollRowAudit(organization_id=actor.organization_id, run_id=run.id, row_id=row.id, account_id=actor.account_id, field_name="status", old_value=before, new_value={"status": "flagged", "flag_reason": data.reason.strip()}, reason=data.reason.strip()))
    await db.commit()
    return _row_out(row)


@router.post("/runs/{run_id}/rows/{employee_id}/unflag")
async def unflag_row(run_id: int, employee_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "create")
    run = await _run(db, actor, run_id, lock=True)
    if run.status != "draft":
        raise HTTPException(status_code=409, detail="Зөвхөн ноорог бодолтын мөрийн тэмдэглэгээг арилгана.")
    row = await db.scalar(select(MonthlyPayrollRunRow).where(
        MonthlyPayrollRunRow.run_id == run.id,
        MonthlyPayrollRunRow.employee_id == employee_id,
        MonthlyPayrollRunRow.organization_id == actor.organization_id,
    ).with_for_update())
    if not row:
        raise HTTPException(status_code=404, detail="Бодолтын ажилтан олдсонгүй.")
    if row.status != "flagged":
        raise HTTPException(status_code=409, detail="Мөр тэмдэглэгдээгүй байна.")
    inputs = dict(row.inputs or {})
    reason = inputs.pop("flag_reason", None)
    row.inputs = inputs
    row.status = "draft"
    row.warnings = [warning for warning in (row.warnings or []) if warning != "row_flagged"]
    db.add(MonthlyPayrollRowAudit(organization_id=actor.organization_id, run_id=run.id, row_id=row.id, account_id=actor.account_id, field_name="status", old_value={"status": "flagged", "flag_reason": reason}, new_value={"status": "draft"}, reason="Тэмдэглэгээ арилгав"))
    await db.commit()
    return _row_out(row)


@router.post("/runs/{run_id}/refresh-time")
async def refresh_run_time(run_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "calculate")
    run = await _run(db, actor, run_id, lock=True)
    if run.status != "draft":
        raise HTTPException(status_code=409, detail="Ноорог бодолтын цагийн мэдээллийг шинэчилнэ.")
    month = await _month(db, actor, run.month_id)
    rows = (await db.execute(select(MonthlyPayrollRunRow).where(
        MonthlyPayrollRunRow.run_id == run.id,
        MonthlyPayrollRunRow.organization_id == actor.organization_id,
    ).with_for_update())).scalars().all()
    changed = 0
    calendar_overrides = {date.fromisoformat(key): value for key, value in month.calendar_snapshot.items()}
    canonical_keys = TIME_INPUT_KEYS
    for row in rows:
        if row.status != "draft" or not row.profile_snapshot.get("complete", True):
            continue
        fresh = await _attendance_inputs(db, actor.organization_id, row.employee_id, month, run.cutoff_date if run.run_type == "advance" else None, amount(row.profile_snapshot.get("daily_norm_hours", 8)), calendar_overrides)
        inputs = dict(row.inputs or {})
        source = dict(inputs.get("_source_snapshot") or {})
        before = dict(inputs)
        new_source = {key: fresh.get(key) for key in canonical_keys}
        for key in canonical_keys:
            if inputs.get(key) == source.get(key):
                inputs[key] = fresh.get(key)
        inputs["_source_snapshot"] = new_source
        if inputs != before:
            row.inputs = inputs
            await _calculate_row(db, month, run, row)
            db.add(MonthlyPayrollRowAudit(organization_id=actor.organization_id, run_id=run.id, row_id=row.id, account_id=actor.account_id, field_name="time_inputs", old_value=source, new_value=new_source, reason="HR цаг, чөлөө, календарийн мэдээллийг шинэчлэв"))
            changed += 1
    await db.commit()
    return {"run_id": run.id, "updated_rows": changed, "rows": [_row_out(row) for row in rows]}


HR_IDENTITY_KEYS = ("name", "last_name", "first_name", "rd", "job_title", "department_id", "department", "pay_date")


@router.post("/runs/{run_id}/sync-workers")
async def sync_run_workers(run_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    """«Ажилчид шинэчлэх»: add newly due workers and flag HR changes.

    Existing rows are snapshots and are never overwritten here; a changed HR
    profile is parked on the row as «HR changed» until the accountant accepts.
    """
    await require_capability(db, actor, "payroll", "create")
    run = await _run(db, actor, run_id, lock=True)
    if run.status != "draft":
        raise HTTPException(status_code=409, detail="Зөвхөн ноорог бодолтын ажилтны жагсаалтыг шинэчилнэ.")
    month = await _month(db, actor, run.month_id, lock=True)
    if month.status != "open":
        raise HTTPException(status_code=409, detail="Хаагдсан сарын бодолтыг шинэчлэх боломжгүй.")
    existing = {row.employee_id: row for row in (await db.execute(select(MonthlyPayrollRunRow).where(
        MonthlyPayrollRunRow.run_id == run.id, MonthlyPayrollRunRow.organization_id == actor.organization_id,
    ).with_for_update())).scalars().all()}
    last = date(month.year, month.month, calendar.monthrange(month.year, month.month)[1])
    workers = await _profiles_for_month(db, actor, month, run.department_id)
    added: list[MonthlyPayrollRunRow] = []
    hr_changed = 0
    for employee, details, profile, history, segments in workers:
        row = existing.get(employee.id)
        if row is not None:
            identity, profile_data = await _fresh_snapshots(db, actor, month, run, employee, details, profile, history, segments)
            current_identity = {key: (row.identity_snapshot or {}).get(key) for key in HR_IDENTITY_KEYS}
            fresh_identity = {key: identity.get(key) for key in HR_IDENTITY_KEYS}
            inputs = dict(row.inputs or {})
            if fresh_identity != current_identity or profile_data != (row.profile_snapshot or {}):
                if inputs.get("_hr_pending") != {"identity": identity, "profile": profile_data}:
                    inputs["_hr_pending"] = {"identity": identity, "profile": profile_data}
                    row.inputs = inputs
                    row.warnings = list(dict.fromkeys([*(row.warnings or []), "hr_changed"]))
                    hr_changed += 1
            elif inputs.pop("_hr_pending", None) is not None:
                row.inputs = inputs
                row.warnings = [warning for warning in (row.warnings or []) if warning != "hr_changed"]
            continue
        if run.run_type == "advance" and (profile is None or profile.payment_frequency == "MONTHLY" or not any(min(int(day), last.day) == run.pay_date.day for day in (profile.pay_days or [])[:-1])):
            continue
        row = await _build_row(db, actor, month, run, employee, details, profile, history, segments)
        await db.flush()
        await _calculate_row_or_422(db, month, run, row)
        db.add(MonthlyPayrollRowAudit(organization_id=actor.organization_id, run_id=run.id, row_id=row.id, account_id=actor.account_id, field_name="worker_sync", old_value=None, new_value={"employee_id": employee.id, "pay_date": row.identity_snapshot["pay_date"]}, reason="Ажилтны жагсаалтыг HR бүртгэлээс шинэчлэв"))
        added.append(row)
    await db.commit()
    return {"run_id": run.id, "added_workers": len(added), "hr_changed_rows": hr_changed, "run": await get_run(run.id, db, actor)}


@router.post("/runs/{run_id}/rows/{employee_id}/accept-hr")
async def accept_hr_change(run_id: int, employee_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "create")
    run = await _run(db, actor, run_id, lock=True)
    if run.status != "draft":
        raise HTTPException(status_code=409, detail="Зөвхөн ноорог бодолтод HR өөрчлөлт хүлээн авна.")
    row = await db.scalar(select(MonthlyPayrollRunRow).where(MonthlyPayrollRunRow.run_id == run.id, MonthlyPayrollRunRow.employee_id == employee_id, MonthlyPayrollRunRow.organization_id == actor.organization_id).with_for_update())
    if row is None:
        raise HTTPException(status_code=404, detail="Бодолтын ажилтан олдсонгүй.")
    inputs = dict(row.inputs or {})
    pending = inputs.pop("_hr_pending", None)
    if not pending:
        raise HTTPException(status_code=409, detail="Хүлээн авах HR өөрчлөлт алга.")
    before = {"identity": row.identity_snapshot, "profile": row.profile_snapshot}
    row.identity_snapshot = pending["identity"]
    row.profile_snapshot = pending["profile"]
    row.inputs = inputs
    _reset_approval(row)
    month = await _month(db, actor, run.month_id)
    await _calculate_row_or_422(db, month, run, row)
    db.add(MonthlyPayrollRowAudit(organization_id=actor.organization_id, run_id=run.id, row_id=row.id, account_id=actor.account_id, field_name="hr_profile_accepted", old_value=before, new_value=pending, reason="HR-ийн өөрчлөлтийг хүлээн авав"))
    await db.commit()
    await db.refresh(row)
    return _row_out(row)


class AddWorkersInput(BaseModel):
    employee_ids: list[int] = Field(min_length=1, max_length=200)


@router.post("/runs/{run_id}/add-workers")
async def add_run_workers(run_id: int, data: AddWorkersInput, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    """«Ажилтан нэмэх»: a one-off advance for a worker not scheduled that day."""
    await require_capability(db, actor, "payroll", "create")
    run = await _run(db, actor, run_id, lock=True)
    if run.run_type != "advance" or run.status != "draft":
        raise HTTPException(status_code=409, detail="Зөвхөн ноорог урьдчилгааны бодолтод ажилтан нэмнэ.")
    month = await _month(db, actor, run.month_id, lock=True)
    if month.status != "open":
        raise HTTPException(status_code=409, detail="Хаагдсан сарын бодолтыг өөрчлөх боломжгүй.")
    existing_ids = set((await db.execute(select(MonthlyPayrollRunRow.employee_id).where(MonthlyPayrollRunRow.run_id == run.id))).scalars().all())
    requested = set(data.employee_ids) - existing_ids
    workers = [worker for worker in await _profiles_for_month(db, actor, month, None) if worker[0].id in requested and worker[2] is not None]
    missing = sorted(requested - {worker[0].id for worker in workers})
    if missing:
        raise HTTPException(status_code=422, detail={"code": "monthly_payroll_advance_worker_invalid", "employee_ids": missing, "message": "Цалингийн профайлтай, энэ сард ажилласан ажилтныг сонгоно уу."})
    added = []
    for employee, details, profile, history, segments in workers:
        row = await _build_row(db, actor, month, run, employee, details, profile, history, segments, one_off=True)
        await db.flush()
        await _calculate_row_or_422(db, month, run, row)
        db.add(MonthlyPayrollRowAudit(organization_id=actor.organization_id, run_id=run.id, row_id=row.id, account_id=actor.account_id, field_name="worker_added", old_value=None, new_value={"employee_id": employee.id, "one_off_advance": True}, reason="Нэг удаагийн урьдчилгаа нэмэв"))
        added.append(row)
    await db.commit()
    return {"run_id": run.id, "added_workers": len(added), "run": await get_run(run.id, db, actor)}


@router.get("/runs/{run_id}/import-template")
async def payroll_input_template(run_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "export")
    run = await _run(db, actor, run_id)
    rows = (await db.execute(select(MonthlyPayrollRunRow).where(
        MonthlyPayrollRunRow.run_id == run.id, MonthlyPayrollRunRow.organization_id == actor.organization_id,
    ))).scalars().all()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Оролтын засвар"
    sheet.append(["employee_id", "employee_name", "worked_normal_hours", "worked_to_date_hours", "overtime_weekday", "overtime_rest_day", "overtime_public_holiday", "leave_pay", "bonus", "deduction_type", "deduction_amount", "deduction_note", "reason"])
    sheet.append(["Ажилтны ID", "Ажилтан (зөвхөн харах)", "Ердийн цаг", "Урьдчилгааны цаг", "Ажлын өдрийн илүү цаг", "Амралтын өдрийн илүү цаг", "Баярын өдрийн илүү цаг", "Ээлжийн амралтын мөнгө", "Урамшуулал", "Суутгалын төрөл", "Суутгалын дүн", "Суутгалын тайлбар", "Засварын шалтгаан"])
    for cell in sheet[2]:
        cell.font = Font(bold=True)
    # One row per worker; blank cells keep the register's current value.
    for row in sorted(rows, key=lambda item: ((item.identity_snapshot or {}).get("department") or "", (item.identity_snapshot or {}).get("name") or "")):
        sheet.append([row.employee_id, (row.identity_snapshot or {}).get("name", "")])
    sheet.freeze_panes = "C3"
    for column in "ABCDEFGHIJKLM":
        sheet.column_dimensions[column].width = 18
    stream = io.BytesIO()
    workbook.save(stream)
    return Response(content=stream.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="monthly-payroll-{run_id}-input-template.xlsx"'})


@router.post("/runs/{run_id}/import-xlsx")
async def import_payroll_inputs(run_id: int, file: UploadFile = File(...), db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "create")
    run = await _run(db, actor, run_id, lock=True)
    if run.status != "draft":
        raise HTTPException(status_code=409, detail="Зөвхөн ноорог бодолтод Excel оролт оруулна.")
    payload = await file.read(5 * 1024 * 1024 + 1)
    if len(payload) > 5 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="Excel файл 5 MB-аас бага байх ёстой.")
    try:
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        headers = [str(value or "").strip() for value in values[0]] if values else []
        records = [dict(zip(headers, row)) for row in values[2:] if any(value not in (None, "") for value in row)]
    except Exception as exc:
        raise HTTPException(status_code=422, detail="Excel загварыг уншиж чадсангүй.") from exc
    required = {"employee_id", "reason"}
    if not required.issubset(headers) or len(headers) != len(set(headers)):
        raise HTTPException(status_code=422, detail="Ажилтны ID, шалтгаан баганатай цалингийн загвар ашиглана уу.")
    employee_ids = []
    for index, record in enumerate(records, start=3):
        try:
            employee_ids.append(int(record.get("employee_id")))
        except (TypeError, ValueError) as exc:
            raise HTTPException(status_code=422, detail=f"{index}-р мөрийн ажилтны ID буруу байна.") from exc
    if len(employee_ids) != len(set(employee_ids)):
        raise HTTPException(status_code=422, detail="Нэг ажилтныг файлд давхар оруулсан байна.")
    rows = (await db.execute(select(MonthlyPayrollRunRow).where(MonthlyPayrollRunRow.run_id == run.id, MonthlyPayrollRunRow.employee_id.in_(employee_ids or [-1]), MonthlyPayrollRunRow.organization_id == actor.organization_id).with_for_update())).scalars().all()
    by_employee = {row.employee_id: row for row in rows}
    if set(employee_ids) != set(by_employee):
        raise HTTPException(status_code=422, detail={"code": "monthly_payroll_import_employee_not_in_run", "employee_ids": sorted(set(employee_ids) - set(by_employee))})
    month = await _month(db, actor, run.month_id)
    changed = 0
    for record, employee_id in zip(records, employee_ids):
        row = by_employee[employee_id]
        def text_value(key: str) -> str | None:
            value = record.get(key)
            return None if value in (None, "") else str(value).strip()
        before = dict(row.inputs or {})
        # A blank cell keeps the row's current value; only filled cells change.
        provided: dict[str, Any] = {"reason": text_value("reason") or ""}
        hour_keys = ("worked_normal_hours", "leave_pay", "bonus") if run.run_type == "final" else ("worked_to_date_hours", "worked_normal_hours", "leave_pay", "bonus")
        for key in hour_keys:
            if text_value(key) is not None:
                provided[key] = text_value(key)
        overtime = {key: text_value(f"overtime_{key}") for key in ("weekday", "rest_day", "public_holiday")}
        if any(value is not None for value in overtime.values()):
            provided["overtime_hours"] = {**(before.get("overtime_hours") or {}), **{key: value for key, value in overtime.items() if value is not None}}
        try:
            if run.run_type == "final" and text_value("deduction_amount") is not None:
                deduction_amount = amount(text_value("deduction_amount"))
                provided["other_deductions"] = [{"type": text_value("deduction_type") or "Бусад", "amount": str(deduction_amount), "note": text_value("deduction_note") or ""}] if deduction_amount > 0 else []
            if set(provided) == {"reason"}:
                continue
            input_data = RowInput.model_validate(provided)
        except (ValueError, TypeError) as exc:
            raise HTTPException(status_code=422, detail={"code": "monthly_payroll_import_value_invalid", "employee_id": employee_id, "message": f"{row.identity_snapshot.get('name') or employee_id}: Excel-ийн утга буруу байна."}) from exc
        if row.status == "flagged":
            raise HTTPException(status_code=409, detail={"code": "monthly_payroll_import_row_locked", "employee_id": employee_id, "message": f"{row.identity_snapshot.get('name')}: шалгах тэмдэглэгээтэй мөрийг эхлээд цэвэрлэнэ үү."})
        if not input_data.reason:
            raise HTTPException(status_code=422, detail={"code": "monthly_payroll_import_reason_required", "employee_id": employee_id, "message": f"{row.identity_snapshot.get('name') or employee_id}: засварын шалтгаан бичнэ үү."})
        row.inputs = {**before, **_json_value(input_data.model_dump(exclude_unset=True, exclude={"reason"}))}
        _reset_approval(row)
        try:
            await _calculate_row(db, month, run, row)
        except ValueError as exc:
            raise HTTPException(status_code=422, detail={"code": "payroll_row_invalid", "employee_id": employee_id, "message": str(exc)}) from exc
        db.add(MonthlyPayrollRowAudit(organization_id=actor.organization_id, run_id=run.id, row_id=row.id, account_id=actor.account_id, field_name="excel_import", old_value=before, new_value=row.inputs, reason=input_data.reason))
        changed += 1
    await db.commit()
    return {"run_id": run.id, "updated_rows": changed}


@router.post("/runs/{run_id}/rows/{employee_id}/approve")
async def approve_row(run_id: int, employee_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "approve")
    run = await _run(db, actor, run_id, lock=True)
    if run.status != "draft":
        raise HTTPException(status_code=409, detail="Ноорог бодолтын мөрийг батална.")
    row = await db.scalar(select(MonthlyPayrollRunRow).where(
        MonthlyPayrollRunRow.run_id == run.id, MonthlyPayrollRunRow.employee_id == employee_id,
        MonthlyPayrollRunRow.organization_id == actor.organization_id,
    ).with_for_update())
    if row is None:
        raise HTTPException(status_code=404, detail="Бодолтын ажилтан олдсонгүй.")
    if row.status != "draft":
        raise HTTPException(status_code=409, detail="Тэмдэглэсэн эсвэл баталсан мөрийг батлах боломжгүй.")
    month = await _month(db, actor, run.month_id)
    try:
        await _calculate_row(db, month, run, row)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail={"code": "payroll_row_invalid", "message": str(exc)}) from exc
    result = row.result or {}
    blocking = sorted(set(row.warnings or []) & BLOCKING_ROW_WARNINGS)
    if blocking:
        raise HTTPException(status_code=422, detail={"code": "monthly_payroll_row_blocked", "issues": blocking})
    if run.run_type == "advance" and Decimal(result.get("advance", "0")) <= 0:
        raise HTTPException(status_code=422, detail={"code": "advance_positive", "message": "Урьдчилгаа 0-ээс их байх ёстой."})
    if run.run_type == "final" and Decimal(result.get("net_pay", "0")) < 0:
        raise HTTPException(status_code=422, detail={"code": "negative_final_pay", "message": "Сүүл цалин сөрөг байна."})
    row.status = "approved"
    row.approved_by_account_id = actor.account_id
    row.approved_at = datetime.now(timezone.utc)
    await db.commit()
    return _row_out(row)


@router.post("/runs/{run_id}/calculate")
async def calculate_run(run_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "calculate")
    run = await _run(db, actor, run_id, lock=True)
    if run.status != "draft":
        raise HTTPException(status_code=409, detail="Ноорог төлөвтэй бодолтыг дахин тооцно.")
    month = await _month(db, actor, run.month_id)
    rows = (await db.execute(select(MonthlyPayrollRunRow).where(MonthlyPayrollRunRow.run_id == run.id).with_for_update())).scalars().all()
    # «Дахин бодох» never touches approved rows (plan §6.4).
    for row in rows:
        if row.status != "approved":
            await _calculate_row_or_422(db, month, run, row)
    await db.commit()
    return await get_run(run.id, db, actor)


async def _company_name(db: AsyncSession, organization_id: int) -> str | None:
    settings = await db.scalar(select(MonthlyPayrollCompanySettings).where(MonthlyPayrollCompanySettings.organization_id == organization_id))
    if settings and settings.legal_company_name:
        return settings.legal_company_name
    return await db.scalar(select(Organization.name).where(Organization.id == organization_id))


@router.get("/runs/{run_id}/export")
async def export_run(run_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "export")
    run = await _run(db, actor, run_id)
    if run.status not in {"approved", "paid", "closed"}:
        raise HTTPException(status_code=409, detail="Баталсан бодолтын дараа Excel экспорт нээгдэнэ.")
    month = await _month(db, actor, run.month_id)
    rows = (await db.execute(select(MonthlyPayrollRunRow).where(
        MonthlyPayrollRunRow.run_id == run.id, MonthlyPayrollRunRow.organization_id == actor.organization_id,
    ))).scalars().all()
    export_rows = [{
        "identity": row.identity_snapshot or {}, "profile": row.profile_snapshot or {}, "inputs": row.inputs or {}, "result": row.result or {},
        "payout": json.loads(decrypt_secret(row.payout_snapshot_ciphertext)) if row.payout_snapshot_ciphertext else None,
    } for row in rows]
    stats = await _closing_stats(db, month, actor.organization_id) if run.run_type == "final" else None
    content = build_run_workbook(
        run_type=run.run_type, pay_date=run.pay_date, year=month.year, month=month.month, rows=export_rows,
        company=await _company_name(db, actor.organization_id), rule_snapshot=month.rule_snapshot or {}, stats=stats,
    )
    kind = "advance" if run.run_type == "advance" else "final"
    filename = f"monthly-payroll-{kind}-{run.pay_date:%Y-%m-%d}.xlsx"
    return Response(content=content, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@router.post("/runs/{run_id}/refresh-advances")
async def refresh_advances(run_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "calculate")
    run = await _run(db, actor, run_id, lock=True)
    if run.run_type != "final" or run.status != "draft":
        raise HTTPException(status_code=409, detail="Зөвхөн ноорог сүүл цалингийн бодолтод урьдчилгаа дахин татна.")
    month = await _month(db, actor, run.month_id)
    rows = (await db.execute(select(MonthlyPayrollRunRow).where(MonthlyPayrollRunRow.run_id == run.id).with_for_update())).scalars().all()
    run.advance_snapshot = await _approved_advance_snapshot(db, month.id, actor.organization_id)
    for row in rows:
        previous_advance = ((row.result or {}).get("advance"), (row.result or {}).get("advance_run_ids"))
        await _calculate_row_or_422(db, month, run, row)
        if (row.result.get("advance"), row.result.get("advance_run_ids")) == previous_advance:
            continue
        _reset_approval(row)
        db.add(MonthlyPayrollRowAudit(organization_id=actor.organization_id, run_id=run.id, row_id=row.id, account_id=actor.account_id, field_name="advances_refreshed", old_value=None, new_value=row.result.get("advance_lines"), reason="Урьдчилгаа дахин татав"))
    await db.commit()
    return await get_run(run.id, db, actor)


@router.post("/runs/{run_id}/approve")
async def approve_run(run_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    """«Бүгдийг батлах»: approve every row without a blocking issue.

    Flagged or blocked rows stay in review and are reported back; the run
    itself becomes Approved only once every row is approved.
    """
    await require_capability(db, actor, "payroll", "approve")
    run = await _run(db, actor, run_id, lock=True)
    if run.status != "draft":
        raise HTTPException(status_code=409, detail="Зөвхөн ноорог бодолтыг батална.")
    month = await _month(db, actor, run.month_id)
    rows = (await db.execute(select(MonthlyPayrollRunRow).where(MonthlyPayrollRunRow.run_id == run.id).with_for_update())).scalars().all()
    approved_now = 0
    skipped: list[dict[str, Any]] = []
    now = datetime.now(timezone.utc)
    for row in rows:
        if row.status == "approved":
            continue
        name = (row.identity_snapshot or {}).get("name")
        if row.status == "flagged":
            skipped.append({"employee_id": row.employee_id, "employee_name": name, "issues": ["row_flagged"]})
            continue
        try:
            await _calculate_row(db, month, run, row)
        except (ValueError, HTTPException) as exc:
            message = exc.detail.get("message") if isinstance(exc, HTTPException) and isinstance(exc.detail, dict) else str(exc)
            skipped.append({"employee_id": row.employee_id, "employee_name": name, "issues": ["invalid_input"], "message": message})
            continue
        blocking = sorted(set(row.warnings or []) & BLOCKING_ROW_WARNINGS)
        if blocking:
            skipped.append({"employee_id": row.employee_id, "employee_name": name, "issues": blocking})
            continue
        row.status = "approved"
        row.approved_by_account_id = actor.account_id
        row.approved_at = now
        approved_now += 1
    if rows and all(row.status == "approved" for row in rows):
        run.status = "approved"
        run.approved_by_account_id = actor.account_id
        run.approved_at = now
        if run.run_type == "advance":
            await _flag_final_advance_changes(db, run.month_id, actor.organization_id)
    await db.commit()
    return {**(await get_run(run.id, db, actor)), "approval_summary": {"approved": approved_now, "skipped": skipped}}


@router.post("/runs/{run_id}/rows/{employee_id}/unapprove")
async def unapprove_row(run_id: int, employee_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "approve")
    run = await _run(db, actor, run_id, lock=True)
    if run.status != "draft":
        raise HTTPException(status_code=409, detail="Бодолтын батлалтыг эхлээд цуцална уу.")
    row = await db.scalar(select(MonthlyPayrollRunRow).where(MonthlyPayrollRunRow.run_id == run.id, MonthlyPayrollRunRow.employee_id == employee_id, MonthlyPayrollRunRow.organization_id == actor.organization_id).with_for_update())
    if row is None or row.status != "approved":
        raise HTTPException(status_code=404 if row is None else 409, detail="Батлагдсан мөр олдсонгүй.")
    _reset_approval(row)
    db.add(MonthlyPayrollRowAudit(organization_id=actor.organization_id, run_id=run.id, row_id=row.id, account_id=actor.account_id, field_name="status", old_value={"status": "approved"}, new_value={"status": "draft"}, reason="Батлалт цуцлав"))
    await db.commit()
    return _row_out(row)


@router.post("/runs/{run_id}/unapprove")
async def unapprove_run(run_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    """«Батлалт цуцлах»: an approved, unpaid run returns to draft; rows keep their approval."""
    await require_capability(db, actor, "payroll", "approve")
    run = await _run(db, actor, run_id, lock=True)
    month = await _month(db, actor, run.month_id)
    if month.status != "open" or run.status != "approved":
        raise HTTPException(status_code=409, detail="Зөвхөн нээлттэй сарын баталсан, төлөөгүй бодолтын батлалтыг цуцална.")
    before = {"status": run.status, "approved_at": run.approved_at.isoformat() if run.approved_at else None}
    run.status = "draft"
    run.approved_by_account_id = None
    run.approved_at = None
    await record_change(db, actor=actor, topic="payroll", aggregate_type="monthly_payroll_run", aggregate_id=run.id, operation="unapproved", before=before, after={"status": run.status})
    if run.run_type == "advance":
        await _flag_final_advance_changes(db, run.month_id, actor.organization_id)
    await db.commit()
    return _run_out(run)


@router.delete("/runs/{run_id}")
async def delete_run(run_id: int, data: UnlockInput, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    """Delete a draft or approved (unpaid) run so it can be created again.

    Rows and their audit lines cascade. Paid runs must be reopened first; a
    final run that pulled a deleted advance is flagged «Урьдчилгаа өөрчлөгдсөн».
    """
    await require_capability(db, actor, "payroll", "create")
    run = await _run(db, actor, run_id, lock=True)
    if run.status not in {"draft", "approved"}:
        raise HTTPException(status_code=409, detail="Төлсөн эсвэл хаасан бодолтыг устгах боломжгүй. Эхлээд «Дахин нээх»-ээр ноорог болгоно уу.")
    if run.status == "approved":
        await require_capability(db, actor, "payroll", "approve")
    row_count = len((await db.execute(select(MonthlyPayrollRunRow.id).where(MonthlyPayrollRunRow.run_id == run.id))).scalars().all())
    before = {**_run_out(run), "rows": row_count}
    month_id, run_type = run.month_id, run.run_type
    await db.delete(run)
    await record_change(db, actor=actor, topic="payroll", aggregate_type="monthly_payroll_run", aggregate_id=run_id, operation="deleted", before=before, after={"reason": data.reason.strip()})
    if run_type == "advance":
        await _flag_final_advance_changes(db, month_id, actor.organization_id)
    await db.commit()
    return {"id": run_id, "deleted": True}


@router.post("/runs/{run_id}/paid")
async def mark_paid(run_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "pay")
    run = await _run(db, actor, run_id, lock=True)
    if run.status != "approved":
        raise HTTPException(status_code=409, detail="Зөвхөн баталсан бодолтыг төлсөн гэж тэмдэглэнэ.")
    run.status = "paid"
    run.paid_by_account_id = actor.account_id
    run.paid_at = datetime.now(timezone.utc)
    await db.commit()
    return _run_out(run)


@router.post("/runs/{run_id}/unpaid")
async def mark_unpaid(run_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "pay")
    run = await _run(db, actor, run_id, lock=True)
    month = await _month(db, actor, run.month_id)
    if month.status != "open" or run.status != "paid":
        raise HTTPException(status_code=409, detail="Зөвхөн нээлттэй сарын төлсөн бодолтыг төлөөгүй болгоно.")
    before = {"status": run.status, "paid_at": run.paid_at.isoformat() if run.paid_at else None}
    run.status = "approved"
    run.paid_by_account_id = None
    run.paid_at = None
    await record_change(db, actor=actor, topic="payroll", aggregate_type="monthly_payroll_run", aggregate_id=run.id, operation="marked_unpaid", before=before, after={"status": run.status})
    await db.commit()
    return _run_out(run)


class CloseInput(BaseModel):
    waivers: dict[int, str] = Field(default_factory=dict)


async def _validated_waivers(runs: list[MonthlyPayrollRun], waivers: dict[int, str]) -> dict[int, str]:
    by_id = {run.id: run for run in runs}
    clean: dict[int, str] = {}
    for run_id, reason in waivers.items():
        run = by_id.get(int(run_id))
        if run is None or run.run_type != "advance" or run.status in APPROVED_RUN_STATUSES:
            raise HTTPException(status_code=422, detail={"code": "monthly_payroll_waiver_invalid", "message": "Зөвхөн батлагдаагүй урьдчилгааны бодолтыг чөлөөлнө."})
        if not str(reason or "").strip():
            raise HTTPException(status_code=422, detail={"code": "monthly_payroll_waiver_reason_required", "message": "Чөлөөлөх шалтгаан бичнэ үү."})
        clean[run.id] = str(reason).strip()
    return clean


@router.post("/months/{month_id}/close")
async def close_month(month_id: int, data: CloseInput | None = None, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "approve")
    await require_capability(db, actor, "payroll", "export")
    month = await _month(db, actor, month_id, lock=True)
    if month.status != "open":
        raise HTTPException(status_code=409, detail={"code": "monthly_payroll_month_closed", "message": "Сар аль хэдийн хаагдсан байна."})
    runs = (await db.execute(select(MonthlyPayrollRun).where(MonthlyPayrollRun.month_id == month.id).order_by(MonthlyPayrollRun.pay_date, MonthlyPayrollRun.run_type).with_for_update())).scalars().all()
    waivers = await _validated_waivers(runs, (data or CloseInput()).waivers)
    issues = await _month_close_issues(db, actor, month, runs, waived_run_ids=set(waivers))
    if issues:
        raise HTTPException(status_code=409, detail={"code": "monthly_payroll_close_blocked", "issues": issues, "message": "Хаалтын шалгалтын алдааг засна уу."})
    closing_stats = await _closing_stats(db, month, actor.organization_id, waived_run_ids=set(waivers))
    closing_stats["waivers"] = {str(run_id): reason for run_id, reason in waivers.items()}
    closing_stats["accounting_summary"] = closing_stats["accounting"]
    snapshots = []
    stored_exports: dict[str, str] = {}
    for run in runs:
        if run.id in waivers:
            run.note = "\n".join(filter(None, [run.note, f"Сар хаахад чөлөөлсөн: {waivers[run.id]}"]))
        # A waived advance run was never approved: it stays draft (never counted
        # as paid) and the closed month blocks any further edits to it.
        if run.id not in waivers:
            run.status = "closed"
        rows = (await db.execute(select(MonthlyPayrollRunRow).where(MonthlyPayrollRunRow.run_id == run.id).order_by(MonthlyPayrollRunRow.id))).scalars().all()
        snapshots.append({"run": {**_run_out(run), "waived": run.id in waivers}, "rows": [_row_out(row, include_payout_snapshot=True) for row in rows]})
        if run.id not in waivers:
            exported = await export_run(run.id, db, actor)
            stored_exports[str(run.id)] = base64.b64encode(exported.body).decode("ascii")
    audits = (await db.execute(select(MonthlyPayrollRowAudit).where(MonthlyPayrollRowAudit.organization_id == actor.organization_id, MonthlyPayrollRowAudit.run_id.in_([run.id for run in runs] or [-1])).order_by(MonthlyPayrollRowAudit.id))).scalars().all()
    previous = await db.scalar(select(MonthlyPayrollArchive.version).where(MonthlyPayrollArchive.month_id == month.id).order_by(MonthlyPayrollArchive.version.desc()).limit(1)) or 0
    closed_at = datetime.now(timezone.utc)
    closing_stats["meta"] = {**closing_stats.get("meta", {}), "closed_by_account_id": actor.account_id, "closed_at": closed_at.isoformat(), "archive_version": previous + 1}
    archive = MonthlyPayrollArchive(
        organization_id=actor.organization_id, month_id=month.id, version=previous + 1,
        closed_by_account_id=actor.account_id,
        snapshot={"month": {"year": month.year, "month": month.month, "rule_snapshot": month.rule_snapshot, "calendar_snapshot": month.calendar_snapshot}, "runs": snapshots, "closing_stats": closing_stats, "accounting_summary": closing_stats["accounting"], "waivers": closing_stats["waivers"], "row_audit": [{"row_id": row.row_id, "field": row.field_name, "old": row.old_value, "new": row.new_value, "reason": row.reason, "account_id": row.account_id, "at": row.created_at.isoformat()} for row in audits], "exports_base64": stored_exports},
    )
    db.add(archive)
    month.status = "closed"
    month.closed_by_account_id = actor.account_id
    month.closed_at = closed_at
    await db.commit()
    return {"month_id": month.id, "archive_version": archive.version, "closed_at": month.closed_at.isoformat(), "closing_stats": closing_stats}


@router.get("/months/{month_id}/closing-stats")
async def get_closing_stats(month_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    """Closing review: statistics plus the checks «Сар хаах» will run."""
    await require_capability(db, actor, "payroll", "view_salary")
    month = await _month(db, actor, month_id)
    runs = (await db.execute(select(MonthlyPayrollRun).where(MonthlyPayrollRun.month_id == month.id, MonthlyPayrollRun.organization_id == actor.organization_id).order_by(MonthlyPayrollRun.pay_date))).scalars().all()
    unapproved_advances = [run for run in runs if run.run_type == "advance" and run.status not in APPROVED_RUN_STATUSES]
    stats = await _closing_stats(db, month, actor.organization_id)
    if month.status == "open":
        stats["close_issues"] = await _month_close_issues(db, actor, month, runs)
        stats["close_issues_with_waivers"] = await _month_close_issues(db, actor, month, runs, waived_run_ids={run.id for run in unapproved_advances})
    stats["unapproved_advance_runs"] = [{"run_id": run.id, "pay_date": run.pay_date.isoformat(), "status": run.status} for run in unapproved_advances]
    return stats


@router.post("/months/{month_id}/unlock")
async def unlock_month(month_id: int, data: UnlockInput, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "administer")
    month = await _month(db, actor, month_id, lock=True)
    if month.status != "closed":
        raise HTTPException(status_code=409, detail="Зөвхөн хаагдсан сарыг нээнэ.")
    runs = (await db.execute(select(MonthlyPayrollRun).where(MonthlyPayrollRun.month_id == month.id).with_for_update())).scalars().all()
    before = {"status": month.status, "run_statuses": {str(run.id): run.status for run in runs}}
    month.status = "open"
    # Runs return to their pre-close state; money already paid stays locked
    # until an administrator reopens that specific run with a reason.
    for run in runs:
        if run.status == "closed":
            run.status = "paid" if run.paid_at else "approved"
    await record_change(db, actor=actor, topic="payroll", aggregate_type="monthly_payroll_month", aggregate_id=month.id, operation="unlocked", before=before, after={"status": month.status, "reason": data.reason.strip()})
    await db.commit()
    return {"month_id": month.id, "status": month.status, "reason": data.reason.strip()}


@router.post("/runs/{run_id}/reopen")
async def reopen_run(run_id: int, data: UnlockInput, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "administer")
    run = await _run(db, actor, run_id, lock=True)
    month = await _month(db, actor, run.month_id)
    if month.status != "open":
        raise HTTPException(status_code=409, detail={"code": "monthly_payroll_month_closed", "message": "Эхлээд сарыг шалтгаантайгаар нээнэ үү."})
    if run.status not in {"approved", "paid"}:
        raise HTTPException(status_code=409, detail={"code": "monthly_payroll_run_not_locked", "message": "Зөвхөн баталсан эсвэл төлсөн бодолтыг нээнэ."})
    before = {"status": run.status, "approved_at": run.approved_at.isoformat() if run.approved_at else None, "paid_at": run.paid_at.isoformat() if run.paid_at else None}
    run.status = "draft"
    run.approved_by_account_id = None
    run.approved_at = None
    run.paid_by_account_id = None
    run.paid_at = None
    rows = (await db.execute(select(MonthlyPayrollRunRow).where(MonthlyPayrollRunRow.run_id == run.id, MonthlyPayrollRunRow.organization_id == actor.organization_id).with_for_update())).scalars().all()
    for row in rows:
        if row.status == "approved":
            row.status = "draft"
            row.approved_by_account_id = None
            row.approved_at = None
    await record_change(db, actor=actor, topic="payroll", aggregate_type="monthly_payroll_run", aggregate_id=run.id, operation="reopened", before=before, after={"status": run.status, "reason": data.reason.strip()})
    if run.run_type == "advance":
        await _flag_final_advance_changes(db, run.month_id, actor.organization_id)
    await db.commit()
    return _run_out(run)


@router.get("/months/{month_id}/archives")
async def list_archives(month_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "view_salary")
    await _month(db, actor, month_id)
    rows = (await db.execute(select(MonthlyPayrollArchive).where(
        MonthlyPayrollArchive.month_id == month_id, MonthlyPayrollArchive.organization_id == actor.organization_id,
    ).order_by(MonthlyPayrollArchive.version.desc()))).scalars().all()
    def public_snapshot(snapshot: dict[str, Any]) -> dict[str, Any]:
        return {key: value for key, value in snapshot.items() if key != "exports_base64"}
    return [{"id": row.id, "version": row.version, "closed_at": row.closed_at.isoformat(), "snapshot": public_snapshot(row.snapshot)} for row in rows]


@router.get("/archives")
async def list_closed_month_archives(db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    """«Цалин → Архив»: latest archive version of every closed month."""
    await require_capability(db, actor, "payroll", "view_salary")
    months = (await db.execute(select(MonthlyPayrollMonth).where(
        MonthlyPayrollMonth.organization_id == actor.organization_id, MonthlyPayrollMonth.status == "closed",
    ).order_by(MonthlyPayrollMonth.year.desc(), MonthlyPayrollMonth.month.desc()))).scalars().all()
    output = []
    for month in months:
        archive = await db.scalar(select(MonthlyPayrollArchive).where(MonthlyPayrollArchive.month_id == month.id, MonthlyPayrollArchive.organization_id == actor.organization_id).order_by(MonthlyPayrollArchive.version.desc()).limit(1))
        if not archive:
            continue
        snapshot = archive.snapshot or {}
        stats = snapshot.get("closing_stats") or {}
        runs = [item.get("run") or {} for item in snapshot.get("runs") or []]
        output.append({
            "month_id": month.id, "month": f"{month.year:04d}-{month.month:02d}", "archive_id": archive.id, "version": archive.version,
            "closed_at": archive.closed_at.isoformat(), "headcount": (stats.get("headcount") or {}).get("on_register", 0),
            "gross": (stats.get("totals") or {}).get("gross", "0"), "company_cost": (stats.get("totals") or {}).get("company_cost", "0"),
            "runs": len(runs), "runs_paid": sum(bool(run.get("paid_at")) for run in runs),
        })
    return output


@router.get("/archives/{archive_id}/runs/{run_id}/export")
async def download_archived_run_export(archive_id: int, run_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "export")
    archive = await db.scalar(select(MonthlyPayrollArchive).where(
        MonthlyPayrollArchive.id == archive_id,
        MonthlyPayrollArchive.organization_id == actor.organization_id,
    ))
    if not archive:
        raise HTTPException(status_code=404, detail="Payroll archive not found")
    encoded = (archive.snapshot or {}).get("exports_base64", {}).get(str(run_id))
    if not encoded:
        raise HTTPException(status_code=404, detail="Archived export not found")
    payload = base64.b64decode(encoded)
    return Response(content=payload, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="monthly-payroll-archive-{archive.version}-{run_id}.xlsx"'})


@router.get("/dashboard")
async def payroll_dashboard(month: str | None = None, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    """«Цалин → Хянах самбар» (plan §9) for one month."""
    await require_capability(db, actor, "payroll", "view_salary")
    try:
        year, month_num = (int(part) for part in (month or date.today().strftime("%Y-%m")).split("-"))
        first = date(year, month_num, 1)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail="Сарыг YYYY-MM форматаар оруулна уу.") from exc
    last = date(year, month_num, calendar.monthrange(year, month_num)[1])
    month_row = await db.scalar(select(MonthlyPayrollMonth).where(MonthlyPayrollMonth.organization_id == actor.organization_id, MonthlyPayrollMonth.year == year, MonthlyPayrollMonth.month == month_num))
    pipeline: list[dict[str, Any]] = []
    alerts = {"blocking_rows": 0, "warning_rows": 0, "incomplete_profiles": 0, "hr_changed": 0, "advance_changed": 0, "advance_not_calculated": 0, "flagged": 0}
    stats: dict[str, Any] | None = None
    runs: list[MonthlyPayrollRun] = []
    if month_row:
        runs = (await db.execute(select(MonthlyPayrollRun).where(MonthlyPayrollRun.month_id == month_row.id, MonthlyPayrollRun.organization_id == actor.organization_id).order_by(MonthlyPayrollRun.run_type, MonthlyPayrollRun.pay_date))).scalars().all()
        runs = sorted(runs, key=lambda run: (run.run_type == "final", run.pay_date))
        for run in runs:
            rows = (await db.execute(select(MonthlyPayrollRunRow).where(MonthlyPayrollRunRow.run_id == run.id))).scalars().all()
            amount_key = "advance" if run.run_type == "advance" else "net_pay"
            for row in rows:
                warnings = set(row.warnings or [])
                alerts["blocking_rows"] += bool(warnings & BLOCKING_ROW_WARNINGS)
                alerts["warning_rows"] += bool(warnings - BLOCKING_ROW_WARNINGS)
                alerts["incomplete_profiles"] += bool(warnings & {"profile_missing", "salary_history_missing_or_incomplete", "profile_incomplete"})
                for key in ("hr_changed", "advance_changed", "advance_not_calculated"):
                    alerts[key] += key in warnings
                alerts["flagged"] += row.status == "flagged"
            pipeline.append({
                **_run_out(run), "workers": len(rows), "approved_rows": sum(row.status == "approved" for row in rows),
                "total": sum((_money((row.result or {}).get(amount_key)) for row in rows), Decimal("0")),
                "gross": sum((_money((row.result or {}).get("gross")) for row in rows), Decimal("0")) if run.run_type == "final" else None,
            })
        stats = await _closing_stats(db, month_row, actor.organization_id)
    # Upcoming payments: unpaid runs plus HR pay days that have no run yet.
    month_open = bool(month_row and month_row.status == "open")
    upcoming = [{"pay_date": item["pay_date"], "run_type": item["run_type"], "run_id": item["id"], "workers": item["workers"], "amount": item["total"], "status": item["status"]} for item in pipeline if month_open and item["status"] in {"draft", "approved"}]
    probe = month_row or MonthlyPayrollMonth(organization_id=actor.organization_id, year=year, month=month_num, calendar_snapshot={}, rule_snapshot={})
    workers = await _profiles_for_month(db, actor, probe, None)
    planned: dict[tuple[str, str], int] = {}
    for _employee, _details, profile, history, _segments in workers:
        if profile is None or history is None:
            if not month_row:
                alerts["incomplete_profiles"] += 1
            continue
        days = [min(int(day), last.day) for day in (profile.pay_days or [])]
        for index, day in enumerate(days):
            kind = "final" if index == len(days) - 1 else "advance"
            planned[(date(year, month_num, day).isoformat(), kind)] = planned.get((date(year, month_num, day).isoformat(), kind), 0) + 1
    existing = {(run.pay_date.isoformat(), run.run_type) for run in runs}
    has_final = any(run.run_type == "final" for run in runs)
    for (pay_date, kind), count in sorted(planned.items()):
        if (month_row and not month_open) or (kind == "final" and has_final) or (pay_date, kind) in existing:
            continue
        upcoming.append({"pay_date": pay_date, "run_type": kind, "run_id": None, "workers": count, "amount": None, "status": "not_created"})
    upcoming.sort(key=lambda item: str(item["pay_date"]))
    # Trend: closed months read their archive; open months are live.
    history_months = (await db.execute(select(MonthlyPayrollMonth).where(
        MonthlyPayrollMonth.organization_id == actor.organization_id,
        (MonthlyPayrollMonth.year * 100 + MonthlyPayrollMonth.month) <= year * 100 + month_num,
    ).order_by(MonthlyPayrollMonth.year.desc(), MonthlyPayrollMonth.month.desc()).limit(12))).scalars().all()
    trend = []
    for item in reversed(history_months):
        if item.status == "closed":
            archive = await db.scalar(select(MonthlyPayrollArchive).where(MonthlyPayrollArchive.month_id == item.id).order_by(MonthlyPayrollArchive.version.desc()).limit(1))
            item_stats = (archive.snapshot or {}).get("closing_stats", {}) if archive else {}
        else:
            item_stats = stats if month_row and item.id == month_row.id else await _closing_stats(db, item, actor.organization_id)
        trend.append({"month": f"{item.year:04d}-{item.month:02d}", "status": item.status, "gross": (item_stats.get("totals") or {}).get("gross", "0"), "company_cost": (item_stats.get("totals") or {}).get("company_cost", "0"), "headcount": (item_stats.get("headcount") or {}).get("on_register", 0)})
    advance_paid = sum((item["total"] for item in pipeline if item["run_type"] == "advance" and item["status"] in {"paid", "closed"}), Decimal("0"))
    advance_planned = sum((item["total"] for item in pipeline if item["run_type"] == "advance"), Decimal("0"))
    return _json_value({
        "month": f"{year:04d}-{month_num:02d}", "month_id": month_row.id if month_row else None, "status": month_row.status if month_row else None,
        "pipeline": pipeline, "stats": stats, "advance": {"paid": advance_paid, "planned": advance_planned},
        "has_final": any(run.run_type == "final" for run in runs), "upcoming": upcoming, "alerts": alerts, "trend": trend,
    })


REPORT_COLUMN_LABELS = {
    "month": "Сар", "employee_id": "Ажилтны ID", "employee_name": "Ажилтан", "department": "Хэлтэс", "run_type": "Бодолт", "pay_date": "Төлбөрийн өдөр",
    "gross": "Олговол зохих", "taxable_income": "Татвар ногдох орлого", "employee_shi": "Ажилтны НДШ", "employer_shi": "БНДШ", "pit": "ХХОАТ", "relief": "ХХОАТ ХӨН",
    "advance": "Урьдчилгаа", "other_deductions": "Бусад суутгал", "net_pay": "Сүүл цалин", "bucket": "Ангилал", "hours": "Цаг", "amount": "Дүн",
    "company_cost": "Нийт зардал", "total": "Дүн", "type": "Төрөл", "note": "Тайлбар",
    "ytd_gross": "Оны эхнээс · олговол зохих", "ytd_employee_shi": "Оны эхнээс · НДШ", "ytd_employer_shi": "Оны эхнээс · БНДШ", "ytd_pit": "Оны эхнээс · ХХОАТ",
}


async def _report_data(db: AsyncSession, actor: ActorContext, report_kind: str, from_month: str, to_month: str, department_id: int | None) -> dict[str, Any]:
    kinds = {"salary-register", "tax-shi", "overtime", "department-cost", "advance-final", "other-deductions"}
    if report_kind not in kinds:
        raise HTTPException(status_code=404, detail="Payroll report not found")
    try:
        start, end = date.fromisoformat(f"{from_month}-01"), date.fromisoformat(f"{to_month}-01")
    except ValueError as exc:
        raise HTTPException(status_code=422, detail="Сарын мужийг YYYY-MM форматаар оруулна уу.") from exc
    if start > end:
        raise HTTPException(status_code=422, detail="Эхлэх сар төгсөх сараас хойш байж болохгүй.")
    query_start = date(start.year, 1, 1) if report_kind == "tax-shi" else start
    month_key = MonthlyPayrollMonth.year * 100 + MonthlyPayrollMonth.month
    months = (await db.execute(select(MonthlyPayrollMonth).where(
        MonthlyPayrollMonth.organization_id == actor.organization_id,
        month_key >= query_start.year * 100 + query_start.month,
        month_key <= end.year * 100 + end.month,
    ).order_by(MonthlyPayrollMonth.year, MonthlyPayrollMonth.month))).scalars().all()
    source_runs: list[dict[str, Any]] = []
    for month in months:
        if month.status == "closed":
            archive = await db.scalar(select(MonthlyPayrollArchive).where(MonthlyPayrollArchive.month_id == month.id).order_by(MonthlyPayrollArchive.version.desc()).limit(1))
            if archive:
                source_runs.extend({**item["run"], "rows": item["rows"], "month": f"{month.year:04d}-{month.month:02d}"} for item in archive.snapshot.get("runs", []))
                continue
        runs = (await db.execute(select(MonthlyPayrollRun).where(MonthlyPayrollRun.month_id == month.id, MonthlyPayrollRun.organization_id == actor.organization_id))).scalars().all()
        for run in runs:
            rows = (await db.execute(select(MonthlyPayrollRunRow).where(MonthlyPayrollRunRow.run_id == run.id, MonthlyPayrollRunRow.organization_id == actor.organization_id))).scalars().all()
            source_runs.append({**_run_out(run), "rows": [_row_out(row) for row in rows], "month": f"{month.year:04d}-{month.month:02d}"})
    output: list[dict[str, Any]] = []
    for run in source_runs:
        if run.get("waived"):
            continue
        if report_kind == "advance-final" and run["run_type"] == "advance":
            output.append({"month": run["month"], "run_type": run["run_type"], "pay_date": run["pay_date"], "total": sum((Decimal(str(row.get("result", {}).get("advance", 0))) for row in run["rows"]), Decimal("0"))})
            continue
        if run["run_type"] != "final" and report_kind != "advance-final":
            continue
        for row in run["rows"]:
            identity, result, inputs = row.get("identity", {}), row.get("result", {}), row.get("inputs", {})
            if department_id and identity.get("department_id") != department_id:
                continue
            base = {"month": run["month"], "employee_id": row.get("employee_id"), "employee_name": identity.get("name", ""), "department": identity.get("department", "Бусад"), "run_type": run["run_type"], "pay_date": run["pay_date"]}
            if report_kind == "salary-register":
                output.append({**base, **{key: result.get(key, "0") for key in ("gross", "taxable_income", "employee_shi", "employer_shi", "pit", "relief", "advance", "other_deductions", "net_pay")}})
            elif report_kind == "tax-shi":
                output.append({**base, **{key: result.get(key, "0") for key in ("gross", "taxable_income", "employee_shi", "employer_shi", "pit", "relief", "net_pay")}})
            elif report_kind == "overtime":
                for bucket, hours in (inputs.get("overtime_hours") or {}).items():
                    output.append({**base, "bucket": bucket, "hours": hours, "amount": result.get("overtime_by_bucket", {}).get(bucket, "0")})
            elif report_kind == "department-cost":
                output.append({**base, "gross": result.get("gross", "0"), "employee_shi": result.get("employee_shi", "0"), "pit": result.get("pit", "0"), "employer_shi": result.get("employer_shi", "0"), "company_cost": str(Decimal(str(result.get("gross", 0))) + Decimal(str(result.get("employer_shi", 0))))})
            elif report_kind == "advance-final":
                output.append({**base, "run_type": "final", "total": result.get("net_pay", "0")})
            elif report_kind == "other-deductions":
                for line in inputs.get("other_deductions", []):
                    output.append({**base, "type": line.get("type", ""), "amount": line.get("amount", "0"), "note": line.get("note", "")})
    if report_kind == "tax-shi":
        ytd: dict[tuple[int, int], dict[str, Decimal]] = {}
        for row in sorted(output, key=lambda item: (item["month"], item.get("employee_id", 0))):
            key = (int(row.get("employee_id") or 0), int(row["month"][:4]))
            totals = ytd.setdefault(key, {field: Decimal("0") for field in ("gross", "employee_shi", "employer_shi", "pit")})
            for field in totals:
                totals[field] += Decimal(str(row.get(field, 0)))
                row[f"ytd_{field}"] = totals[field]
        output = [row for row in output if row["month"] >= from_month]
    return {"report_kind": report_kind, "from_month": from_month, "to_month": to_month, "rows": _json_value(output)}


@router.get("/reports/{report_kind}")
async def payroll_monthly_report(report_kind: str, from_month: str, to_month: str, department_id: int | None = None, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "view_salary")
    return await _report_data(db, actor, report_kind, from_month, to_month, department_id)


@router.get("/reports/{report_kind}/export")
async def export_monthly_report(report_kind: str, from_month: str, to_month: str, department_id: int | None = None, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "export")
    report = await _report_data(db, actor, report_kind, from_month, to_month, department_id)
    rows = report["rows"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Тайлан"
    titles = {"salary-register": "Цалингийн бүртгэл", "tax-shi": "Татвар, НДШ (оны эхнээс)", "overtime": "Илүү цаг", "department-cost": "Хэлтсийн зардал", "advance-final": "Урьдчилгаа ба сүүл цалин", "other-deductions": "Бусад суутгал"}
    sheet.append([f"{titles.get(report_kind, report_kind)} · {from_month} — {to_month}"])
    sheet.cell(1, 1).font = Font(bold=True, size=13)
    keys = list(rows[0]) if rows else ["month", "employee_name"]
    sheet.append([REPORT_COLUMN_LABELS.get(key, key) for key in keys])
    for cell in sheet[2]:
        cell.font = Font(bold=True)
    run_labels = {"advance": "Урьдчилгаа", "final": "Сүүл цалин"}
    for row in rows:
        values = []
        for key in keys:
            value = row.get(key, "")
            if isinstance(value, (dict, list)):
                value = json.dumps(value, ensure_ascii=False)
            elif key == "run_type":
                value = run_labels.get(value, value)
            elif key not in {"month", "pay_date", "employee_id", "employee_name", "department", "bucket", "type", "note"} and isinstance(value, str):
                try:
                    number = Decimal(value)
                    value = int(number) if number == number.to_integral_value() else float(number)
                except ArithmeticError:
                    pass
            values.append(value)
        sheet.append(values)
    for row_cells in sheet.iter_rows(min_row=3):
        for cell in row_cells:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"
    sheet.freeze_panes = "A3"
    output = io.BytesIO()
    workbook.save(output)
    filename = f"payroll-{report_kind}-{from_month}-{to_month}.xlsx"
    return Response(content=output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{filename}"'})


async def _worker_history(db: AsyncSession, actor: ActorContext, employee_id: int) -> list[dict[str, Any]]:
    employee = await db.scalar(select(Employee.id).where(Employee.id == employee_id, Employee.organization_id == actor.organization_id))
    if not employee:
        raise HTTPException(status_code=404, detail="Ажилтан олдсонгүй.")
    months = (await db.execute(select(MonthlyPayrollMonth).where(MonthlyPayrollMonth.organization_id == actor.organization_id, MonthlyPayrollMonth.status == "closed").order_by(MonthlyPayrollMonth.year, MonthlyPayrollMonth.month))).scalars().all()
    history = []
    for month in months:
        archive = await db.scalar(select(MonthlyPayrollArchive).where(MonthlyPayrollArchive.month_id == month.id).order_by(MonthlyPayrollArchive.version.desc()).limit(1))
        if not archive:
            continue
        for archived_run in archive.snapshot.get("runs", []):
            for row in archived_run.get("rows", []):
                if row.get("employee_id") == employee_id:
                    history.append({"month": f"{month.year:04d}-{month.month:02d}", "run_type": archived_run["run"]["run_type"], "pay_date": archived_run["run"]["pay_date"], **row.get("identity", {}), **row.get("result", {})})
    return history


@router.get("/employees/{employee_id}/history")
async def payroll_worker_history(employee_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "view_salary")
    return await _worker_history(db, actor, employee_id)


@router.get("/employees/{employee_id}/history/export")
async def export_payroll_worker_history(employee_id: int, db: AsyncSession = Depends(get_db), actor: ActorContext = Depends(get_actor)):
    await require_capability(db, actor, "payroll", "export")
    rows = await _worker_history(db, actor, employee_id)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ажилтны түүх"
    keys = ["month", "run_type", "pay_date", "name", "department", "gross", "employee_shi", "pit", "advance", "other_deductions", "net_pay"]
    sheet.append(["Архивлагдсан цалингийн түүх"])
    sheet.append(keys)
    for cell in sheet[2]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([row.get(key, "") for key in keys])
    output = io.BytesIO()
    workbook.save(output)
    filename = f"payroll-worker-{employee_id}-history.xlsx"
    return Response(content=output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{filename}"'})
