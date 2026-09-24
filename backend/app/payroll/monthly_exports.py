"""Excel workbooks for monthly payroll runs (plan §12).

Pure functions over plain row dictionaries: the caller resolves rows,
decrypted payout details, company name and closing statistics. Values are
exported as approved, including manual edits; totals are SUBTOTAL formulas
so department subtotals and the grand total reconcile inside Excel.
"""

from __future__ import annotations

import io
from collections import OrderedDict
from datetime import date
from decimal import Decimal
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


MONEY = "#,##0"
HOURS = "#,##0.0#"
THIN = Side(style="thin", color="A6A6A6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="E7E6E6")
GROUP_FILL = PatternFill("solid", fgColor="F2F2F2")
TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")
BUCKET_LABELS = {"weekday": "Ажлын өдрийн илүү цаг", "rest_day": "Амралтын өдөр", "public_holiday": "Нийтийн амралтын өдөр"}
DAY_TYPE_LABELS = {"working": "Ажлын өдөр", "weekly_rest": "Долоо хоногийн амралт", "public_holiday": "Нийтийн амралт"}
WEEKDAYS = ("Даваа", "Мягмар", "Лхагва", "Пүрэв", "Баасан", "Бямба", "Ням")
FUND_LABELS = {"pension": "Тэтгэврийн даатгал", "benefit": "Тэтгэмжийн даатгал", "unemployment": "Ажилгүйдлийн даатгал", "health": "Эрүүл мэндийн даатгал", "injury": "ҮОМШӨ"}
BASIS_LABELS = {"FIXED": "Тогтмол дүн", "PERCENT": "Үндсэн цалингийн хувь", "WORKED-TO-DATE": "Ажилласан цагаар"}
SALARY_TYPE_LABELS = {"PRORATION": "Цагаар", "FIXED": "Тогтмол"}


def _d(value: Any) -> Decimal:
    try:
        return Decimal(str(value if value not in (None, "") else 0))
    except Exception:  # noqa: BLE001 - malformed snapshot values export as zero
        return Decimal("0")


def _num(value: Any) -> int | float:
    number = _d(value)
    return int(number) if number == number.to_integral_value() else float(number)


def company_title(company: str | None, text: str) -> str:
    name = (company or "").strip() or "Байгууллага"
    prefix = f"{name}-ийн" if "ХХК" in name.upper() else f'"{name}" ХХК-ийн'
    return f"{prefix} {text}"


def split_name(identity: Mapping[str, Any]) -> tuple[str, str]:
    last, first = (identity.get("last_name") or "").strip(), (identity.get("first_name") or "").strip()
    if last or first:
        return last, first or (identity.get("name") or "")
    parts = str(identity.get("name") or "").split(maxsplit=1)
    return (parts[0], parts[1]) if len(parts) == 2 else ("", parts[0] if parts else "")


def _by_department(rows: Sequence[Mapping[str, Any]]) -> "OrderedDict[str, list[Mapping[str, Any]]]":
    groups: OrderedDict[str, list[Mapping[str, Any]]] = OrderedDict()
    for row in sorted(rows, key=lambda item: ((item["identity"].get("department") or "Бусад"), item["identity"].get("name") or "")):
        groups.setdefault(row["identity"].get("department") or "Бусад", []).append(row)
    return groups


def _style_range(sheet: Worksheet, row: int, first_col: int, last_col: int, *, bold: bool = False, fill: PatternFill | None = None) -> None:
    for column in range(first_col, last_col + 1):
        cell = sheet.cell(row, column)
        cell.border = BORDER
        if bold:
            cell.font = Font(bold=True)
        if fill:
            cell.fill = fill


def _title(sheet: Worksheet, text: str, width: int, subtitle: str | None = None) -> int:
    sheet.cell(1, 1, text).font = Font(bold=True, size=14)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, width))
    sheet.cell(1, 1).alignment = Alignment(horizontal="center")
    if subtitle:
        sheet.cell(2, 1, subtitle).font = Font(italic=True, color="595959")
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(1, width))
    return 3


def _page_setup(sheet: Worksheet, freeze: str) -> None:
    sheet.freeze_panes = freeze
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0


def _widths(sheet: Worksheet, widths: Mapping[int, float]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width


def _write_grouped_table(
    sheet: Worksheet,
    start_row: int,
    headers: Sequence[str],
    groups: Mapping[str, Sequence[Sequence[Any]]],
    *,
    sum_columns: Iterable[int],
    formats: Mapping[int, str],
    label_span: int,
    group_label: str = "{name}",
    subtotal_label: str = "{name} — дүн",
    grand_label: str = "НИЙТ ДҮН",
) -> int:
    """Write group rows, data rows and SUBTOTAL rows; return the next free row."""
    sum_columns = list(sum_columns)
    width = len(headers)
    row_index = start_row
    data_start = row_index
    for name, rows in groups.items():
        sheet.cell(row_index, 1, group_label.format(name=name))
        sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=width)
        _style_range(sheet, row_index, 1, width, bold=True, fill=GROUP_FILL)
        row_index += 1
        first = row_index
        for values in rows:
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row_index, column, value)
                cell.border = BORDER
                if column in formats:
                    cell.number_format = formats[column]
            row_index += 1
        last = row_index - 1
        sheet.cell(row_index, 1, subtotal_label.format(name=name))
        if label_span > 1:
            sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=label_span)
        for column in sum_columns:
            letter = get_column_letter(column)
            cell = sheet.cell(row_index, column, f"=SUBTOTAL(9,{letter}{first}:{letter}{last})" if last >= first else 0)
            cell.number_format = formats.get(column, MONEY)
        _style_range(sheet, row_index, 1, width, bold=True)
        row_index += 1
    data_end = row_index - 1
    sheet.cell(row_index, 1, grand_label)
    if label_span > 1:
        sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=label_span)
    for column in sum_columns:
        letter = get_column_letter(column)
        cell = sheet.cell(row_index, column, f"=SUBTOTAL(9,{letter}{data_start}:{letter}{data_end})" if data_end >= data_start else 0)
        cell.number_format = formats.get(column, MONEY)
    _style_range(sheet, row_index, 1, width, bold=True, fill=TOTAL_FILL)
    return row_index + 1


FINAL_COLUMNS = (
    "№", "Овог", "Нэр", "РД", "Албан тушаал", "Үндсэн цалин", "Өдөр", "Цаг", "Ажилласан цаг", "Тооцсон цалин",
    "Илүү цаг", "Илүү цагийн хөлс", "Ээлжийн амралтын мөнгө", "Хоол унаа", "Урамшуулал", "Олговол зохих цалин",
    "НДШ", "ХХОАТ ХӨН", "ХХОАТ", "Урьдчилгаа", "Бусад суутгал", "Суутгалын дүн", "Гарт олгох цалин", "БНДШ",
)
FINAL_GROUPS = {(7, 8): "Ажиллах", (17, 21): "Суутгалууд"}


def final_register_values(index: int, row: Mapping[str, Any]) -> list[Any]:
    identity, profile, inputs, result = row["identity"], row["profile"], row["inputs"], row["result"]
    last, first = split_name(identity)
    overtime_hours = sum((_d(value) for value in (inputs.get("overtime_hours") or {}).values()), Decimal("0"))
    return [
        index, last, first, identity.get("rd") or "", identity.get("job_title") or "",
        _num(profile.get("base_salary")), _num(result.get("planned_days")), _num(result.get("planned_hours")),
        _num(inputs.get("worked_normal_hours")), _num(result.get("base_pay")), _num(overtime_hours), _num(result.get("overtime_pay")),
        _num(inputs.get("leave_pay")), _num(result.get("meal_commute")), _num(inputs.get("bonus")), _num(result.get("gross")),
        _num(result.get("employee_shi")), _num(result.get("relief")), _num(result.get("pit")), _num(result.get("advance")),
        _num(result.get("other_deductions")), _num(result.get("total_deductions")), _num(result.get("net_pay")), _num(result.get("employer_shi")),
    ]


def _final_register(workbook: Workbook, rows: Sequence[Mapping[str, Any]], company: str | None, year: int, month: int) -> None:
    sheet = workbook.create_sheet("Цалингийн хүснэгт")
    width = len(FINAL_COLUMNS)
    header_row = _title(sheet, company_title(company, f"{year} оны {month:02d} сарын цалингийн хүснэгт"), width)
    grouped = {column for span in FINAL_GROUPS for column in range(span[0], span[1] + 1)}
    for column, label in enumerate(FINAL_COLUMNS, start=1):
        if column in grouped:
            sheet.cell(header_row + 1, column, label)
        else:
            sheet.cell(header_row, column, label)
            sheet.merge_cells(start_row=header_row, start_column=column, end_row=header_row + 1, end_column=column)
    for (start, end), label in FINAL_GROUPS.items():
        sheet.cell(header_row, start, label)
        sheet.merge_cells(start_row=header_row, start_column=start, end_row=header_row, end_column=end)
    for row_index in (header_row, header_row + 1):
        _style_range(sheet, row_index, 1, width, bold=True, fill=HEADER_FILL)
        for column in range(1, width + 1):
            sheet.cell(row_index, column).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[header_row + 1].height = 42
    counter = 0
    groups: OrderedDict[str, list[list[Any]]] = OrderedDict()
    for department, department_rows in _by_department(rows).items():
        groups[department] = []
        for row in department_rows:
            counter += 1
            groups[department].append(final_register_values(counter, row))
    formats = {column: MONEY for column in range(6, width + 1)}
    formats.update({7: "0", 8: HOURS, 9: HOURS, 11: HOURS})
    next_row = _write_grouped_table(sheet, header_row + 2, FINAL_COLUMNS, groups, sum_columns=range(6, width + 1), formats=formats, label_span=5)
    sheet.cell(next_row + 1, 1, "БНДШ – Байгууллагын төлөх нийгмийн даатгалын шимтгэл").font = Font(italic=True)
    _widths(sheet, {1: 5, 2: 14, 3: 14, 4: 12, 5: 18, **{column: 13 for column in range(6, width + 1)}, 7: 7, 8: 8})
    _page_setup(sheet, f"D{header_row + 2}")


def _advance_register(workbook: Workbook, rows: Sequence[Mapping[str, Any]], company: str | None, pay_date: date) -> None:
    headers = ("№", "Овог", "Нэр", "РД", "Албан тушаал", "Үндсэн цалин", "Цалингийн төрөл", "Урьдчилгааны суурь", "Хувь / дүн", "Ажилласан цаг", "Урьдчилгаа", "Төлбөрийн өдөр")
    sheet = workbook.create_sheet("Урьдчилгаа")
    header_row = _title(sheet, company_title(company, f"{pay_date.month:02d} сарын урьдчилгаа цалин ({pay_date.day:02d} өдөр)"), len(headers))
    for column, label in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column, label)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    _style_range(sheet, header_row, 1, len(headers), bold=True, fill=HEADER_FILL)
    counter = 0
    groups: OrderedDict[str, list[list[Any]]] = OrderedDict()
    for department, department_rows in _by_department(rows).items():
        groups[department] = []
        for row in department_rows:
            counter += 1
            identity, profile, inputs, result = row["identity"], row["profile"], row["inputs"], row["result"]
            last, first = split_name(identity)
            basis = result.get("advance_basis") or profile.get("advance_basis")
            value = result.get("advance_value")
            groups[department].append([
                counter, last, first, identity.get("rd") or "", identity.get("job_title") or "",
                _num(profile.get("base_salary")), SALARY_TYPE_LABELS.get(profile.get("salary_type"), profile.get("salary_type") or ""),
                BASIS_LABELS.get(basis, basis or ""), (f"{_num(value)}%" if basis == "PERCENT" else _num(value)) if value not in (None, "") else "",
                _num(inputs.get("worked_to_date_hours")) if basis == "WORKED-TO-DATE" else "",
                _num(result.get("advance")), identity.get("pay_date") or pay_date.isoformat(),
            ])
    _write_grouped_table(sheet, header_row + 1, headers, groups, sum_columns=(6, 11), formats={6: MONEY, 9: MONEY, 10: HOURS, 11: MONEY}, label_span=5)
    _widths(sheet, {1: 5, 2: 14, 3: 14, 4: 12, 5: 18, 6: 14, 7: 12, 8: 18, 9: 11, 10: 12, 11: 14, 12: 13})
    _page_setup(sheet, f"D{header_row + 1}")


def _payment_list(workbook: Workbook, rows: Sequence[Mapping[str, Any]], amount_key: str, title: str) -> None:
    headers = ("№", "Ажилтан", "РД", "Банк", "Дансны дугаар", "Дүн", "Тэмдэглэл")
    sheet = workbook.create_sheet("Төлбөрийн жагсаалт")
    header_row = _title(sheet, title, len(headers))
    for column, label in enumerate(headers, start=1):
        sheet.cell(header_row, column, label)
    _style_range(sheet, header_row, 1, len(headers), bold=True, fill=HEADER_FILL)
    groups: OrderedDict[str, list[list[Any]]] = OrderedDict()
    counter = 0
    for row in sorted(rows, key=lambda item: (item["identity"].get("pay_date") or "", item["identity"].get("name") or "")):
        payout = row.get("payout") or {}
        counter += 1
        groups.setdefault(row["identity"].get("pay_date") or "", []).append([
            counter, row["identity"].get("name") or "", row["identity"].get("rd") or "", payout.get("bank_code", ""),
            payout.get("account_number", ""), _num(row["result"].get(amount_key)), "" if payout else "Банкны мэдээлэл дутуу",
        ])
    _write_grouped_table(sheet, header_row + 1, headers, groups, sum_columns=(6,), formats={6: MONEY}, label_span=5, group_label="Төлбөрийн өдөр: {name}", subtotal_label="{name} — дүн")
    _widths(sheet, {1: 5, 2: 26, 3: 12, 4: 10, 5: 20, 6: 14, 7: 24})
    _page_setup(sheet, f"A{header_row + 1}")


def _summary(workbook: Workbook, stats: Mapping[str, Any], title: str) -> None:
    sheet = workbook.create_sheet("Дүн")
    row_index = _title(sheet, title, 4)
    totals, headcount = stats.get("totals") or {}, stats.get("headcount") or {}

    def section(label: str) -> None:
        nonlocal row_index
        row_index += 1
        sheet.cell(row_index, 1, label).font = Font(bold=True, size=12)
        row_index += 1

    def line(label: str, value: Any, *, indent: int = 0, fmt: str = MONEY, bold: bool = False) -> None:
        nonlocal row_index
        sheet.cell(row_index, 1, ("    " * indent) + label).font = Font(bold=bold)
        cell = sheet.cell(row_index, 2, _num(value) if not isinstance(value, str) else value)
        cell.number_format = fmt
        cell.font = Font(bold=bold)
        row_index += 1

    section("Мөнгөн дүн")
    line("Олговол зохих цалин", totals.get("gross"), bold=True)
    for label, key in (("Тооцсон цалин", "base_pay"), ("Илүү цагийн хөлс", "overtime_pay"), ("Ээлжийн амралтын мөнгө", "leave_pay"), ("Хоол унаа", "meal_commute"), ("Урамшуулал", "bonus")):
        line(label, totals.get(key), indent=1)
    line("ХХОАТ (хөнгөлөлтийн өмнө)", totals.get("pit_before_relief"))
    line("ХХОАТ ХӨН", totals.get("relief"))
    line("ХХОАТ", totals.get("pit"))
    line("Ажилтны НДШ", totals.get("employee_shi"))
    line("БНДШ", totals.get("employer_shi"))
    line("Урьдчилгаа (батлагдсан бодолтууд)", totals.get("advance_total"))
    for item in stats.get("advance_by_run") or []:
        line(f"{item.get('pay_date')} · {item.get('status')}{' · чөлөөлсөн' if item.get('waived') else ''}", item.get("total"), indent=1)
    line("Бусад суутгал", totals.get("other_deductions"))
    for label, value in (stats.get("other_deductions_by_type") or {}).items():
        line(label, value, indent=1)
    line("Сүүл цалин", totals.get("net_pay"), bold=True)
    line("Нийт зардал (олговол зохих + БНДШ)", totals.get("company_cost"), bold=True)

    section("Толгой тоо")
    for label, key in (("Хүснэгтэд", "on_register"), ("Шинээр орсон", "new"), ("Гарсан", "left"), ("Илүү цагтай", "with_extra_work"), ("Гараар зассан", "manually_edited")):
        line(label, headcount.get(key), fmt="0")
    averages = stats.get("averages") or {}
    section("Дундаж ба муж")
    for label, key in (("Дундаж олговол зохих", "average_gross"), ("Медиан олговол зохих", "median_gross"), ("Дундаж гарт олгох", "average_take_home"), ("Хамгийн их", "highest_gross"), ("Хамгийн бага", "lowest_gross")):
        line(label, averages.get(key))

    accounting = stats.get("accounting") or {}
    section("Нягтлан бодох бүртгэл")

    def account_name(item: Mapping[str, Any] | None, fallback: str) -> str:
        account = (item or {}).get("account")
        return f"{fallback} · {account.get('code')} {account.get('name')}" if account and account.get("code") else f"{fallback} · данс сонгоогүй"

    account_a = accounting.get("account_a") or {}
    line(account_name(account_a, "A · Цалингийн зардал"), account_a.get("total"), bold=True)
    lines_a = account_a.get("lines") or {}
    for label, key in (("Урьдчилгаа", "advance"), ("ХХОАТ", "pit"), ("Ажилтны НДШ", "employee_shi")):
        line(label, lines_a.get(key), indent=1)
    for label, value in (lines_a.get("other_deductions") or {}).items():
        line(f"Бусад суутгал · {label}", value, indent=1)
    line("Сүүл цалин", lines_a.get("net_pay"), indent=1)
    line(account_name(accounting.get("account_b"), "B · Байгууллагын НДШ"), (accounting.get("account_b") or {}).get("total"), bold=True)
    if accounting.get("account_c"):
        line(account_name(accounting.get("account_c"), "C · Урьдчилгаа цалин (тооцоо)"), accounting["account_c"].get("total"), bold=True)
    reconciliation = accounting.get("advance_reconciliation") or {}
    line("Урьдчилгааны тулгалт", "Таарсан" if reconciliation.get("matches", True) else "ТААРАХГҮЙ", fmt="@")

    departments = stats.get("by_department") or {}
    if departments:
        section("Хэлтсээр")
        headers = ("Хэлтэс", "Ажилтан", "Олговол зохих", "НДШ", "ХХОАТ", "БНДШ", "Нийт зардал", "Хувь (%)")
        for column, label in enumerate(headers, start=1):
            sheet.cell(row_index, column, label)
        _style_range(sheet, row_index, 1, len(headers), bold=True, fill=HEADER_FILL)
        row_index += 1
        for name, group in departments.items():
            values = [name, _num(group.get("headcount")), _num(group.get("gross")), _num(group.get("employee_shi")), _num(group.get("pit")), _num(group.get("employer_shi")), _num(group.get("company_cost")), _num(group.get("share"))]
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row_index, column, value)
                cell.border = BORDER
                cell.number_format = "0" if column == 2 else "0.0" if column == 8 else MONEY
            row_index += 1

    comparison = stats.get("comparison")
    if comparison:
        section(f"Өмнөх хаасан сартай харьцуулалт ({comparison.get('month')})")
        labels = {"gross": "Олговол зохих", "company_cost": "Нийт зардал", "headcount": "Ажилтан", "overtime_amount": "Илүү цагийн хөлс"}
        for key, metric in (comparison.get("metrics") or {}).items():
            pct = metric.get("change_pct")
            line(f"{labels.get(key, key)}: {_num(metric.get('previous')):,} → {_num(metric.get('current')):,}{f' ({pct}%)' if pct is not None else ''}", metric.get("change"))
    _widths(sheet, {1: 46, 2: 16, 3: 14, 4: 14, 5: 14, 6: 14, 7: 14, 8: 10})


def _overtime_sheet(workbook: Workbook, rows: Sequence[Mapping[str, Any]], title: str) -> None:
    headers = ("Ажилтан", "Огноо", "Гараг", "Өдрийн төрөл", "Ангилал", "Цаг", "Үржүүлэгч", "Цагийн үнэлгээ", "Дүн", "Эх сурвалж")
    sheet = workbook.create_sheet("Илүү цаг")
    header_row = _title(sheet, title, len(headers), "Цагийн үнэлгээ = үндсэн цалин ÷ сарын ажиллах цаг (Хөдөлмөрийн тухай хууль 109.1, 109.2, 109.4)")
    for column, label in enumerate(headers, start=1):
        sheet.cell(header_row, column, label)
    _style_range(sheet, header_row, 1, len(headers), bold=True, fill=HEADER_FILL)
    groups: OrderedDict[str, list[list[Any]]] = OrderedDict()
    for row in sorted(rows, key=lambda item: item["identity"].get("name") or ""):
        for line in row["result"].get("overtime_lines") or []:
            weekday = line.get("weekday")
            groups.setdefault(row["identity"].get("name") or "", []).append([
                row["identity"].get("name") or "", line.get("date") or "Гараар", WEEKDAYS[int(weekday)] if weekday not in (None, "") else "",
                DAY_TYPE_LABELS.get(line.get("day_type"), line.get("day_type") or ""), BUCKET_LABELS.get(line.get("bucket"), line.get("bucket")),
                _num(line.get("hours")), _num(line.get("multiplier")), _num(line.get("rate")), _num(line.get("amount")),
                "Гараар" if line.get("source") == "manual" else "Цагийн бүртгэл",
            ])
    _write_grouped_table(sheet, header_row + 1, headers, groups, sum_columns=(6, 9), formats={6: HOURS, 7: "0.0#", 8: "#,##0.00", 9: MONEY}, label_span=5, subtotal_label="{name} — дүн")
    _widths(sheet, {1: 24, 2: 12, 3: 10, 4: 20, 5: 22, 6: 8, 7: 10, 8: 14, 9: 14, 10: 16})
    _page_setup(sheet, f"A{header_row + 1}")


def _shi_sheet(workbook: Workbook, rows: Sequence[Mapping[str, Any]], rule_snapshot: Mapping[str, Any], title: str) -> None:
    employee_rates = rule_snapshot.get("employee_rates") or {}
    employer_rates = rule_snapshot.get("employer_rates") or {}
    headers = ["Ажилтан", "НДШ суурь", *[f"Ажилтан · {FUND_LABELS.get(code, code)} ({_num(_d(rate) * 100)}%)" for code, rate in employee_rates.items()], "Ажилтны НДШ", *[f"Байгууллага · {FUND_LABELS.get(code, code)} ({_num(_d(rate) * 100)}%)" for code, rate in employer_rates.items()], "БНДШ"]
    sheet = workbook.create_sheet("НДШ задаргаа")
    header_row = _title(sheet, title, len(headers), "Сан тус бүрийн дүнг тусад нь бүхэлчилсэн тул нийлбэр нь нийт дүнгээс ±1₮ зөрж болно.")
    for column, label in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column, label)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    _style_range(sheet, header_row, 1, len(headers), bold=True, fill=HEADER_FILL)
    sheet.row_dimensions[header_row].height = 45
    data = []
    for row in sorted(rows, key=lambda item: item["identity"].get("name") or ""):
        base = _d(row["result"].get("shi_base"))
        data.append([
            row["identity"].get("name") or "", _num(base),
            *[_num((base * _d(rate)).quantize(Decimal("1"))) for rate in employee_rates.values()], _num(row["result"].get("employee_shi")),
            *[_num((base * _d(rate)).quantize(Decimal("1"))) for rate in employer_rates.values()], _num(row["result"].get("employer_shi")),
        ])
    _write_grouped_table(sheet, header_row + 1, headers, {"Бүх ажилтан": data}, sum_columns=range(2, len(headers) + 1), formats={column: MONEY for column in range(2, len(headers) + 1)}, label_span=1)
    _widths(sheet, {1: 26, **{column: 15 for column in range(2, len(headers) + 1)}})
    _page_setup(sheet, f"B{header_row + 1}")


def _deductions_sheet(workbook: Workbook, rows: Sequence[Mapping[str, Any]], title: str) -> None:
    headers = ("Ажилтан", "Хэлтэс", "Төрөл", "Дүн", "Тайлбар")
    sheet = workbook.create_sheet("Бусад суутгал")
    header_row = _title(sheet, title, len(headers), "Татвар, НДШ-ийн дараа суутгана; НДШ, ХХОАТ-ын суурийг өөрчлөхгүй.")
    for column, label in enumerate(headers, start=1):
        sheet.cell(header_row, column, label)
    _style_range(sheet, header_row, 1, len(headers), bold=True, fill=HEADER_FILL)
    groups: OrderedDict[str, list[list[Any]]] = OrderedDict()
    for row in sorted(rows, key=lambda item: item["identity"].get("name") or ""):
        for line in row["inputs"].get("other_deductions") or []:
            groups.setdefault(line.get("type") or "Төрөлгүй", []).append([row["identity"].get("name") or "", row["identity"].get("department") or "", line.get("type") or "", _num(line.get("amount")), line.get("note") or ""])
    _write_grouped_table(sheet, header_row + 1, headers, groups, sum_columns=(4,), formats={4: MONEY}, label_span=3)
    _widths(sheet, {1: 26, 2: 18, 3: 28, 4: 14, 5: 36})


def build_run_workbook(
    *,
    run_type: str,
    pay_date: date,
    year: int,
    month: int,
    rows: Sequence[Mapping[str, Any]],
    company: str | None,
    rule_snapshot: Mapping[str, Any],
    stats: Mapping[str, Any] | None = None,
) -> bytes:
    """Build the advance (§12.1) or final (§12.2) workbook.

    Each row is ``{"identity", "profile", "inputs", "result", "payout"}`` with
    ``payout`` the decrypted bank details or ``None``.
    """
    workbook = Workbook()
    workbook.remove(workbook.active)
    if run_type == "advance":
        _advance_register(workbook, rows, company, pay_date)
        _payment_list(workbook, rows, "advance", company_title(company, f"{month:02d} сарын урьдчилгааны төлбөрийн жагсаалт ({pay_date.day:02d} өдөр)"))
    else:
        _final_register(workbook, rows, company, year, month)
        _summary(workbook, stats or {}, company_title(company, f"{year} оны {month:02d} сарын цалингийн дүн"))
        _overtime_sheet(workbook, rows, company_title(company, f"{month:02d} сарын илүү цаг, амралт, баярын өдрийн ажил"))
        _shi_sheet(workbook, rows, rule_snapshot, company_title(company, f"{month:02d} сарын НДШ задаргаа"))
        _payment_list(workbook, rows, "net_pay", company_title(company, f"{month:02d} сарын сүүл цалингийн төлбөрийн жагсаалт"))
        _deductions_sheet(workbook, rows, company_title(company, f"{month:02d} сарын бусад суутгал"))
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()
